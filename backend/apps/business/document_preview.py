"""Bounded, read-only attachment previews. Never evaluate formulas or follow links."""

import csv
import io
import math
import re
import zipfile
from contextlib import ExitStack
from datetime import date, datetime, time, timedelta
from itertools import islice
from pathlib import Path

import xlrd
from defusedxml import ElementTree
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from rest_framework.exceptions import NotFound, ValidationError

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000
MAX_XML_ELEMENTS = 1000000
MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_BLOCKS = 5000
MAX_TEXT = 200000
MAX_CELL_TEXT = 4000
PAGE_SIZE = 100
WORD = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
SPREADSHEET = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
LAYOUT_WARNING = '预览仅展示文字和单元格值，不还原原始排版、合并单元格、图表、图片或嵌入对象；完整内容请下载原文件。'
FORMULA_WARNING = '公式仅显示文件中已保存的计算结果，可能不是最新值；预览不会执行公式或更新外部链接。'
UNSUPPORTED_MESSAGE = '暂不支持此格式在线预览。请使用原软件导出 PDF 后上传，或下载原文件查看。'


def _warn(warnings, message):
    if message not in warnings:
        warnings.append(message)


class TextBudget:
    """Limit both individual cells and the total text returned in one response."""

    def __init__(self, warnings):
        self.remaining = MAX_TEXT
        self.warnings = warnings

    def take(self, value, limit=MAX_CELL_TEXT):
        value = str(value)
        size = min(limit, self.remaining)
        if len(value) > size:
            _warn(self.warnings, '部分文字过长，预览已截断；完整内容请下载原文件。')
        result = value[:size]
        self.remaining -= len(result)
        return result


def read_content(document):
    try:
        with document.file.open('rb') as stream:
            content = stream.read(MAX_BYTES + 1)
    except (OSError, ValueError) as exc:
        raise NotFound('附件文件暂不可用，请联系管理员检查或恢复备份。') from exc
    if not content or len(content) > MAX_BYTES:
        raise ValidationError('附件为空或超过 20 MB，无法预览；请检查原文件。')
    return content


def _integer(query, key, default, minimum=0):
    value = query.get(key)
    if value is None:
        return default
    if hasattr(query, 'getlist') and len(query.getlist(key)) != 1:
        raise ValidationError('工作表和页码参数只能提交一次。')
    if not re.fullmatch(r'[0-9]{1,6}', str(value)) or int(value) < minimum:
        raise ValidationError('工作表必须是非负整数，页码必须是从 1 开始的正整数。')
    return int(value)


def _pagination(page, count):
    pages = max(1, math.ceil(count / PAGE_SIZE))
    if page > pages:
        raise ValidationError('页码超出可预览范围，请返回第一页或选择有效页码。')
    return {'page': page, 'page_size': PAGE_SIZE, 'total_pages': pages}


def _sheet(sheets, index):
    if not sheets:
        raise ValidationError('文件没有可见工作表，请在原软件中取消隐藏后重新上传。')
    if index is None:
        return sheets[0]['index']
    if index not in {sheet['index'] for sheet in sheets}:
        raise ValidationError('工作表不存在或已隐藏，请选择列表中的工作表。')
    return index


def _check_zip(content, warnings):
    """Validate sizes and every XML part before any office parser reads it."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if (
            len(infos) > MAX_ZIP_MEMBERS
            or sum(info.file_size for info in infos) > MAX_EXPANDED_BYTES
            or any(info.file_size > MAX_BYTES for info in infos)
        ):
            raise ValidationError('文件解压后过大或部件过多，无法安全预览；请拆分文件或导出 PDF。')
        if len({info.filename for info in infos}) != len(infos):
            raise ValidationError('文件包含重复部件，无法预览；请重新保存后上传。')
        elements = 0
        for info in infos:
            if info.flag_bits & 1:
                raise ValidationError('加密文件无法预览，请去除密码后重新上传。')
            if not info.filename.lower().endswith(('.xml', '.rels')):
                continue
            with archive.open(info) as stream:
                for _, element in ElementTree.iterparse(
                    stream, forbid_dtd=True, forbid_entities=True, forbid_external=True
                ):
                    elements += 1
                    if elements > MAX_XML_ELEMENTS:
                        raise ValidationError('文档结构过于复杂，无法安全预览；请拆分文件或导出 PDF。')
                    if element.tag == SPREADSHEET + 'c':
                        coordinate = re.fullmatch(r'([A-Z]+)[0-9]+', element.get('r', ''))
                        if coordinate and column_index_from_string(coordinate[1]) > MAX_COLUMNS:
                            _warn(warnings, f'表格超过 {MAX_COLUMNS} 列，预览仅显示前 {MAX_COLUMNS} 列。')
                    element.clear()


def _display(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, datetime):
        if value.time() == time.min and not re.search(r'[hs]', number_format, re.I):
            return value.date().isoformat()
        return value.isoformat(sep=' ', timespec='seconds')
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        # Common Excel formats only; arbitrary custom formats remain plain numeric text.
        pattern = number_format.split(';')[0]
        if len(pattern) > 100:
            return str(value)
        if re.fullmatch(r'0+', pattern):
            return f'{value:0{len(pattern)}.0f}'
        if re.fullmatch(r'[#0,]+(?:\.0+)?%?', pattern):
            decimals = len(pattern.split('.')[1].rstrip('%')) if '.' in pattern else 0
            suffix = '%' if pattern.endswith('%') else ''
            grouping = ',' if ',' in pattern else ''
            return f'{value * (100 if suffix else 1):{grouping}.{decimals}f}{suffix}'
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    return str(value)


def _spreadsheet_result(rows, columns, total, sheets, index, page):
    return {
        'kind': 'spreadsheet',
        'sheets': sheets,
        'sheet': index,
        **_pagination(page, total),
        'total_rows': total,
        'columns': [get_column_letter(column) for column in range(1, columns + 1)],
        'rows': [{'number': number, 'cells': cells + [''] * (columns - len(cells))} for number, cells in rows],
    }


def _xlsx(content, index, page, warnings):
    _check_zip(content, warnings)
    _warn(warnings, FORMULA_WARNING)
    budget = TextBudget(warnings)
    with ExitStack() as stack:
        books = []
        for data_only in (False, True):
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=data_only, keep_links=False)
            stack.callback(book.close)
            books.append(book)
        sheets = [
            {'index': number, 'name': sheet.title}
            for number, sheet in enumerate(books[0].worksheets)
            if sheet.sheet_state == 'visible'
        ]
        index = _sheet(sheets, index)
        iterators = []
        for book in books:
            sheet = book.worksheets[index]
            # Some exporters write incorrect dimensions; count the actual bounded rows.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(max_col=MAX_COLUMNS)
            stack.callback(iterator.close)
            iterators.append(iterator)
        total, width, rows = 0, 0, []
        for number, (source, cached) in enumerate(islice(zip(*iterators, strict=True), MAX_ROWS + 1), 1):
            if number > MAX_ROWS:
                _warn(warnings, f'表格超过 {MAX_ROWS} 行，预览仅显示前 {MAX_ROWS} 行。')
                break
            total = number
            for column, (cell, value) in enumerate(zip(source, cached, strict=True), 1):
                if cell.value is not None:
                    width = max(width, column)
                if cell.data_type == 'f' and value.value is None:
                    _warn(warnings, '工作表含尚未计算的公式，相关单元格显示为空；请在 Excel 中重新计算并保存后上传。')
            if (page - 1) * PAGE_SIZE < number <= page * PAGE_SIZE:
                rows.append((number, [budget.take(_display(cell.value, cell.number_format)) for cell in cached]))
        width = max(1, width)
        return _spreadsheet_result([(n, cells[:width]) for n, cells in rows], width, total, sheets, index, page)


def _xls(content, index, page, warnings):
    _warn(warnings, FORMULA_WARNING)
    _warn(warnings, '旧版 XLS 仅能读取缓存值，无法识别所有未计算公式；如有空值，请重新计算并保存为 XLSX 后上传。')
    budget = TextBudget(warnings)
    book = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True, ragged_rows=True)
    try:
        sheets = []
        for number, name in enumerate(book.sheet_names()):
            sheet = book.sheet_by_index(number)
            if sheet.visibility == 0:
                sheets.append({'index': number, 'name': name})
            book.unload_sheet(number)
        index = _sheet(sheets, index)
        sheet = book.sheet_by_index(index)
        total, width = min(sheet.nrows, MAX_ROWS), max(1, min(sheet.ncols, MAX_COLUMNS))
        if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
            _warn(warnings, f'表格过大，预览仅显示前 {MAX_ROWS} 行、{MAX_COLUMNS} 列。')
        _pagination(page, total)
        rows = []
        for row in range((page - 1) * PAGE_SIZE, min(page * PAGE_SIZE, total)):
            cells = []
            for column in range(min(width, sheet.ncols)):
                if column >= sheet.row_len(row):
                    cells.append('')
                    continue
                cell = sheet.cell(row, column)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(value, '单元格错误')
                number_format = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                cells.append(budget.take(_display(value, number_format)))
            rows.append((row + 1, cells))
        return _spreadsheet_result(rows, width, total, sheets, index, page)
    finally:
        book.release_resources()


def _decode(content, warnings):
    try:
        result = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            result = content.decode('gb18030')
        except UnicodeDecodeError as exc:
            raise ValidationError('无法识别文件文字编码，请保存为 UTF-8 或 GB18030 后重新上传。') from exc
        _warn(warnings, '文件按 GB18030 编码读取；如有乱码，请保存为 UTF-8 后重新上传。')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', result):
        raise ValidationError('文件包含非文本内容，请检查格式或导出 PDF 后上传。')
    return result


def _csv(content, index, page, warnings):
    sheets = [{'index': 0, 'name': 'CSV'}]
    index = _sheet(sheets, index)
    budget = TextBudget(warnings)
    rows, total, width = [], 0, 1
    reader = csv.reader(io.StringIO(_decode(content, warnings), newline=''), strict=True)
    for number, row in enumerate(islice(reader, MAX_ROWS + 1), 1):
        if number > MAX_ROWS:
            _warn(warnings, f'表格超过 {MAX_ROWS} 行，预览仅显示前 {MAX_ROWS} 行。')
            break
        total = number
        if len(row) > MAX_COLUMNS:
            _warn(warnings, f'表格超过 {MAX_COLUMNS} 列，预览仅显示前 {MAX_COLUMNS} 列。')
        width = max(width, min(len(row), MAX_COLUMNS))
        if (page - 1) * PAGE_SIZE < number <= page * PAGE_SIZE:
            rows.append((number, [budget.take(cell) for cell in row[:MAX_COLUMNS]]))
    return _spreadsheet_result(rows, width, total, sheets, index, page)


def _word_text(element):
    return ''.join(
        node.text or '' if node.tag == WORD + 't' else '\t' if node.tag == WORD + 'tab' else '\n'
        for node in element.iter()
        if node.tag in {WORD + 't', WORD + 'tab', WORD + 'br', WORD + 'cr'}
    )


def _docx(content, page, warnings):
    _check_zip(content, warnings)
    budget = TextBudget(warnings)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = ElementTree.fromstring(
            archive.read('word/document.xml'), forbid_dtd=True, forbid_entities=True, forbid_external=True
        )
    body = root.find(WORD + 'body')
    if body is None:
        raise ValidationError('DOCX 缺少正文，无法预览；请重新保存后上传。')
    blocks = [child for child in body if child.tag in {WORD + 'p', WORD + 'tbl'}]
    if len(blocks) > MAX_BLOCKS:
        _warn(warnings, f'文档过长，预览仅包含前 {MAX_BLOCKS} 个段落或表格。')
        blocks = blocks[:MAX_BLOCKS]
    result = {'kind': 'document', **_pagination(page, len(blocks)), 'blocks': []}
    for block in blocks[(page - 1) * PAGE_SIZE : page * PAGE_SIZE]:
        if block.tag == WORD + 'p':
            result['blocks'].append({'kind': 'paragraph', 'text': budget.take(_word_text(block), MAX_TEXT)})
        else:
            rows = block.findall(WORD + 'tr')
            if len(rows) > PAGE_SIZE:
                _warn(warnings, f'部分表格过长，单个表格仅展示前 {PAGE_SIZE} 行。')
            values = []
            for row in rows[:PAGE_SIZE]:
                cells = row.findall(WORD + 'tc')
                if len(cells) > MAX_COLUMNS:
                    _warn(warnings, f'部分表格过宽，仅展示前 {MAX_COLUMNS} 列。')
                values.append(
                    [
                        budget.take('\n'.join(_word_text(p) for p in cell.iter(WORD + 'p')))
                        for cell in cells[:MAX_COLUMNS]
                    ]
                )
            result['blocks'].append({'kind': 'table', 'rows': values})
    return result


def inline_type(name, content):
    """Allow only known passive formats; never infer inline HTML/SVG from a filename."""
    extension = Path(name).suffix.lower()
    valid = {
        '.pdf': ('application/pdf', bool(re.match(rb'%PDF-[12]\.[0-9](?:\s|$)', content))),
        '.png': ('image/png', content.startswith(b'\x89PNG\r\n\x1a\n')),
        '.jpg': ('image/jpeg', content.startswith(b'\xff\xd8\xff')),
        '.jpeg': ('image/jpeg', content.startswith(b'\xff\xd8\xff')),
        '.gif': ('image/gif', content.startswith((b'GIF87a', b'GIF89a'))),
        '.webp': ('image/webp', content.startswith(b'RIFF') and content[8:12] == b'WEBP'),
        '.bmp': ('image/bmp', len(content) >= 14 and content.startswith(b'BM')),
    }
    if extension not in valid:
        raise ValidationError('此格式不能直接预览，请下载原文件或导出 PDF 后上传。')
    mime, matches = valid[extension]
    if not matches:
        raise ValidationError('文件内容与扩展名不一致或文件已损坏，请检查原文件后重新上传。')
    return mime


def preview(document, query):
    content = read_content(document)
    page = _integer(query, 'page', 1, minimum=1)
    index = _integer(query, 'sheet', None)
    extension = Path(document.original_name).suffix.lower()
    warnings = []
    result = {'name': document.original_name, 'warnings': warnings}
    if index is not None and extension not in {'.xlsx', '.xls', '.csv'}:
        raise ValidationError('此文件没有可选择的工作表，请移除工作表参数后重试。')
    try:
        if extension in {'.xlsx', '.xls', '.csv'}:
            _warn(warnings, LAYOUT_WARNING)
            result.update({'.xlsx': _xlsx, '.xls': _xls, '.csv': _csv}[extension](content, index, page, warnings))
        elif extension == '.docx':
            _warn(warnings, LAYOUT_WARNING)
            result.update(_docx(content, page, warnings))
        elif extension == '.txt':
            result.update(kind='text', text=TextBudget(warnings).take(_decode(content, warnings), MAX_TEXT))
            _pagination(page, 1)
        elif extension in {'.pdf', '.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp'}:
            mime = inline_type(document.original_name, content)
            _pagination(page, 1)
            result.update(
                kind='pdf' if mime == 'application/pdf' else 'image',
                file_url=f'/api/business/documents/{document.pk}/preview-file/',
            )
        else:
            _pagination(page, 1)
            result.update(kind='unsupported', message=UNSUPPORTED_MESSAGE)
    except ValidationError:
        raise
    except Exception as exc:
        # Office parsers expose many implementation-specific errors for malformed input.
        # Keep this boundary around parsing only, never around permission checks or DB access.
        raise ValidationError('文件已损坏、已加密或格式不受支持，无法预览；请在原软件中重新保存或导出 PDF。') from exc
    return result
