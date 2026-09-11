import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from django.db.models import F, Sum
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

from ..models import BOMLine, Item, PurchaseLine
from .bom import incoming, issued, revision
from .common import day, number, state

MAX_BYTES = 5 * 1024 * 1024
LEGACY_HEADERS = ['物料编码', '数量', '变更说明', '单元']
PREVIOUS_HEADERS = ['物料编码', '单元', '需求数量', '变更说明']
HEADERS = PREVIOUS_HEADERS + ['需求日期', '申请日期', '申请人']


def read_file(upload, headers=None, legacy_headers=None, *, return_headers=False):
    if headers is None:
        legacy_headers = [PREVIOUS_HEADERS, LEGACY_HEADERS, LEGACY_HEADERS[:3]]
    headers = HEADERS if headers is None else headers
    if upload is None or not hasattr(upload, 'read'):
        raise ValidationError({'file': '请选择 CSV 或 XLSX 文件。'})
    content = upload.read(MAX_BYTES + 1)
    if not content or len(content) > MAX_BYTES:
        raise ValidationError({'file': '文件不能为空且不能超过 5 MB。'})
    extension = Path(upload.name).suffix.lower()
    try:
        if extension == '.csv':
            result = []
            for index, row in enumerate(csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True)):
                if index > 1000:
                    raise ValidationError({'file': '一次最多导入 1000 行。'})
                result.append(row)
        elif extension == '.xlsx':
            result = read_xlsx(content, len(headers))
        else:
            raise ValidationError({'file': '仅支持 UTF-8 CSV 或 XLSX。'})
    except (
        UnicodeDecodeError,
        csv.Error,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        DefusedXmlException,
        ParseError,
        InvalidFileException,
        IndexError,
        TypeError,
        RuntimeError,
        NotImplementedError,
        OSError,
        EOFError,
    ) as exc:
        raise ValidationError({'file': '文件损坏、编码不正确或不是受支持的表格。'}) from exc
    if not result:
        raise ValidationError({'file': '表格不能为空。'})
    header = [str(value).strip() if value is not None else '' for value in result[0]]
    while header and not header[-1]:
        header.pop()
    if (
        header != headers
        and header != legacy_headers
        and not (legacy_headers and isinstance(legacy_headers[0], list) and header in legacy_headers)
    ):
        raise ValidationError({'file': '表头必须依次为：' + '、'.join(headers) + '。'})
    if len(result) > 1001:
        raise ValidationError({'file': '一次最多导入 1000 行。'})
    normalized = []
    for row in result[1:]:
        if any(value not in (None, '') for value in row[len(header) :]):
            raise ValidationError({'file': '数据列数不能超过表头。'})
        normalized.append(list(row[: len(header)]) + [''] * max(0, len(header) - len(row)))
    return (header, normalized) if return_headers else normalized


def read_xlsx(content, columns=3):
    # Bound decompression and reject entity declarations before openpyxl sees XML.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > 200 or sum(info.file_size for info in infos) > 20 * 1024 * 1024:
            raise ValidationError({'file': 'XLSX 解压后过大或包含过多部件。'})
        for info in infos:
            if info.filename.lower().endswith(('.xml', '.rels')):
                root = ElementTree.fromstring(
                    archive.read(info), forbid_dtd=True, forbid_entities=True, forbid_external=True
                )
                if info.filename.startswith('xl/worksheets/'):
                    for row in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row'):
                        if int(row.get('r', '0')) > 1001:
                            raise ValidationError({'file': '一次最多导入 1000 行，请移除表格末尾多余行。'})
                    for cell in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c'):
                        column = re.match(r'[A-Z]+', cell.get('r', ''))
                        if (
                            column
                            and column.group() not in {get_column_letter(i) for i in range(1, columns + 1)}
                            and any(child.tag.rsplit('}', 1)[-1] in {'v', 'is', 'f'} for child in cell)
                        ):
                            raise ValidationError({'file': f'表格只能包含{columns}列。'})
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValidationError({'file': '请保留一个工作表后再导入。'})
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        result = []
        for cells in sheet.iter_rows(max_row=1002, max_col=columns + 1):
            if any(cell.data_type == 'f' for cell in cells):
                raise ValidationError({'file': '导入不接受公式，请粘贴为值后再导入。'})
            result.append([cell.value for cell in cells])
        while result and all(value in (None, '') for value in result[-1]):
            result.pop()
        return result
    finally:
        workbook.close()


def message(detail):
    if isinstance(detail, dict):
        return '；'.join(f'{key}：{message(value)}' for key, value in detail.items())
    if isinstance(detail, (list, tuple)):
        return '；'.join(message(value) for value in detail)
    return str(detail)


def preview(project, upload):
    state(project, {'draft', 'quoted', 'active', 'delivering'})
    expected_revision = revision(project)
    headers, raw = read_file(upload, return_headers=True)
    if headers in (HEADERS, PREVIOUS_HEADERS):
        raw = [[row[0], row[2], row[3], row[1], *row[4:]] for row in raw]
    codes = {str(row[0]).strip() for row in raw if row and row[0] is not None}
    items = {item.code: item for item in Item.objects.filter(code__in=codes, is_active=True)}
    existing = list(BOMLine.objects.filter(project=project))
    current = {(line.item_id, line.assembly_unit): line for line in existing}
    lines, errors, seen = [], [], set()
    for row_number, row in enumerate(raw, 2):
        if all(value in (None, '') for value in row):
            continue
        try:
            values = list(row[:3]) + [''] * max(0, 3 - len(row))
            code = str(values[0] or '').strip()
            item = items.get(code)
            if item is None:
                raise ValidationError('物料编码不存在或已停用。')
            note = str(values[2] or '').strip()
            assembly_unit = (
                str(row[3] or '').strip()
                if len(row) > 3
                else next((line.assembly_unit for line in existing if line.item_id == item.pk), '')
            )
            if len(row) < 4 and sum(line.item_id == item.pk for line in existing) > 1:
                raise ValidationError('此物料分属多个单元，请使用含单元列的新模板。')
            row_key = (item.pk, assembly_unit)
            if row_key in seen:
                raise ValidationError('同一文件中的物料编码和单元不能重复。')
            seen.add(row_key)
            qty = number(values[1], '数量', 3, positive=True)
            details = {}
            if headers == HEADERS:
                for field, value in zip(('required_date', 'application_date'), row[4:6], strict=True):
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    details[field] = day({field: value}, field, optional=True)
                    if details[field]:
                        details[field] = details[field].isoformat()
                details['applicant'] = str(row[6] or '').strip()
                if len(details['applicant']) > 80:
                    raise ValidationError('申请人不能超过80字符。')
            if len(assembly_unit) > 100:
                raise ValidationError('单元不能超过100字符。')
            if len(note) > 500 or (row_key in current and not note):
                raise ValidationError('修改已有 BOM 必须填写变更说明，且不超过 500 字符。')
            if row_key in current:
                pending = (
                    PurchaseLine.objects.filter(bom_line=current[row_key])
                    .exclude(purchase__status='cancelled')
                    .aggregate(total=Sum(F('quantity') - F('received_quantity') - F('cancelled_quantity')))['total']
                    or 0
                )
                if qty < pending:
                    raise ValidationError('单元用量不能小于该行在途采购数量。')
            lines.append(
                {
                    **details,
                    'row': row_number,
                    'item': item.pk,
                    'item_code': code,
                    'item_name': item.name,
                    'brand': item.brand,
                    'specification': item.specification,
                    'drawing_number': item.drawing_number,
                    'drawing_revision': item.drawing_revision,
                    'product_category': item.product_category,
                    'unit': item.unit,
                    'part_type': item.part_type,
                    'quantity': str(qty),
                    'change_note': note,
                    'assembly_unit': assembly_unit,
                }
            )
        except ValidationError as exc:
            errors.append({'row': row_number, 'message': message(exc.detail)})
    if not lines and not errors:
        errors.append({'row': 2, 'message': '文件没有可导入的明细。'})
    proposed = {key: line.quantity for key, line in current.items()}
    proposed.update({(line['item'], line['assembly_unit']): number(line['quantity'], 'quantity', 3) for line in lines})
    for item_id in {line['item'] for line in lines}:
        total = sum(qty for (pk, unit), qty in proposed.items() if pk == item_id)
        if total < issued(project, item_id) + incoming(project, item_id):
            errors.append(
                {
                    'row': next(line['row'] for line in lines if line['item'] == item_id),
                    'message': '物料各单元总量不能小于已领用和在途采购数量。',
                }
            )
    return {
        'expected_revision': expected_revision,
        'lines': lines,
        'errors': errors,
        'can_import': bool(lines) and not errors,
    }
