"""One employee-facing guide shared by template downloads and import dialogs."""

import io
from html import escape

from django.http import HttpResponse
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from rest_framework.exceptions import ValidationError

from ..models import Item, Partner, Task
from .tabular import document

GUIDE_SHEET = '填写说明'
GUIDE_MARKER = 'ERP导入模板填写说明'
COMMON = [
    '只在“数据”页第2行起填写，一行一条明细；保留表头名称和顺序。填写说明不参与导入，请勿在说明页填业务数据。',
    '优先使用Excel模板。选中数据单元格可看到填写提示；单击第2行起的空白格，再点右侧▼展开下拉，'
    '不是点击表头，也不需要双击。说明页的“点此选择”可直接跳到对应填写位置；单位也可自行输入。',
    '红色表头为必填，黄色表头为条件必填，蓝色表头为可选。开始填写一行后，已知必填的空格会标红；空白行不提示。',
    '下拉选项、输入提示和缺项标红覆盖第2～1001行。粘贴可能覆盖或绕过Excel校验，上传后仍须检查系统预览。',
    '说明中的姓名、账号、编码、ID和日期均为格式示例，必须换成实际值；示例不代表系统中已有记录。',
    '最多1000条数据、5MB。Excel可保留系统附带的填写说明页；不要增加其他工作表。旧版单工作表模板仍可用。',
    '日期填写YYYY-MM-DD，例如2026-10-20。不能填写20天、20或46290等普通数字；Excel日期单元格请设置为日期格式。',
    '人员可填姓名或登录账号，重名请填姓名（用户名）；编码、账号、电话、流水号按文本填写，保留前导零，不要把编号与记录ID混用。',
    '金额使用人民币含税金额，最多2位小数；数量最多3位小数。只填数值，不加元、件、千分位或计算公式。',
    'CSV使用UTF-8编码，只保存数据页；CSV不保留下拉选项、提示或标红，请配合下载的填写说明使用。'
    'Excel打开CSV时注意编码与长数字按文本读取。',
    '上传后先检查预览，再确认导入；任一行失败整批不保存。不要重复上传已经确认成功的文件，避免重复新增。',
]

# title, entry, purpose / prerequisites. These are documentation, not alternative validation rules.
RESOURCES = {
    'items': (
        '物料',
        '基础资料 → 物料 → 导入导出 → 导入',
        '新增物料，不覆盖已有物料。编码可自定义或留空自动生成；相同物料优先复用，确需独立建码应填写原因。',
    ),
    'partners': (
        '往来单位',
        '基础资料 → 往来单位 → 导入导出 → 导入',
        '新增客户、供应商或双用途往来单位，编号自动生成。采购账期只适用于供应商；销售经理单岗只能新增纯客户。',
    ),
    'sales': (
        '销售',
        '销售 → 导入导出 → 导入',
        '新增销售草稿，后续仍需报价和签约。仅销售经理权限时，负责人填写本人姓名或账号。',
    ),
    'projects': (
        '项目',
        '项目 → 项目列表 → 导入导出 → 导入',
        '新增执行项目。已签约销售请通过销售签约生成项目，避免重复；此入口不能导入任务或BOM。',
    ),
    'purchases': (
        '采购',
        '采购 → 导入导出 → 导入',
        '同一分组号的行合为一张采购草稿，仍需提交审批。同组的项目、供应商、交期、账期和说明必须一致；同组物料不能重复。'
        '该模板不关联BOM行；需要按BOM单元下单时请使用“从BOM多选下单”。',
    ),
    'tasks': (
        '项目任务',
        '打开对应项目 → 任务与工时 → 项目任务 → 导入导出 → 导入',
        '只新增设计、装配、调试任务，项目须为执行中或交付中。采购在采购模块办理，发货在交付模块办理；'
        '安装和验收任务随发货创建，售后从交付批次创建。T0、T1等编号不是阶段选项。',
    ),
    'stocks': (
        '期初库存',
        '设置 → 管理员初始化 → 期初库存 → 导入导出 → 导入',
        '仅管理员录入期初库存，只允许尚无库存流水的物料和库位；已有库存请使用盘点，不可重复导入期初。',
    ),
    'entries': (
        '项目费用',
        '收付款 → 应收应付与费用 → 导入导出 → 导入',
        '仅新增项目费用。合同应收、采购应付由原单据产生，不要通过此模板重复登记；项目须为执行中、交付中或质保阶段。',
    ),
    'payments': (
        '收付款记录',
        '收付款 → 收付款流水 → 导入导出 → 导入',
        '对已有款项登记正常收付款，金额须为正且不超过可收付余额。采购付款需要有效已确认对账单；'
        '退款、冲销从原记录操作。导入不发起银行转账。',
    ),
    'time': (
        '工时',
        '打开对应项目 → 任务与工时 → 工时记录 → 导入导出 → 导入',
        '为已有任务登记工时，不创建任务。任务须为待完成或已完成；人员须在项目范围内。'
        '普通成员只能登记本人任务的本人工时，管理岗位仍按任务管理范围校验。',
    ),
    'deliveries': (
        '发货',
        '打开对应项目 → 交付与售后 → 导入导出 → 导入',
        '登记实际发货并创建安装、验收任务。设计、装配、调试任务须全部完成，BOM及领料须满足本批交付。'
        '多单元项目请从页面发货并明确本批配套物料，不能用此模板省略配套清单。',
    ),
    'moves': (
        '项目领料',
        '库存 → 库存流水 → 导入导出 → 导入',
        '登记项目领料，数量不能超过库存及BOM剩余需求；退料、退货从原流水操作。'
        '需要确认占用其他项目到货时，请改用页面领料并核对提示。',
    ),
    'bom': (
        'BOM',
        'BOM → 选择项目 → 导入 BOM → 上传并检查问题；也可在项目详情的BOM页操作',
        '已有物料填写系统物料编码，后方物料资料可以留空；填写时必须与现有资料一致，不会覆盖主数据。'
        '新增物料编码留空，填写名称、规格、产品编码类别和单位；有图产品再填图号。'
        '完全相同物料自动复用；多个匹配须明确编码，新编号只在确认后生成。'
        '同一文件中“物料编码＋单元”不能重复，空单元也算同一单元；先核实重复粘贴、独立用量或不同单元，不能盲目累加。',
    ),
}

# Each template starts with a short task-specific route, before the detailed field reference.
# These instructions describe existing rules; they do not invent default business records.
QUICK_START = {
    'items': (
        '先在基础资料 → 物料搜索，确认没有同一物料；准备型号、图纸和计量单位。',
        '物料名称必填。选择产品编码类别后还要填规格；有图类别再填图号。编码留空可自动编号。',
        '例如新增外购滑块：名称填滑块，规格填厂家真实型号，类别选无图·标准件，单位按实物填写；'
        '不要把这份表传到BOM入口。',
    ),
    'partners': (
        '准备单位全称和联系方式，先确认是客户、供应商还是两者兼有。',
        '先填往来单位，再下拉选择类型。纯客户不用填采购账期和月结天数；供应商按实际采购约定填写。',
        '例如新增客户：填真实公司全称，类型选客户，其余无资料的可选项留空。'
        '选择自定义月结的供应商还要核实约定天数，不能用销售回款约定代替采购账期。',
    ),
    'sales': (
        '先建好客户；到基础资料 → 往来单位复制客户编码，确认销售负责人已启用。',
        '销售名称、客户编码、负责人姓名或账号必填；有约定交期时填具体日期。',
        '例如销售一台设备：填合同业务名称、真实客户编码和负责人姓名。'
        '导入后是销售草稿，仍需报价、签约；不会直接生成执行项目。',
    ),
    'projects': (
        '先确认是否已有销售单：已签约的销售请从销售签约生成项目，避免建两次。',
        '项目名称、客户编码、负责人姓名或账号必填；负责人须是项目经理或管理员。',
        '例如直接新建执行项目：填真实项目名称、客户编码、项目经理姓名。'
        '要导入设计/装配/调试任务，请使用项目任务模板；要导入物料需求，请使用BOM模板。',
    ),
    'purchases': (
        '准备项目编号、供应商编码和已建档物料编码。按BOM单元采购请使用页面“从BOM多选下单”。',
        '每行填分组号、项目编号、供应商编码、交期、物料编码、数量、含税单价；同张单的分组号相同。',
        '例如同一采购单有两种物料：分成两行，分组号都填G001，物料编码分别填写，其他单头资料保持一致。'
        '另一张单改填G002；这些分组号由填表人编写，不是系统采购编号。',
    ),
    'tasks': (
        '先确认项目处于执行中或交付中，执行人员已加入该项目。',
        '任务名称、阶段、执行人姓名或账号、项目编号必填；期限有要求时填截止日期。',
        '例如设计任务：名称填机械设计，阶段下拉选设计，执行人填真实姓名，项目编号从项目列表复制。'
        '期限填2026-10-20这种日期，不填20；采购、发货和T0/T1不是任务阶段。',
    ),
    'stocks': (
        '仅管理员从设置 → 管理员初始化使用。先建好物料，确认这个物料和库位还没有库存流水；已有库存用盘点调整。',
        '物料编码、库存数量、单位成本、原因必填；库位留空默认主仓。',
        '例如录入某库位10件物料：数量填10，单位成本填每件成本，不能填10件的总成本。'
        '同一物料在不同库位要分行填写，不要重复导入期初。',
    ),
    'entries': (
        '准备项目编号、真实费用依据和付款期限；这里只登记项目费用。',
        '项目编号、款项、原金额、最近待付期限必填；这里的原金额就是本笔费用金额。',
        '例如项目差旅费：款项填项目差旅费，原金额填本笔费用，期限填付款到期日期。'
        '合同应收、采购应付由业务单据生成，不要在这里再录一次。',
    ),
    'payments': (
        '先到收付款 → 应收应付与费用导出款项记录ID；采购付款还要核实已确认的业务对账单ID。',
        '款项ID、金额、日期、说明必填；采购付款还须填对账单ID，金额不能超过可收付余额。',
        '例如登记实际到账：款项ID从导出文件复制，金额填本次到账金额，日期填实际到账日。'
        '这里登记流水，不会发起银行转账；退款和冲销从原记录操作。',
    ),
    'time': (
        '先打开项目 → 任务与工时 → 项目任务，导出取得任务ID；任务须已存在。',
        '任务ID、日期、工时、说明必填；人员留空表示本人。普通成员仅能登记本人任务的本人工时。',
        '例如当天完成7.5小时设计：工时填7.5，说明填写实际工作内容。'
        '不要在任务ID里填任务名称；同人同日累计净工时不能超过24小时。',
    ),
    'deliveries': (
        '先核实设计/装配/调试已完成，BOM和领料满足本批交付；多单元项目从页面发货并指定配套清单。',
        '项目编号、设备数量、发货日必填；安装人、验收人留空默认项目负责人。',
        '例如实际发出1台设备：设备数量填1，发货日填实际发出日期。不能超过剩余未发货台数；导入后会创建安装和验收任务。',
    ),
    'moves': (
        '到库存 → 现有库存导出库存记录ID，核对物料和库位，再准备项目编号。',
        '库存ID、数量、说明、项目编号必填；已验收项目还须填该项目未完成的售后任务ID。',
        '例如装配领料：库存ID从对应物料和库位的导出行复制，数量填本次实际领用数量。'
        '不要用物料编码代替库存ID；库存不足、超BOM或需确认占用其他项目到货时，按系统提示处理。',
    ),
    'bom': (
        '先在BOM页面选对项目；再分清本行是引用已有物料，还是新增物料。',
        '已有物料：填系统物料编码和需求数量。新增物料：编码留空，填需求数量、名称、规格、产品编码类别、单位；'
        '有图物料还填图号。修改已有BOM行还要填变更说明。',
        '例如新增外购滑块：编码空着，数量填本项目需求量，规格填真实厂家型号，类别选无图·标准件，'
        '单位按实物填写。加工件选实际加工类别，规格写真实材质/尺寸，图号另填，不能用图号代替规格。',
    ),
}

SOURCES = {
    'customer': '基础资料 → 往来单位：查到客户后复制“编码”。',
    'supplier': '基础资料 → 往来单位：查到供应商后复制“编码”。',
    'project': '项目 → 项目列表：复制对应项目的“编号”，例如ATM2601。',
    'item': '基础资料 → 物料：核对名称和规格，再复制“编码”。',
    'entry': '收付款 → 应收应付与费用 → 导入导出 → 导出：复制对应行的“ID”。',
    'stock': '库存 → 现有库存 → 导入导出 → 导出：核对物料和库位，复制对应行的“ID”。',
    'task': '打开对应项目 → 任务与工时 → 项目任务 → 导入导出 → 导出：复制任务“ID”。',
    'document': '核对已上传的同项目收付凭证记录ID；不清楚时请财务核对，不能填文件名或路径。',
    'reconciliation': '收付款 → 业务对账：核对已确认且仍有效的同款项对账单，导出取得“ID”。',
    'specification': '标准件查厂家型号；加工件查设计图纸的材质和尺寸，资料缺失先找设计人员确认。',
    'drawing_number': '从实际设计图纸复制图号，与规格分开填写。',
    'drawing_revision': '查图纸上的实际版本；没有版本可空，不是计量单位。',
    'brand': '按实物铭牌、厂家资料或原物料档案填写。',
    'group': '自己为本次采购文件编写分组号：同张单相同，不同单不同。',
    'phone': '向业务联系人核实号码；保留区号和前导零。',
}


def field_source(resource, key, options):
    if key in ('manager', 'members', 'assignee', 'user', 'installer', 'acceptor'):
        return '填写系统中已启用人员的姓名；重名时请管理员在设置 → 用户管理核对账号，填“姓名（用户名）”。'
    if key == 'code':
        return (
            '已有物料到基础资料 → 物料复制编码；新增物料此格留空。'
            if resource == 'bom'
            else '新增编码可留空自动编号；自定义编码先核实未使用，已有物料请直接复用。'
        )
    if options:
        return (
            '单击第2行起的空白格，再点右侧▼选择；单位允许手填实际单位。'
            if key == 'unit'
            else ('单击第2行起的空白格，再点右侧▼选择；不用记代码，也不要改写选项名称。')
        )
    if key.endswith('_date') or key == 'date':
        return '从实际业务记录或约定中确认具体日期；填写年-月-日，例如2026-10-20。'
    return SOURCES.get(key, '按本列规则填写真实业务资料；没有资料的可选项留空，不复制示例凑数。')


PERSON_HINT = '可填姓名或登录账号；重名或姓名与他人账号相同时，填写姓名（用户名）明确人员。'
DATE_HINT = '填写具体日期YYYY-MM-DD，不是天数，也不能填Excel日期序号。'
PRODUCT_HINT = '；'.join(f'{value}＝{label}' for value, label in Item.ProductCategory.choices) + (
    '。用于自动生成物料编码；Excel从下拉框选择完整中文名称，CSV可填上述编码或完整名称，不能填1或2。'
    '“物料类别”只区分标准件/非标件，不要把机加、钣金填到那一列。'
)
COMMON_UNITS = [
    '件',
    'PCS',
    '套',
    '台',
    '个',
    '只',
    '支',
    '根',
    '米',
    '平方米',
    '千克',
    '克',
    '升',
    '毫升',
    '瓶',
    '桶',
    '盒',
    '箱',
    '包',
    '卷',
    '张',
]

# required, hint, example. Every schema key must have an explicit description.
FIELDS = {
    'name': ('必填', '填写名称，不填写系统编号。', '输送线支架'),
    'code': (
        '可空',
        '新增物料可自定义唯一编码，最多30字符；留空自动生成。已用编码（含停用或已删除）不能重复使用。',
        'MAT-001',
    ),
    'specification': (
        '条件必填',
        '填写产品编码类别时必填真实型号/规格，不能用图号替代。标准件填厂家型号；'
        '加工件可填实际材质、尺寸等识别信息，例如6061铝板 200×100×10mm。最多2000字符。',
        'HGW20CC',
    ),
    'drawing_number': (
        '条件必填',
        '产品编码类别为11/12/13/19时必填图号。与型号/规格分别填写，最多100字符。',
        'A2608-01-P001',
    ),
    'drawing_revision': ('可空', '填写实际图档版本，最多30字符；没有版本可留空。不要填写PCS、品牌或数量。', 'A'),
    'product_category': ('可空', PRODUCT_HINT, '无图·标准件'),
    'brand': ('可空', '填写品牌，最多80字符；不要填写图档版本。', '上银'),
    'part_type': (
        '可空',
        '只区分标准件/非标件；通常外购标准产品选标准件，按项目图纸定制的零件选非标件，不确定可先留空。'
        '机加、钣金填在产品编码类别中。旧文件兼容standard/custom。',
        '标准件',
    ),
    'unit': ('可空，默认件', '填写计量单位，不要填入图档版本列；同一物料沿用其原单位。', 'PCS'),
    'duplicate_reason': (
        '条件必填',
        '相同名称、规格等资料的物料确需独立建码时，填写真实原因；不覆盖原物料。',
        '客户要求单独管理',
    ),
    'contact': ('可空', '填写业务联系人姓名。', '张工'),
    'phone': ('可空', '按文本填写电话，保留区号和前导零。', '0571-12345678'),
    'address': ('可空', '填写联系或收货地址。', '杭州市某路1号'),
    'customer': ('必填', '填写基础资料中已启用客户或双用途单位的编码，不能填写客户名称或记录ID。', 'C001'),
    'supplier': ('必填', '填写基础资料中已启用供应商或双用途单位的编码，不能填写名称或记录ID。', 'S001'),
    'manager': ('必填', PERSON_HINT + '人员须已启用。', '刘经理'),
    'requirements': ('可空', '填写需求内容，最多20000字符。', '输送后自动定位并焊接'),
    'due_date': ('可空', DATE_HINT, '2026-10-20'),
    'equipment_quantity': ('可空，默认1', '填写大于0的整数，单位为台。', '1'),
    'warranty_months': ('可空，默认12', '填写0～120的整数，单位为月。', '12'),
    'members': ('可空', PERSON_HINT + '多个成员用分号（;或；）分隔，人员须已启用。', '张工;李工'),
    'group': ('必填', '同一采购草稿使用同一分组号；不同采购单用不同分组号，不是系统采购编号。', 'G001'),
    'project': (
        '必填',
        '填写可操作的执行项目编号，不填名称或记录ID。文件编号决定归属，不自动使用当前页面筛选的项目。',
        'ATM2601',
    ),
    'item': ('必填', '填写已启用物料的系统编码，不能填名称、图号或记录ID。', 'MAT-001'),
    'quantity': ('必填', '填写大于0的数量，最多3位小数，不加单位文字。', '2.125'),
    'unit_price': ('必填', '填写大于或等于0的人民币含税单价，最多2位小数；不填未税价或整行总额。', '120.50'),
    'line_due_date': ('可空', DATE_HINT + '留空沿用该采购单交期。', '2026-10-20'),
    'payment_term': (
        '可空',
        '填写指定日期、现付、月结30天、月结60天、月结90天、月结120天、自定义月结；'
        '兼容manual/cash/month30/month60/month90/month120/custom。月结按收货当月月底加天数。',
        '月结30天',
    ),
    'payment_days': ('自定义月结时填写', '填写0～365的整数，仅用于自定义月结，其他账期请留空。', '45'),
    'payment_due_date': ('可空', DATE_HINT + '仅指定日期账期可填；自动账期必须留空，按实际收货计算。', '2026-10-20'),
    'note': ('可空', '填写本单据的补充说明。', '按约定交付'),
    'title': ('必填', '填写任务的具体名称，最多150字符。', '机械设计'),
    'description': ('可空', '填写任务工作内容，最多20000字符。', '完成装配图并校核'),
    'assignee': ('必填', PERSON_HINT + '须已启用且属于该项目的负责人或可访问成员。', '张工'),
    'location': ('可空，默认主仓', '填写库位名称；期初允许建立新库位，同物料不同库位分别填行。', '主仓'),
    'unit_cost': ('必填', '填写大于或等于0的单位成本，人民币含税，最多2位小数，不填库存总金额。', '20.00'),
    'reason': ('必填', '填写实际业务原因或工作内容。', '期初盘点录入'),
    'amount': ('必填', '填写大于0的人民币含税金额，最多2位小数，不带货币符号或千分位。', '100.00'),
    'entry': ('必填', '从应收应付列表导出获取款项的记录ID，不填项目编号或款项名称。', '123'),
    'date': ('必填', DATE_HINT + '填写实际发生日期，不能晚于今天。', '2026-10-19'),
    'method': ('可空，默认银行转账', '填写银行转账、现金、其他（兼容bank/cash/other）。', '银行转账'),
    'account': ('可空', '填写收付账户标识，按文本保留前导零。', '基本户'),
    'reference': ('可空', '填写银行流水号，按文本保留完整号码，不用Excel科学计数法。', '000012345678'),
    'document': ('可空', '填写已上传的同项目收付凭证记录ID，不填文件名或文件路径。', '456'),
    'reconciliation': ('采购付款必填', '填写已确认且仍有效的同款项对账单ID；客户收款、项目费用可留空。', '789'),
    'task': ('必填', '从项目任务列表导出获取任务记录ID，不填任务名称或项目编号。', '101'),
    'user': ('可空，默认本人', PERSON_HINT + '须已启用且在项目范围内。', '张工'),
    'hours': ('必填', '填写大于0的小时数，最多2位小数；同一人员同一天累计净工时不能超过24小时。', '7.50'),
    'installer': ('可空，默认项目负责人', PERSON_HINT + '安装人员须已启用且属于该项目。', '张工'),
    'acceptor': ('可空，默认项目负责人', PERSON_HINT + '验收人员须已启用且属于该项目。', '刘经理'),
    'stock': ('必填', '从共享库存列表导出获取库存记录ID，代表具体物料和库位，不是物料编码。', '202'),
    'assembly_unit': (
        '可空',
        '填写实际装配单元，最多100字符。同料可分属不同单元；相同编码和单元不可重复，空单元也会判重。',
        '上料单元',
    ),
    'change_note': ('修改已有BOM时必填', '填写变更原因，最多500字符；新行可空，不是物料备注。', '设计变更，增加支撑件'),
    'required_date': ('可空', DATE_HINT + '表示物料需求日期。', '2026-10-20'),
    'application_date': ('可空', DATE_HINT + '表示申请日期。', '2026-10-19'),
    'applicant': ('可空', '填写纸面申请人姓名，最多80字符；本列不是任务的执行人账号。', '张工'),
}

OVERRIDES = {
    ('items', 'name'): ('必填', '填写物料名称，不填写编码或规格。', '滑块'),
    ('partners', 'name'): ('必填', '填写往来单位名称，不填写联系人姓名。', '示例机械有限公司'),
    ('sales', 'name'): ('必填', '填写本笔销售业务名称，最多150字符。', '输送线设备销售'),
    ('projects', 'name'): ('必填', '填写执行项目名称，最多150字符；项目编号自动生成。', '输送线改造项目'),
    ('partners', 'kind'): ('必填', '填写客户、供应商、客户及供应商（兼容customer/supplier/both）。', '供应商'),
    ('partners', 'payment_term'): (
        '可空，默认指定日期',
        FIELDS['payment_term'][1] + '纯客户留空或填指定日期。',
        '月结30天',
    ),
    ('purchases', 'payment_term'): ('可空，沿用供应商账期', FIELDS['payment_term'][1], '月结30天'),
    ('purchases', 'due_date'): ('必填', DATE_HINT + '表示整张采购单交期，同组保持一致。', '2026-10-20'),
    ('projects', 'manager'): ('必填', FIELDS['manager'][1] + '须具备项目经理或管理员岗位。', '刘经理'),
    ('sales', 'manager'): ('必填', FIELDS['manager'][1] + '须具备销售经理、项目经理或管理员岗位。', '刘经理'),
    ('tasks', 'kind'): (
        '必填',
        '填写设计、装配、调试（兼容design/assembly/test）。不能填T0/T1/T2/T3、采购或发货。',
        '设计',
    ),
    ('tasks', 'due_date'): ('可空', DATE_HINT + '“期限”表示截止日期，不是工期天数。', '2026-10-20'),
    ('entries', 'due_date'): ('必填', DATE_HINT + '表示费用付款到期日。', '2026-10-20'),
    ('entries', 'title'): ('必填', '填写项目费用名称；不能用来新增合同应收或采购应付。', '项目差旅费'),
    ('time', 'reason'): ('必填', '填写当日实际工作内容。', '完成机械装配图'),
    ('payments', 'reason'): ('必填', '填写本次实际收付款说明。', '支付项目差旅费'),
    ('moves', 'reason'): ('必填', '填写实际领料用途。', '装配领料'),
    ('moves', 'task'): ('条件必填', '可关联同项目未完成任务。已验收项目领料必须填未完成的售后任务ID。', '101'),
    ('deliveries', 'quantity'): ('必填', '填写大于0的设备台数整数，不能超过项目剩余未发货台数。', '1'),
    ('bom', 'code'): (
        '已有物料必填；新增留空',
        '已有编码必须在系统中存在且启用；后方物料资料可空，填了须与主数据一致。'
        '新增物料请留空，不能直接填外部系统的未建档编码。',
        'MAT-001',
    ),
    ('bom', 'name'): ('新增物料必填', '编码留空时填写物料名称；已有物料可空，填写时须与系统一致。', '滑块'),
    ('bom', 'specification'): (
        '新增物料必填',
        '编码留空时必须填写真实型号/规格，即使已填写图号也不能省略。标准件填厂家型号，例如HGW20CC；'
        '加工件填实际材质、尺寸等识别信息，例如6061铝板 200×100×10mm。资料不全请向设计人员确认。',
        'HGW20CC',
    ),
    ('bom', 'drawing_number'): (
        '新增有图物料必填',
        '新建11/12/13/19类物料时填写图号，最多100字符；已有物料可空，填写时须与系统一致。',
        'A2608-01-P001',
    ),
    ('bom', 'product_category'): ('新增物料必填', PRODUCT_HINT, '无图·标准件'),
    ('bom', 'unit'): ('新增物料必填', '填写计量单位；已有物料可空，填写时须与系统一致。', 'PCS'),
}

BOM_COLUMNS = [
    ('物料编码', 'code'),
    ('单元', 'assembly_unit'),
    ('需求数量', 'quantity'),
    ('变更说明', 'change_note'),
    ('需求日期', 'required_date'),
    ('申请日期', 'application_date'),
    ('申请人', 'applicant'),
    ('物料名称', 'name'),
    ('规格', 'specification'),
    ('图号', 'drawing_number'),
    ('图档版本', 'drawing_revision'),
    ('产品编码类别', 'product_category'),
    ('品牌', 'brand'),
    ('物料类别', 'part_type'),
    ('单位', 'unit'),
]


def choices(resource, key):
    # Reuse the importer's accepted Chinese values; avoid a second enum vocabulary.
    from .transfers import ENUMS

    if key == 'product_category':
        return list(Item.ProductCategory.labels)
    if key in ('part_type', 'payment_term', 'method'):
        return list(ENUMS[key])
    if (resource, key) == ('partners', 'kind'):
        return list(Partner.Kind.labels)
    if (resource, key) == ('tasks', 'kind'):
        return [Task.Kind(value).label for value in ('design', 'assembly', 'test')]
    return COMMON_UNITS if key == 'unit' else []


def guide(resource, columns):
    title, entry, note = RESOURCES[resource]
    fields = []
    for index, (label, key) in enumerate(columns, 1):
        required, hint, example = OVERRIDES[(resource, key)] if (resource, key) in OVERRIDES else FIELDS[key]
        options = choices(resource, key)
        if options:
            hint = (
                '可从下拉选择常用单位，也可直接输入实际单位。' if key == 'unit' else 'Excel请从下拉选项中选择。'
            ) + hint
        fields.append(
            {
                'key': key,
                'label': label,
                'required': required,
                'hint': hint,
                'example': example,
                'options': options,
                'allow_custom': key == 'unit',
                'source': field_source(resource, key, options),
                'cell': f'{get_column_letter(index)}2',
            }
        )
    return {
        'resource': resource,
        'title': title,
        'entry': entry,
        'note': note,
        'quick_start': dict(zip(('prepare', 'minimum', 'example'), QUICK_START[resource], strict=True)),
        'instructions': COMMON,
        'columns': fields,
    }


def required_condition(resource, column, letters):
    """Only highlight conditions knowable from this row; server checks project facts."""
    key = column['key']
    if column['required'] == '必填':
        return 'TRUE'
    if resource == 'bom' and key in ('name', 'specification', 'product_category', 'unit'):
        return f'LEN(TRIM(${letters["code"]}2&""))=0'
    if resource == 'items' and key == 'specification':
        return f'LEN(TRIM(${letters["product_category"]}2&""))>0'
    if resource in ('bom', 'items') and key == 'drawing_number':
        category = f'${letters["product_category"]}2'
        drawing = f'OR(LEFT({category}&"",1)="1",LEFT({category}&"",2)="有图")'
        return f'AND(LEN(TRIM(${letters["code"]}2&""))=0,{drawing})' if resource == 'bom' else drawing
    return ''


def input_controls(data, layout):
    letters = {column['key']: get_column_letter(index) for index, column in enumerate(layout['columns'], 1)}
    active_row = f'COUNTA($A2:${get_column_letter(len(letters))}2)>0'
    for column in layout['columns']:
        letter = letters[column['key']]
        cells = f'{letter}2:{letter}1001'
        options = column['options']
        validation = DataValidation(
            type='list' if options else 'custom',
            formula1='"' + ','.join(options) + '"' if options else 'TRUE',
            allow_blank=True,
            showDropDown=False,
            showInputMessage=True,
            showErrorMessage=bool(options) and not column['allow_custom'],
            errorStyle='stop',
            errorTitle='请使用下拉选项',
            error=f'“{column["label"]}”请从右侧▼选择，不要自行改写名称。可选内容：' + '、'.join(options),
            promptTitle=column['label'][:32],
            prompt=(
                f'{column["required"]}。{column["source"]} 示例：{column["example"]}。'
                + ('完整规则见“填写说明”页。' if options else column['hint'])
            )[:255],
        )
        data.add_data_validation(validation)
        validation.add(cells)
        condition = required_condition(layout['resource'], column, letters)
        if condition:
            data.conditional_formatting.add(
                cells,
                FormulaRule(
                    formula=[f'AND({active_row},{condition},LEN(TRIM({letter}2&""))=0)'],
                    fill=PatternFill('solid', fgColor='FDE2E2'),
                    font=Font(color='9B1C1C'),
                ),
            )
        if column['key'].endswith('_date') or column['key'] == 'date':
            data.column_dimensions[letter].number_format = 'yyyy-mm-dd'
        elif column['key'] in (
            'code',
            'item',
            'customer',
            'supplier',
            'project',
            'phone',
            'reference',
            'account',
            'manager',
            'members',
            'assignee',
            'user',
            'installer',
            'acceptor',
            'drawing_number',
            'drawing_revision',
            'group',
            'entry',
            'stock',
            'task',
            'document',
            'reconciliation',
        ):
            data.column_dimensions[letter].number_format = '@'


def html_guide(layout):
    fields = ''.join(
        '<tr>'
        + f'<td>{escape(column["label"])}<br><small>数据页 {column["cell"]} 起'
        + (' · ▼下拉选择' if column['options'] else ' · 手动填写')
        + '</small></td>'
        + ''.join(f'<td>{escape(column[key])}</td>' for key in ('required', 'source', 'hint', 'example'))
        + '</tr>'
        for column in layout['columns']
    )
    rules = ''.join(f'<li>{escape(rule)}</li>' for rule in layout['instructions'])
    quick = layout['quick_start']
    dropdowns = '、'.join(
        f'{column["label"]}（{column["cell"]}起）' for column in layout['columns'] if column['options']
    )
    dropdown_hint = (
        dropdowns + '。单击数据格，再点右侧▼；表头不显示下拉箭头。'
        if dropdowns
        else '本模板没有固定选项，请按真实资料填写；选中数据格可查看填写提示。'
    )
    return (
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>{escape(layout["title"])}导入模板填写说明</title>'
        '<style>body{font:16px/1.7 system-ui,sans-serif;color:#233044;max-width:1100px;margin:32px auto;padding:0 20px}'
        'h1{font-size:26px}li{margin:8px 0}table{border-collapse:collapse;width:100%}th,td{border:1px solid #ccd5df;'
        'padding:12px;text-align:left;vertical-align:top}th{background:#edf3f8}td:first-child{min-width:100px}'
        '.quick{background:#edf5ff;border:1px solid #b9d2f2;border-radius:12px;padding:16px 20px}'
        '.table{overflow:auto}small{color:#546377}@media print{body{font-size:11px;margin:0}thead{display:table-header-group}'
        'tr{break-inside:avoid}}</style>'
        f'<h1>{escape(layout["title"])}导入模板填写说明</h1>'
        f'<p><strong>导入入口：</strong>{escape(layout["entry"])}</p>'
        '<section class="quick"><h2>第一次填写，照这三步做</h2>'
        f'<p><strong>① 先准备：</strong>{escape(quick["prepare"])}</p>'
        f'<p><strong>② 开始填写：</strong>{escape(quick["minimum"])}</p>'
        '<p><strong>③ 上传预览：</strong>保存Excel，回到上方导入入口选择文件，核对项目、行数、数量和金额。'
        '有错误先按行号修改，重新上传，校验通过后再确认导入。</p>'
        f'<p><strong>下拉在哪里：</strong>{escape(dropdown_hint)}</p></section>'
        f'<h2>照着理解，不要照抄</h2><p>{escape(quick["example"])}</p>'
        f'<p>{escape(layout["note"])}</p>'
        f'<h2>填写前请阅读</h2><ol>{rules}</ol>'
        '<h2>逐列填写说明</h2><p>示例仅说明格式，请替换为真实业务数据。</p><div class="table"><table>'
        '<thead><tr><th>模板字段 / 填写位置</th><th>是否必填</th><th>去哪找 / 怎么选</th><th>填写规则</th>'
        '<th>格式示例</th></tr></thead>'
        f'<tbody>{fields}</tbody></table></div></html>'
    )


def workbook_guide(book, layout):
    sheet = book.create_sheet(GUIDE_SHEET)
    quick = layout['quick_start']
    sheet.append([GUIDE_MARKER])
    sheet.append([f'{layout["title"]}导入 · 从这里开始'])
    sheet.merge_cells('A2:C2')
    sheet.append(['用在哪个入口', layout['entry']])
    sheet.append(['① 先准备', quick['prepare']])
    sheet.append(['② 最少填什么', quick['minimum']])
    sheet.append(['开始填写 →', '点击这里，到“数据”页第2行开始填写；一行一条，不改表头。'])
    sheet['B6'].hyperlink = Hyperlink(ref='B6', location="'数据'!A2")
    sheet.append(
        [
            '下拉怎么打开',
            '单击空白数据格，再点右侧▼；不要点表头，也不需要双击。下面标有“▼点此选择”的字段可以直接跳转。'
            if any(column['options'] for column in layout['columns'])
            else '本模板没有固定下拉选项，按真实资料填写即可。下方“点此填写”可跳到对应输入格，并查看填写提示。',
        ]
    )
    sheet.append(
        [
            '③ 上传并预览',
            '保存Excel，回到上面的导入入口选择文件。核对项目、行数、数量和金额；'
            '有错误按文件行号修改并重新上传，通过后才确认导入。',
        ]
    )
    sheet.append(['先看一个例子', quick['example']])
    sheet.append(
        [
            '遇到错误怎么办',
            '不要猜规格、姓名、编号或ID。按报错行号核对资料；表头不符时先确认导入入口，'
            '重新下载对应模板，只把业务数据按列粘贴为值。没有错误才确认，成功文件不要重复导入。',
        ]
    )
    sheet.append(['这份表会做什么', layout['note']])
    sheet.append(
        [
            '颜色怎么认',
            '红色表头＝必填；黄色＝满足条件才必填；蓝色＝可选。填写后缺项标红。'
            '说明页不参与导入，示例不要复制为真实业务。',
        ]
    )
    section_rows = [13]
    sheet.append(['字段 / 点此去填', '是否必填 · 去哪里找 · 填写规则', '示例（必须换成真实资料）'])
    for column in layout['columns']:
        action = '▼点此选择' if column['options'] else '点此填写'
        sheet.append(
            [
                f'{column["label"]}\n{action} · {column["cell"]}',
                f'{column["required"]}\n{column["source"]}\n{column["hint"]}',
                column['example'],
            ]
        )
        cell = sheet.cell(sheet.max_row, 1)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'数据'!{column['cell']}")
    section_rows.append(sheet.max_row + 1)
    sheet.append(['保存前再核对', '操作提醒'])
    for rule in layout['instructions']:
        sheet.append(['填写提醒', rule])
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='等线', size=11, color='233044')
            if cell.hyperlink:
                cell.font = Font(name='等线', size=11, bold=True, color='175CD3', underline='single')
        lines = max(
            sum(max(1, (len(line) + width - 1) // width) for line in str(cell.value or '').split('\n'))
            for cell, width in zip(row, (12, 42, 18), strict=True)
        )
        sheet.row_dimensions[row[0].row].height = max(40, 18 * lines + 12)
    for letter, width in [('A', 26), ('B', 86), ('C', 34)]:
        sheet.column_dimensions[letter].width = width
    for row in (2, 6, *section_rows):
        for cell in sheet[row]:
            cell.fill = PatternFill('solid', fgColor='E8F0F8')
    sheet['A2'].font = Font(name='等线', size=20, bold=True, color='17365D')
    sheet.row_dimensions[1].hidden = True
    sheet.row_dimensions[2].height = 42
    sheet.freeze_panes = 'B3'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    book.active = sheet


def template_document(name, layout, file_format):
    headers = [column['label'] for column in layout['columns']]
    if file_format == 'csv':
        response = document(name, headers, [], 'csv')
    elif file_format == 'html':
        response = HttpResponse(html_guide(layout), content_type='text/html; charset=utf-8')
    elif file_format == 'xlsx':
        book = Workbook()
        data = book.active
        data.title = '数据'
        data.append(headers)
        data.freeze_panes = 'A2'
        data.sheet_properties.pageSetUpPr.fitToPage = True
        data.row_dimensions[1].height = 32
        for index, column in enumerate(layout['columns'], 1):
            header = data.cell(1, index)
            header.font = Font(bold=True, color='233044')
            required = column['required']
            header.fill = PatternFill(
                'solid', fgColor='FDE2E2' if required == '必填' else 'FFF2CC' if '必填' in required else 'E8F0F8'
            )
            header.alignment = Alignment(vertical='center', wrap_text=True)
            header.comment = Comment(
                f'{column["required"]}\n{column["hint"]}\n格式示例：{column["example"]}\n完整说明见“填写说明”页。',
                '项目ERP',
            )
            header.comment.width = 380
            header.comment.height = 220
            data.column_dimensions[get_column_letter(index)].width = max(18, len(column['label']) * 2 + 2)
        input_controls(data, layout)
        workbook_guide(book, layout)
        output = io.BytesIO()
        book.save(output)
        book.close()
        response = HttpResponse(
            output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    else:
        raise ValidationError({'file_format': '请选择csv、xlsx或html填写说明。'})
    title = layout['title'] + ('导入填写说明' if file_format == 'html' else '导入模板')
    response['Content-Disposition'] = content_disposition_header(True, f'{title}.{file_format}')
    response['Cache-Control'] = 'no-store'
    return response
