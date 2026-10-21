"""Presentation-only import feedback; never persists uploaded files or business data."""

import base64
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def template_hint(headers):
    # Only compare public template definitions, never inspect references in file rows.
    from . import bom_import, transfers
    from .import_guides import BOM_COLUMNS, RESOURCES

    candidates = {
        resource
        for resource in transfers.SCHEMAS
        if headers in [[label for label, _ in columns] for columns in transfers.accepted_schemas(resource)]
    }
    if headers in [
        [label for label, _ in BOM_COLUMNS],
        bom_import.HEADERS,
        bom_import.PREVIOUS_HEADERS,
        bom_import.LEGACY_HEADERS,
        bom_import.LEGACY_HEADERS[:3],
    ]:
        candidates.add('bom')
    if len(candidates) != 1:
        return None
    resource = candidates.pop()
    title, entry, _ = RESOURCES[resource]
    return {'code': resource, 'title': '物料清单（BOM）' if resource == 'bom' else title, 'entry': entry}


def locate(errors, columns):
    """Infer which template columns an error message refers to."""
    for error in errors:
        text = error['message']
        if not error.get('fields'):
            error['fields'] = [
                key
                for label, key in columns
                if label in text
                or (label.endswith('ID') and label[:-2] in text)
                or (key == 'quantity' and ('数量' in text or '用量' in text))
                or (key == 'amount' and ('金额' in text or '余额' in text))
                or (key == 'hours' and ('工时' in text or '小时' in text))
            ]
    return errors


def feedback(result, headers, rows, columns, title):
    """Add locations and a re-uploadable workbook without changing the import schema."""
    errors = locate(result.get('errors', []), columns)
    if not errors:
        return result
    book = Workbook()
    sheet = book.active
    sheet.title = '数据'
    sheet.freeze_panes = 'A2'
    for index, label in enumerate(headers, 1):
        cell = sheet.cell(1, index, label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476A')
        cell.comment = Comment(
            '红色单元格的批注包含校验问题。修改后保存并重新上传，系统会重新校验整批数据。', '导入校验'
        )
        sheet.column_dimensions[get_column_letter(index)].width = 24
    for row_number, row in enumerate(rows, 2):
        for index, value in enumerate(row, 1):
            if isinstance(value, datetime):
                value = value.date().isoformat()
            elif isinstance(value, date):
                value = value.isoformat()
            cell = sheet.cell(row_number, index, '' if value is None else str(value))
            # Preserve literal input, including leading zeros and formula-like text.
            cell.data_type = 's'
    labels = dict((key, label) for label, key in columns)
    for error in errors:
        row_number = error.get('row')
        if not isinstance(row_number, int) or not 2 <= row_number <= len(rows) + 1:
            row_number = 1
        targets = [
            index for index, label in enumerate(headers, 1) if label in {labels.get(key) for key in error['fields']}
        ]
        # Legacy headers may differ; a whole-row issue is still explicitly marked.
        for index in targets or list(range(1, len(headers) + 1)):
            cell = sheet.cell(row_number, index)
            previous = cell.comment.text + '\n' if cell.comment else ''
            cell.comment = Comment(previous + error['message'], '导入校验')
            cell.fill = PatternFill('solid', fgColor='FDE2E2')
    output = io.BytesIO()
    book.save(output)
    book.close()
    result['error_file'] = {
        'name': f'{title}待修改.xlsx',
        'content': base64.b64encode(output.getvalue()).decode('ascii'),
    }
    return result
