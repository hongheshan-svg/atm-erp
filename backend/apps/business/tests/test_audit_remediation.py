"""审计整改 1–17 中涉及后端行为的回归：报表口径、BOM 按行采购、账期到期、主数据、导出、日期与库位等。"""

import csv
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.business.models import (
    BOMLine,
    Entry,
    Item,
    PurchaseLine,
    PurchaseOrder,
    PurchaseWarranty,
    SalesOrder,
    Stock,
    StockMove,
    Task,
)
from apps.business.services.tabular import cell, export_rows
from apps.core.models import AuditLog

from .test_commercial_chain import TODAY, BusinessFixtures
from .test_execution import ExecutionFixtures


class AuditRemediationTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()

    def get(self, role, url, params=None, status=200):
        response = self.clients[role].get('/api/' + url, params or {})
        self.assertEqual(response.status_code, status, getattr(response, 'data', response))
        return response

    def two_unit_bom(self, project):
        self.post(
            'manager',
            f'projects/{project.pk}/revise-bom/',
            {
                'expected_revision': self.bom_version(project.pk),
                'lines': [
                    {'item': self.item.pk, 'quantity': '10', 'assembly_unit': 'A单元'},
                    {'item': self.item.pk, 'quantity': '10', 'assembly_unit': 'B单元'},
                ],
            },
        )
        return (
            BOMLine.objects.get(project=project, assembly_unit='A单元'),
            BOMLine.objects.get(project=project, assembly_unit='B单元'),
        )

    def linked_purchase(self, project, bom, qty, *, status=201, from_demand=False):
        return self.post(
            'purchaser',
            'purchases/',
            {
                'project': project.pk,
                'supplier': self.supplier.pk,
                'due_date': TODAY,
                'from_demand': from_demand,
                'lines': [{'item': self.item.pk, 'bom_line': bom.pk, 'quantity': qty, 'unit_price': '10'}],
            },
            status=status,
        )

    def receive(self, purchase, qty, status=200, **extra):
        line = purchase.lines.get()
        return self.post(
            'warehouse',
            f'purchases/{purchase.pk}/receive/',
            {'lines': [{'line': line.pk, 'quantity': qty}], 'reason': '到货', **extra},
            status=status,
        )

    def demand(self, project):
        rows = self.get('manager', f'business/projects/{project.pk}/demand/').data['lines']
        return {row['bom_line']: row for row in rows}

    # 1
    def test_cancelled_project_contract_is_excluded_from_report(self):
        self.active_project('10000.00')
        cancelled = self.active_project('5000.00')
        self.post('manager', f'projects/{cancelled.pk}/cancel/', {'reason': '客户取消'})
        report = self.get('admin', 'business/reports/').data
        self.assertEqual(report['summary']['contract_amount'], '10000.00')
        row = next(row for row in report['results'] if row['id'] == cancelled.pk)
        self.assertEqual(row['contract_amount'], '0.00')
        # 销售单保留签约事实，列表通过执行项目状态看出已取消。
        sale = self.get('manager', f'business/sales/{SalesOrder.objects.get(project=cancelled).pk}/').data
        self.assertEqual((sale['status'], sale['project_status']), ('signed', 'cancelled'))
        self.post('manager', f'projects/{cancelled.pk}/reopen/', {'reason': '客户恢复'})
        self.assertEqual(self.get('admin', 'business/reports/').data['summary']['contract_amount'], '15000.00')

    # 2
    def test_purchase_cannot_borrow_another_units_bom_quantity(self):
        project = self.active_project()
        unit_a, unit_b = self.two_unit_bom(project)
        conflict = self.linked_purchase(project, unit_a, '20', status=409)
        self.assertIn('A单元', str(conflict['detail']))
        draft = PurchaseOrder.objects.get(pk=self.linked_purchase(project, unit_a, '10')['id'])
        demand = self.demand(project)
        self.assertEqual(Decimal(demand[unit_b.pk]['shortage']), Decimal('10'))
        # 缺料列表显示的缺口可以直接按 B 行下单，不再出现「看得到缺料却下不了单」。
        self.linked_purchase(project, unit_b, demand[unit_b.pk]['shortage'], from_demand=True)
        detail = self.get('purchaser', f'business/purchases/{draft.pk}/').data
        line = detail['lines'][0]
        self.post(
            'purchaser',
            f'purchases/{draft.pk}/edit/',
            {
                'expected_updated_at': detail['updated_at'],
                'reason': '改数量',
                'due_date': TODAY,
                'lines': [{'id': line['id'], 'quantity': '15', 'unit_price': '10', 'due_date': TODAY}],
            },
            status=409,
        )

    def test_legacy_overlinked_incoming_is_shared_with_other_units(self):
        project = self.active_project()
        unit_a, unit_b = self.two_unit_bom(project)
        purchase = self.purchase(project, qty='1', approve=False)
        # 按行校验上线前的历史数据：同一物料 20 件全挂在 A 行上。
        PurchaseLine.objects.filter(purchase=purchase).update(bom_line=unit_a, quantity=Decimal('20'))
        demand = self.demand(project)
        self.assertEqual(Decimal(demand[unit_a.pk]['incoming']), Decimal('10'))
        self.assertEqual(Decimal(demand[unit_b.pk]['incoming']), Decimal('10'))
        self.assertEqual(Decimal(demand[unit_b.pk]['shortage']), Decimal('0'))

    # 3
    def test_manual_term_payable_is_due_only_for_received_goods(self):
        project = self.active_project()
        purchase = self.purchase(project)  # 指定日期账期，到期日为已过去的 TODAY，尚未收货
        entry = Entry.objects.get(purchase=purchase)
        row = next(row for row in self.get('admin', 'business/reports/').data['results'] if row['id'] == project.pk)
        self.assertEqual((row['payable'], row['overdue_payable']), ('500.00', '0.00'))
        listed = self.get('finance', f'business/entries/{entry.pk}/').data
        self.assertEqual((listed['status'], listed['due_amount'], listed['payment_schedule']), ('open', '0.00', []))
        settlements = self.get('finance', 'business/workbench/', {'bucket': 'settlements'}).data['settlements']
        self.assertNotIn(entry.pk, [item['id'] for item in settlements['results']])
        aging = self.get('admin', 'business/reports/', {'view': 'aging'}).data['results']
        self.assertNotIn(entry.pk, [item['id'] for item in aging])

        self.receive(purchase, '2')
        row = next(row for row in self.get('admin', 'business/reports/').data['results'] if row['id'] == project.pk)
        self.assertEqual((row['payable'], row['overdue_payable']), ('500.00', '200.00'))
        listed = self.get('finance', f'business/entries/{entry.pk}/').data
        self.assertEqual((listed['status'], Decimal(listed['due_amount'])), ('overdue', Decimal('200')))
        settlements = self.get('finance', 'business/workbench/', {'bucket': 'settlements'}).data['settlements']
        self.assertIn(entry.pk, [item['id'] for item in settlements['results']])

    # 4
    def patch_item(self, item, data, status):
        response = self.clients['purchaser'].patch(
            f'/api/business/items/{item.pk}/', data, format='json', HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4())
        )
        self.assertEqual(response.status_code, status, response.data)
        return response.data

    def test_referenced_item_cannot_change_unit_or_specification(self):
        self.patch_item(self.item, {'unit': '台'}, 200)
        self.post(
            'admin',
            'stocks/opening/',
            {'item': self.item.pk, 'quantity': '5', 'unit_cost': '100', 'reason': '期初'},
        )
        error = self.patch_item(self.item, {'unit': '米'}, 400)
        self.assertIn('单位', str(error))
        # 原来为空的规格可以补填，已有值不能改写；大小写一致的写法不算改动。
        self.patch_item(self.item, {'specification': '750W'}, 200)
        self.patch_item(self.item, {'specification': '750w', 'unit': '台'}, 200)
        self.patch_item(self.item, {'specification': '1kW'}, 400)
        self.item.refresh_from_db()
        self.assertEqual((self.item.unit, self.item.specification), ('台', '750w'))

    # 5
    def test_exports_keep_numbers_numeric_and_escape_formulas(self):
        self.assertEqual(cell('-100.00'), '-100.00')
        self.assertEqual(cell('-2+3'), "'-2+3")
        self.assertEqual(cell('=SUM(A1)'), "'=SUM(A1)")
        response = export_rows(
            'probe',
            [
                {'code': '0012', 'amount': '-100.00', 'quantity': '5.000', 'reason': '-备注'},
                {'code': '1', 'amount': '1234567890123456.78', 'quantity': '0.001', 'reason': ''},
            ],
            'xlsx',
        )
        sheet = load_workbook(io.BytesIO(response.content)).active
        code, amount, quantity, reason = next(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual((code, amount, quantity, reason), ('0012', -100, 5, "'-备注"))
        # 超过 Excel 15 位有效数字的金额保留文本，不被截断。
        self.assertEqual(sheet.cell(row=3, column=2).value, '1234567890123456.78')
        project = self.active_project()
        self.post(
            'manager',
            f'projects/{project.pk}/revise-bom/',
            {'expected_revision': self.bom_version(project.pk), 'lines': [{'item': self.item.pk, 'quantity': '5'}]},
        )
        self.post(
            'admin', 'stocks/opening/', {'item': self.item.pk, 'quantity': '5', 'unit_cost': '10', 'reason': '期初'}
        )
        stock = Stock.objects.get()
        self.post(
            'warehouse', 'stocks/issue/', {'project': project.pk, 'stock': stock.pk, 'quantity': '2', 'reason': '领料'}
        )
        content = self.get('warehouse', 'business/moves/export/', {'file_format': 'csv'}).content.decode('utf-8-sig')
        quantities = [row['数量'] for row in csv.DictReader(io.StringIO(content))]
        self.assertIn('-2.000', quantities)

    # 7 与 9
    def test_project_edit_audits_full_changes_and_syncs_signed_sale(self):
        project = self.active_project()
        self.post(
            'manager',
            f'projects/{project.pk}/edit/',
            {'name': '二期装配线', 'due_date': '2026-12-31', 'requirements': '增加视觉检测', 'reason': '客户确认变更'},
        )
        log = AuditLog.objects.filter(operation='project.edit', resource=f'project:{project.pk}').get()
        self.assertEqual(log.detail['before'], {'name': '自动装配线', 'due_date': None, 'requirements': ''})
        self.assertEqual(
            log.detail['after'], {'name': '二期装配线', 'due_date': '2026-12-31', 'requirements': '增加视觉检测'}
        )
        self.assertEqual(sorted(log.detail['sale_synced']), ['due_date', 'name', 'requirements'])
        sale = SalesOrder.objects.get(project=project)
        self.assertEqual(
            (sale.name, sale.due_date, sale.requirements), ('二期装配线', date(2026, 12, 31), '增加视觉检测')
        )

    # 8
    def test_payment_and_bank_dates_cannot_be_in_future_and_reversal_follows_original(self):
        project = self.active_project()
        entry = Entry.objects.get(project=project, kind='receivable')
        today = timezone.localdate()
        tomorrow = (today + timedelta(days=1)).isoformat()
        self.post(
            'finance', f'entries/{entry.pk}/pay/', {'amount': '100', 'date': tomorrow, 'reason': '收款'}, status=400
        )
        payment = self.post(
            'finance',
            f'entries/{entry.pk}/pay/',
            {'amount': '100', 'date': (today - timedelta(days=5)).isoformat(), 'reason': '收款'},
        )['id']
        earlier = (today - timedelta(days=10)).isoformat()
        error = self.post('finance', f'payments/{payment}/reverse/', {'date': earlier, 'reason': '错账'}, status=400)
        self.assertIn('date', error)
        self.post('finance', f'payments/{payment}/reverse/', {'date': today.isoformat(), 'reason': '错账'})
        bank = {
            'amount': '100',
            'date': tomorrow,
            'account': 'A1',
            'reference': 'R1',
            'counterparty': '客户',
            'reason': '到账',
        }
        self.post('finance', 'bank-records/', bank, status=400)
        self.post('finance', 'bank-records/', {**bank, 'date': today.isoformat()}, status=201)

    # 10
    def test_issuing_another_projects_received_stock_requires_confirmation(self):
        mine, theirs = self.active_project(), self.active_project()
        for project in (mine, theirs):
            self.post(
                'manager',
                f'projects/{project.pk}/revise-bom/',
                {'expected_revision': self.bom_version(project.pk), 'lines': [{'item': self.item.pk, 'quantity': '5'}]},
            )
        purchase = self.purchase(theirs)
        self.receive(purchase, '5')
        stock = Stock.objects.get()
        issue = {'stock': stock.pk, 'reason': '领料'}
        # 为自己采购的到货不需要确认。
        self.post('warehouse', 'stocks/issue/', {**issue, 'project': theirs.pk, 'quantity': '2'})
        conflict = self.post('warehouse', 'stocks/issue/', {**issue, 'project': mine.pk, 'quantity': '1'}, status=409)
        self.assertIn(theirs.code, str(conflict['detail']))
        self.post('warehouse', 'stocks/issue/', {**issue, 'project': mine.pk, 'quantity': '1', 'confirm_shared': True})
        log = AuditLog.objects.filter(operation='stock.issue').first()
        self.assertEqual(log.detail['shared_from'], [{'project': theirs.code, 'quantity': '3.000'}])
        self.post(
            'warehouse',
            'stocks/issue/',
            {**issue, 'project': mine.pk, 'quantity': '1', 'confirm_shared': 'yes'},
            status=400,
        )

    # 12
    def test_purchase_warranty_months_drive_warranty_and_contract(self):
        project = self.active_project()
        payload = {
            'project': project.pk,
            'supplier': self.supplier.pk,
            'due_date': TODAY,
            'lines': [{'item': self.item.pk, 'quantity': '5', 'unit_price': '100'}],
        }
        self.post('purchaser', 'purchases/', {**payload, 'warranty_months': '0'}, status=400)
        purchase = PurchaseOrder.objects.get(
            pk=self.post('purchaser', 'purchases/', {**payload, 'warranty_months': '24'}, status=201)['id']
        )
        default = PurchaseOrder.objects.get(pk=self.post('purchaser', 'purchases/', payload, status=201)['id'])
        self.assertEqual((purchase.warranty_months, default.warranty_months), (24, 12))
        preview = self.get(
            'purchaser', f'business/purchases/{purchase.pk}/contract-preview/', {'version': 'current'}
        ).data
        self.assertIn('验收合格之日起24个月', preview['clauses']['quality'][2])
        default_preview = self.get(
            'purchaser', f'business/purchases/{default.pk}/contract-preview/', {'version': 'current'}
        ).data
        self.assertIn('验收合格之日起一年', default_preview['clauses']['quality'][2])
        self.post('purchaser', f'purchases/{purchase.pk}/submit/')
        self.post('manager', f'purchases/{purchase.pk}/approve/')
        self.receive(purchase, '5', received_date=TODAY, backfill_reason='根据历史签收单补录质保起算日期')
        listing = self.get('purchaser', f'business/purchases/{purchase.pk}/warranty/').data
        self.assertEqual(listing['warranty_months'], 24)
        self.assertEqual(listing['receipts'][0]['warranty_end'], date(2028, 9, 9))

    # BP-03
    def test_purchase_return_of_shared_arrival_requires_confirmation(self):
        mine, theirs = self.active_project(), self.active_project()
        for project in (mine, theirs):
            self.post(
                'manager',
                f'projects/{project.pk}/revise-bom/',
                {'expected_revision': self.bom_version(project.pk), 'lines': [{'item': self.item.pk, 'quantity': '5'}]},
            )
        own = self.purchase(mine)
        self.receive(own, '5')
        stock = Stock.objects.get()
        self.post(
            'warehouse', 'stocks/issue/', {'project': mine.pk, 'stock': stock.pk, 'quantity': '5', 'reason': '领料'}
        )
        other = self.purchase(theirs)
        self.receive(other, '5')
        receipt = StockMove.objects.get(kind='receipt', project=mine)
        url = f'moves/{receipt.pk}/return-purchase/'
        conflict = self.post('warehouse', url, {'quantity': '2', 'reason': '退货'}, status=409)
        self.assertIn(theirs.code, str(conflict['detail']))
        self.assertIn('确认占用其他项目到货', str(conflict['detail']))
        self.post('warehouse', url, {'quantity': '2', 'reason': '退货', 'confirm_shared': 'yes'}, status=400)
        self.assertEqual(Stock.objects.get().quantity, Decimal('5'))
        self.post('warehouse', url, {'quantity': '2', 'reason': '退货', 'confirm_shared': True})
        self.assertEqual(Stock.objects.get().quantity, Decimal('3'))
        log = AuditLog.objects.filter(operation='stock.return_purchase').first()
        self.assertEqual(log.detail['shared_from'], [{'project': theirs.code, 'quantity': '5.000'}])

    # BP-03
    def test_purchase_return_of_own_unissued_arrival_needs_no_confirmation(self):
        project = self.active_project()
        self.post(
            'manager',
            f'projects/{project.pk}/revise-bom/',
            {'expected_revision': self.bom_version(project.pk), 'lines': [{'item': self.item.pk, 'quantity': '5'}]},
        )
        self.receive(self.purchase(project), '5')
        receipt = StockMove.objects.get(kind='receipt')
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '5', 'reason': '退货'})
        log = AuditLog.objects.filter(operation='stock.return_purchase').first()
        self.assertEqual(log.detail['shared_from'], [])

    def warranty_receipt(self):
        project = self.active_project()
        purchase = self.purchase(project)
        self.receive(purchase, '5')
        return purchase, StockMove.objects.get(kind='receipt')

    def warranty_case(self, purchase, receipt, qty):
        self.post(
            'purchaser',
            f'purchases/{purchase.pk}/warranty/',
            {
                'receipt': receipt.pk,
                'date': timezone.localdate().isoformat(),
                'quantity': qty,
                'description': '故障',
            },
        )
        return PurchaseWarranty.objects.latest('pk')

    def update_case(self, purchase, case, status=200, role='purchaser', **extra):
        case.refresh_from_db()
        return self.post(
            role,
            f'purchases/{purchase.pk}/warranty/',
            {
                'case': case.pk,
                'expected_updated_at': case.updated_at.isoformat().replace('+00:00', 'Z'),
                'status': 'repairing',
                'response': '处理中',
                **extra,
            },
            status=status,
        )

    # BP-06
    def test_warranty_quantity_excludes_returned_goods(self):
        purchase, receipt = self.warranty_receipt()
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '3', 'reason': '退货'})
        url = f'purchases/{purchase.pk}/warranty/'
        body = {
            'receipt': receipt.pk,
            'date': timezone.localdate().isoformat(),
            'quantity': '3',
            'description': '故障',
        }
        error = self.post('purchaser', url, body, status=409)
        self.assertIn('扣除已退货', str(error['detail']))
        self.post('purchaser', url, {**body, 'quantity': '2'})

    # BP-06
    def test_warranty_returned_move_cannot_be_shared_or_exceed_case_quantity(self):
        purchase, receipt = self.warranty_receipt()
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '1', 'reason': '退货'})
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '2', 'reason': '退货'})
        small, large = StockMove.objects.filter(kind='purchase_return').order_by('pk')
        first = self.warranty_case(purchase, receipt, '1')
        second = self.warranty_case(purchase, receipt, '1')
        error = self.update_case(purchase, first, status=400, returned=large.pk)
        self.assertIn('returned', error)
        self.update_case(purchase, first, returned=small.pk)
        error = self.update_case(purchase, second, status=409, returned=small.pk)
        self.assertIn('已关联其他质保单', str(error['detail']))
        # 同一质保单可重复保存自己已关联的流水。
        self.update_case(purchase, first, returned=small.pk)

    # BP-06
    def test_warranty_case_may_link_partial_return(self):
        purchase, receipt = self.warranty_receipt()
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '1', 'reason': '退货'})
        returned = StockMove.objects.get(kind='purchase_return')
        case = self.warranty_case(purchase, receipt, '2')
        self.update_case(purchase, case, returned=returned.pk)

    # BP-06
    def test_warranty_expense_cannot_be_shared_between_cases(self):
        purchase, receipt = self.warranty_receipt()
        expense = Entry.objects.create(
            project=purchase.project,
            kind='expense',
            title='返修运费',
            amount=Decimal('10'),
            due_date=timezone.localdate(),
        )
        first = self.warranty_case(purchase, receipt, '1')
        second = self.warranty_case(purchase, receipt, '1')
        self.update_case(purchase, first, expense=expense.pk, role='admin')
        error = self.update_case(purchase, second, status=409, expense=expense.pk, role='admin')
        self.assertIn('已关联其他质保单', str(error['detail']))

    # 13
    def test_new_location_requires_confirmation_and_is_normalized(self):
        project = self.active_project()
        purchase = self.purchase(project)
        error = self.receive(purchase, '1', status=400, location='Ａ区 ')
        self.assertIn('A区', str(error['location']))
        self.receive(purchase, '1', location='Ａ区 ', new_location=True)
        self.assertTrue(Stock.objects.filter(location='A区').exists())
        self.receive(purchase, '1', location='A区')
        self.assertEqual(Stock.objects.get().quantity, Decimal('2'))
        self.assertEqual(self.get('warehouse', 'business/stocks/locations/').data, ['A区'])
        # 期初由管理员在上线时建立库位，不需要逐个确认。
        other = Item.objects.create(code='I2', name='气缸')
        self.post(
            'admin',
            'stocks/opening/',
            {'item': other.pk, 'location': 'B 区', 'quantity': '1', 'unit_cost': '1', 'reason': '期初'},
        )
        self.assertEqual(self.get('warehouse', 'business/stocks/locations/').data, ['A区', 'B 区'])

    # 15
    def test_audit_log_shows_actor_name_and_filters(self):
        self.users['manager'].display_name = '王经理'
        self.users['manager'].save()
        project = self.active_project()
        self.post('manager', f'projects/{project.pk}/edit/', {'name': '改名', 'reason': '更正名称'})
        results = self.get('admin', 'core/audit/', {'search': '王经理', 'operation': 'project.edit'}).data['results']
        self.assertEqual([(row['operation'], row['actor_name']) for row in results], [('project.edit', '王经理')])
        today = timezone.localdate().isoformat()
        self.assertTrue(self.get('admin', 'core/audit/', {'date_from': today, 'date_to': today}).data['count'])
        self.assertEqual(self.get('admin', 'core/audit/', {'date_to': '2000-01-01'}).data['count'], 0)
        self.get('manager', 'core/audit/', status=403)


class ServiceWarrantyTests(ExecutionFixtures, TestCase):
    def setUp(self):
        self.setup_execution()

    # 11
    def test_out_of_warranty_free_service_requires_reason(self):
        project = self.ready_project(warranty=0)
        self.stock_and_issue(project)
        delivery = self.deliver(project)
        self.accept(delivery)
        tomorrow = date.fromisoformat(TODAY) + timedelta(days=1)
        payload = {
            'delivery': delivery.pk,
            'date': tomorrow.isoformat(),
            'title': '过保维修',
            'assignee': self.users['member'].pk,
            'fee': '0',
        }
        with patch('apps.business.services.execution.timezone.localdate', return_value=tomorrow):
            error = self.post('manager', f'projects/{project.pk}/service/', payload, status=400)
            self.assertIn('free_reason', error)
            task_id = self.post(
                'manager', f'projects/{project.pk}/service/', {**payload, 'free_reason': '商务赠送一次免费维修'}
            )['id']
        task = Task.objects.get(pk=task_id)
        self.assertFalse(Entry.objects.filter(task=task).exists())
        log = AuditLog.objects.get(operation='service.create', resource=f'task:{task.pk}')
        self.assertEqual(log.detail['free_reason'], '商务赠送一次免费维修')
