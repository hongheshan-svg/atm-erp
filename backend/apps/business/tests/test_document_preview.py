import io
import re
import struct
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.http import QueryDict
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import Workbook
from rest_framework.exceptions import NotFound, ValidationError
from rest_framework.test import APIClient

from apps.business import document_preview as preview
from apps.business.models import Document, SalesOrder
from apps.core.models import ActionReceipt, AuditLog, CodeRule

from .test_commercial_chain import BusinessFixtures

PDF = b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\n%%EOF\n'


def xlsx_content(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def zip_content(parts):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return stream.getvalue()


def docx_content(body):
    return zip_content(
        {
            'word/document.xml': (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f'<w:body>{body}</w:body></w:document>'
            )
        }
    )


def xls_content():
    """Small BIFF8 workbook, avoiding a writer dependency solely for this fixture."""

    def record(code, data=b''):
        return struct.pack('<HH', code, len(data)) + data

    def bof(kind):
        return record(0x0809, struct.pack('<HHHHII', 0x0600, kind, 0x0DBB, 1996, 0x41, 6))

    globals_ = bof(5) + record(0x0042, struct.pack('<H', 1200))
    globals_ += record(0x00E0, b'\0' * 20)
    name = b'Data'
    offset = len(globals_) + 4 + 8 + len(name) + 4
    globals_ += record(0x0085, struct.pack('<IBBBB', offset, 0, 0, len(name), 0) + name)
    globals_ += record(0x000A)
    sheet = bof(0x10) + record(0x0200, struct.pack('<IIHHH', 0, 1, 0, 2, 0))
    sheet += record(0x0203, struct.pack('<HHHd', 0, 0, 0, 12.5))
    sheet += record(0x0006, struct.pack('<HHHdHIH', 0, 1, 0, 42, 0, 0, 0))
    return globals_ + sheet + record(0x000A)


class DocumentPreviewParsingTests(SimpleTestCase):
    def render(self, name, content, query=None):
        document = SimpleNamespace(pk=7, original_name=name, file=ContentFile(content))
        return preview.preview(document, query or {})

    def assert_invalid(self, name, content, query=None, message=None):
        with self.assertRaises(ValidationError) as caught:
            self.render(name, content, query)
        text = str(caught.exception.detail)
        if message:
            self.assertIn(message, text)
        self.assertNotIn('Traceback', text)
        self.assertNotIn('BadZipFile', text)
        self.assertTrue(any('\u4e00' <= char <= '\u9fff' for char in text))

    def test_xlsx_visible_sheets_cached_formulas_formats_and_pagination(self):
        workbook = Workbook()
        workbook.active.title = '隐藏页'
        workbook.active.sheet_state = 'hidden'
        sheet = workbook.create_sheet('明细')
        sheet.append([123, 12.5, 0.125, datetime(2026, 10, 26), '=6*7', '=1+1', '<script>示例</script>'])
        for cell, number_format in [('A1', '000000'), ('B1', '0.00'), ('C1', '0.0%'), ('D1', 'yyyy-mm-dd')]:
            sheet[cell].number_format = number_format
        for number in range(2, 206):
            sheet.append([number])
        workbook.create_sheet('汇总').append(['另一页'])
        original = xlsx_content(workbook)
        with zipfile.ZipFile(io.BytesIO(original)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts['xl/worksheets/sheet2.xml'] = re.sub(
            rb'<f>6\*7</f><v(?:\s*/>|></v>)', b'<f>6*7</f><v>42</v>', parts['xl/worksheets/sheet2.xml']
        )
        self.assertIn(b'<v>42</v>', parts['xl/worksheets/sheet2.xml'])
        content = zip_content(parts)
        result = self.render('报价.XLSX', content)
        self.assertEqual(result['name'], '报价.XLSX')
        self.assertEqual(result['kind'], 'spreadsheet')
        self.assertEqual(result['sheets'], [{'index': 1, 'name': '明细'}, {'index': 2, 'name': '汇总'}])
        self.assertEqual((result['sheet'], result['page'], result['page_size']), (1, 1, 100))
        self.assertEqual((result['total_rows'], result['total_pages']), (205, 3))
        self.assertEqual(result['columns'], list('ABCDEFG'))
        self.assertEqual(
            result['rows'][0],
            {'number': 1, 'cells': ['000123', '12.50', '12.5%', '2026-10-26', '42', '', '<script>示例</script>']},
        )
        self.assertIn('尚未计算', ''.join(result['warnings']))
        self.assertIn('图表', ''.join(result['warnings']))
        self.assertEqual(len(result['rows']), 100)
        second = self.render('报价.xlsx', content, {'sheet': '1', 'page': '2'})
        self.assertEqual([row['number'] for row in second['rows']], list(range(101, 201)))
        self.assertEqual(second['rows'][0]['cells'], ['101', '', '', '', '', '', ''])
        last = self.render('报价.xlsx', content, {'page': '3'})
        self.assertEqual([row['number'] for row in last['rows']], list(range(201, 206)))
        summary = self.render('报价.xlsx', content, {'sheet': '2'})
        self.assertEqual(summary['rows'], [{'number': 1, 'cells': ['另一页']}])
        self.assert_invalid('报价.xlsx', content, {'sheet': '0'}, '已隐藏')
        self.assert_invalid('报价.xlsx', content, {'page': '4'}, '页码')

    def test_xlsx_incorrect_dimensions_and_bounded_rows_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet['A1'] = '首行'
        sheet['A101'] = '后页'
        sheet['CW102'] = '超出列数'
        content = xlsx_content(workbook)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(b'A1:CW102', b'A1:A1')
        result = self.render('sparse.xlsx', zip_content(parts), {'page': 2})
        self.assertEqual(result['total_rows'], 102)
        self.assertEqual(result['rows'][0], {'number': 101, 'cells': ['后页']})
        self.assertIn('100 列', ''.join(result['warnings']))
        with patch.object(preview, 'MAX_ROWS', 100):
            result = self.render('bounded.xlsx', content)
        self.assertEqual(result['total_rows'], 100)
        self.assertEqual(result['total_pages'], 1)
        self.assertIn('100 行', ''.join(result['warnings']))

    def test_xls_reads_cached_formula_values(self):
        result = self.render('旧报表.xls', xls_content())
        self.assertEqual(result['sheets'], [{'index': 0, 'name': 'Data'}])
        self.assertEqual(result['rows'], [{'number': 1, 'cells': ['12.5', '42']}])
        self.assertIn('未计算公式', ''.join(result['warnings']))

    def test_csv_gb18030_quotes_leading_zero_and_formula_is_plain_text(self):
        content = '编码,说明,备注\r\n000123,"夹具,型号",=1+1\r\n000124,"第一行\n第二行",<img src=x>\r\n'
        result = self.render('物料.csv', content.encode('gb18030'))
        self.assertEqual(result['total_rows'], 3)
        self.assertEqual(result['rows'][1]['cells'], ['000123', '夹具,型号', '=1+1'])
        self.assertEqual(result['rows'][2]['cells'], ['000124', '第一行\n第二行', '<img src=x>'])
        self.assertIn('GB18030', ''.join(result['warnings']))
        with patch.object(preview, 'MAX_ROWS', 2), patch.object(preview, 'MAX_COLUMNS', 2):
            bounded = self.render('large.csv', content.encode())
        self.assertEqual(bounded['total_rows'], 2)
        self.assertEqual(bounded['columns'], ['A', 'B'])
        self.assertIn('2 行', ''.join(bounded['warnings']))
        self.assertIn('2 列', ''.join(bounded['warnings']))

    def test_docx_plain_text_tables_tabs_breaks_and_pagination(self):
        paragraph = '<w:p><w:r><w:t>&lt;script&gt;测试&lt;/script&gt;</w:t><w:tab/><w:t>后文</w:t><w:br/></w:r></w:p>'
        table = '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>物料</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>001</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        content = docx_content(paragraph + table + paragraph * 100)
        result = self.render('合同.docx', content)
        self.assertEqual(result['kind'], 'document')
        self.assertEqual((result['page'], result['total_pages']), (1, 2))
        self.assertEqual(result['blocks'][0], {'kind': 'paragraph', 'text': '<script>测试</script>\t后文\n'})
        self.assertEqual(result['blocks'][1], {'kind': 'table', 'rows': [['物料', '001']]})
        self.assertEqual(len(self.render('合同.docx', content, {'page': '2'})['blocks']), 2)
        self.assert_invalid('合同.docx', content, {'page': '3'}, '页码')
        with patch.object(preview, 'MAX_BLOCKS', 1):
            limited = self.render('合同.docx', content)
        self.assertEqual(len(limited['blocks']), 1)
        self.assertIn('前 1', ''.join(limited['warnings']))

    def test_text_utf8_gb18030_and_text_budgets(self):
        self.assertEqual(self.render('说明.txt', b'\xef\xbb\xbf<script>text</script>')['text'], '<script>text</script>')
        result = self.render('说明.txt', '中文资料'.encode('gb18030'))
        self.assertEqual(result['text'], '中文资料')
        with patch.object(preview, 'MAX_TEXT', 5):
            result = self.render('说明.txt', b'123456789')
        self.assertEqual(result['text'], '12345')
        self.assertIn('截断', ''.join(result['warnings']))
        with patch.object(preview, 'MAX_TEXT', 5):
            result = self.render('说明.csv', b'123456,abcdef\nmore,text\n')
        self.assertEqual(sum(len(c) for row in result['rows'] for c in row['cells']), 5)
        self.assertIn('截断', ''.join(result['warnings']))
        self.assert_invalid('binary.txt', b'abc\x00def', message='非文本')
        self.assert_invalid('unknown.txt', b'\xff', message='编码')

    def test_binary_whitelist_mime_and_root_relative_file_url(self):
        cases = [
            ('a.pdf', PDF, 'application/pdf', 'pdf'),
            ('a.png', b'\x89PNG\r\n\x1a\ncontent', 'image/png', 'image'),
            ('a.jpg', b'\xff\xd8\xff\xe0content', 'image/jpeg', 'image'),
            ('a.jpeg', b'\xff\xd8\xff\xe0content', 'image/jpeg', 'image'),
            ('a.gif', b'GIF89acontent', 'image/gif', 'image'),
            ('a.webp', b'RIFF\x04\x00\x00\x00WEBPVP8 content', 'image/webp', 'image'),
            ('a.bmp', b'BM' + b'\0' * 12, 'image/bmp', 'image'),
        ]
        for name, content, mime, kind in cases:
            with self.subTest(name=name):
                self.assertEqual(preview.inline_type(name, content), mime)
                result = self.render(name, content)
                self.assertEqual(result['kind'], kind)
                self.assertEqual(result['file_url'], '/api/business/documents/7/preview-file/')
                self.assertEqual(set(result), {'name', 'kind', 'warnings', 'file_url'})
                self.assert_invalid(name, b'<html><script>alert(1)</script></html>', message='扩展名')
        self.assert_invalid('a.png', PDF, message='扩展名')

    def test_unsupported_formats_offer_export_pdf_without_inline_url(self):
        for name in ('legacy.doc', 'drawing.dwg', 'active.svg', 'page.html', 'archive.zip'):
            with self.subTest(name=name):
                result = self.render(name, b'<script>not served inline</script>')
                self.assertEqual(result['kind'], 'unsupported')
                self.assertIn('导出 PDF', result['message'])
                self.assertIn('下载原文件', result['message'])
                self.assertNotIn('file_url', result)
                with self.assertRaises(ValidationError):
                    preview.inline_type(name, b'<script>')

    def test_invalid_queries_are_chinese_400(self):
        for key, value in [
            ('page', '0'),
            ('page', '-1'),
            ('page', '1.5'),
            ('page', '9999999'),
            ('sheet', '-1'),
            ('sheet', 'a'),
            ('sheet', '999'),
            ('page', '2'),
        ]:
            with self.subTest(key=key, value=value):
                self.assert_invalid('a.csv', b'1,2\n', {key: value})
        self.assert_invalid('a.csv', b'1', QueryDict('page=1&page=2'), '只能提交一次')
        self.assert_invalid('a.csv', b'1', QueryDict('sheet=0&sheet=0'), '只能提交一次')

    def test_missing_empty_oversized_corrupt_and_encrypted_input(self):
        file = SimpleNamespace(open=lambda mode: (_ for _ in ()).throw(FileNotFoundError('/private/location')))
        with self.assertRaises(NotFound) as caught:
            preview.read_content(SimpleNamespace(file=file))
        self.assertNotIn('/private', str(caught.exception))
        self.assert_invalid('a.csv', b'', message='附件为空')
        with patch.object(preview, 'MAX_BYTES', 4):
            self.assert_invalid('a.txt', b'12345', message='20 MB')
        for name in ('a.xlsx', 'a.xls', 'a.docx'):
            self.assert_invalid(name, b'not an office file', message='损坏')
        self.assert_invalid('a.csv', b'"unclosed', message='损坏')
        self.assert_invalid('a.docx', zip_content({'word/document.xml': '<broken>'}))
        workbook = Workbook()
        valid = xlsx_content(workbook)
        with patch.object(preview.xlrd, 'open_workbook', side_effect=preview.xlrd.XLRDError('Workbook is encrypted')):
            self.assert_invalid('a.xls', b'fake encrypted input', message='加密')
        info = zipfile.ZipInfo('word/document.xml')
        info.flag_bits = 1
        with patch.object(zipfile.ZipFile, 'infolist', return_value=[info]):
            self.assert_invalid('a.xlsx', valid, message='加密')

    def test_zip_limits_and_xml_entities_are_rejected(self):
        valid = docx_content('<w:p/>')
        with patch.object(preview, 'MAX_EXPANDED_BYTES', 10):
            self.assert_invalid('a.docx', valid, message='解压后过大')
        with patch.object(preview, 'MAX_ZIP_MEMBERS', 1):
            self.assert_invalid('a.docx', zip_content({'a.xml': '<a/>', 'b.xml': '<b/>'}), message='部件过多')
        with patch.object(preview, 'MAX_XML_ELEMENTS', 1):
            self.assert_invalid('a.docx', valid, message='过于复杂')
        hostile = b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'
        for name in ('a.docx', 'a.xlsx'):
            self.assert_invalid(name, zip_content({'word/document.xml': hostile}), message='损坏')
        with patch('socket.create_connection', side_effect=AssertionError('network access')) as connect:
            result = self.render('a.docx', docx_content('<w:p><w:r><w:t>safe</w:t></w:r></w:p>'))
            self.assertEqual(result['blocks'][0]['text'], 'safe')
            connect.assert_not_called()


class DocumentPreviewAPITests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()
        self.project = self.active_project()
        self.directory = tempfile.TemporaryDirectory(prefix='lean-preview-test-')
        self.addCleanup(self.directory.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.directory.name)
        self.override.enable()
        self.addCleanup(self.override.disable)

    def document(self, *, name='技术资料.pdf', content=PDF, category='drawing', **owner):
        return Document.objects.create(
            **(owner or {'project': self.project}),
            category=category,
            original_name=name,
            file=ContentFile(content, name=name),
            size=len(content),
            sha256='test',
            created_by=self.users['manager'],
        )

    def read(self, document, action='preview', role='member', query=None, status=200):
        response = self.clients[role].get(f'/api/business/documents/{document.pk}/{action}/', query or {})
        self.assertEqual(response.status_code, status, getattr(response, 'data', None))
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
        if action == 'preview-file':
            self.assertEqual(response['Content-Security-Policy'], "default-src 'none'; frame-ancestors 'none'; sandbox")
        if response.streaming:
            # Consume via the test client's wrapper, which safely closes the stream
            # without sending request_finished inside this TestCase's transaction.
            response.preview_bytes = b''.join(response.streaming_content)
        return response

    def test_all_attachment_owners_and_roles_share_download_scope(self):
        self.project.members.add(
            *(self.users[role] for role in ('mechanical_engineer', 'electrical_engineer', 'production_manager'))
        )
        purchase = self.purchase(self.project)
        sale = SalesOrder.objects.get(project=self.project)
        documents = [
            self.document(),
            self.document(category='contract'),
            self.document(category='receipt'),
            self.document(category='other', sale=sale),
            self.document(category='contract', purchase=purchase),
        ]
        for role, client in self.clients.items():
            for document in documents:
                with self.subTest(role=role, document=document.pk):
                    downloaded = client.get(f'/api/business/documents/{document.pk}/download/')
                    status = downloaded.status_code
                    self.assertIn(status, (200, 404))
                    if downloaded.streaming:
                        b''.join(downloaded.streaming_content)
                    self.read(document, role=role, status=status)
                    self.read(document, 'preview-file', role=role, status=status)
        self.read(documents[0], role='mechanical_engineer')
        self.read(documents[1], role='member', status=404)
        self.read(documents[3], role='sales_manager', status=404)
        self.read(documents[4], role='purchase_manager')

    def test_cross_project_and_combined_roles_do_not_expand_preview_scope(self):
        other = self.active_project()
        other.manager = self.users['admin']
        other.save(update_fields=['manager'])
        other.members.clear()
        technical = self.document(project=other)
        contract = self.document(project=other, category='contract')
        for role in ('manager', 'member', 'mechanical_engineer', 'production_manager'):
            self.users[role].additional_roles = ['purchaser']
            self.users[role].save(update_fields=['additional_roles'])
            for document in (technical, contract):
                downloaded = self.clients[role].get(f'/api/business/documents/{document.pk}/download/')
                status = downloaded.status_code
                if downloaded.streaming:
                    b''.join(downloaded.streaming_content)
                self.read(document, role=role, status=status)
                self.read(document, 'preview-file', role=role, status=status)
            self.read(contract, role=role, status=404)

    def test_anonymous_revoked_membership_removed_and_soft_deleted(self):
        document = self.document()
        anonymous = APIClient()
        denied = anonymous.get(f'/api/business/documents/{document.pk}/download/').status_code
        self.assertIn(denied, (401, 403))
        for action in ('download', 'preview', 'preview-file'):
            response = anonymous.get(f'/api/business/documents/{document.pk}/{action}/')
            self.assertEqual(response.status_code, denied)
            if action != 'download':
                self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.project.members.clear()
        for action in ('preview', 'preview-file'):
            self.read(document, action, status=404)
        self.post('manager', f'documents/{document.pk}/remove/', {'reason': '误传'})
        for action in ('preview', 'preview-file'):
            self.read(document, action, role='admin', status=404)
        self.assertTrue(Path(document.file.path).exists())
        second = self.document()
        second.soft_delete(self.users['admin'])
        self.read(second, role='admin', status=404)
        self.read(second, 'preview-file', role='admin', status=404)

    def test_preview_file_mime_inline_and_original_download_unchanged(self):
        document = self.document()
        data = self.read(document).data
        self.assertEqual(data['file_url'], f'/api/business/documents/{document.pk}/preview-file/')
        response = self.read(document, 'preview-file')
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('inline;', response['Content-Disposition'])
        self.assertEqual(response.preview_bytes, PDF)
        downloaded = self.read(document, 'download')
        self.assertEqual(downloaded['Content-Type'], 'application/octet-stream')
        self.assertIn('attachment;', downloaded['Content-Disposition'])
        self.assertEqual(downloaded.preview_bytes, PDF)
        image = self.document(name='照片.png', content=b'\x89PNG\r\n\x1a\ncontent')
        response = self.read(image, 'preview-file')
        self.assertEqual(response['Content-Type'], 'image/png')

    def test_unsupported_disguised_missing_and_invalid_queries(self):
        unsupported = self.document(name='零件.dwg')
        self.assertEqual(self.read(unsupported).data['kind'], 'unsupported')
        self.read(unsupported, 'preview-file', status=400)
        self.read(unsupported, 'download')
        disguised = self.document(content=b'<html><script>alert(1)</script></html>')
        for action in ('preview', 'preview-file'):
            self.read(disguised, action, status=400)
        missing = self.document()
        missing.file.storage.delete(missing.file.name)
        for action in ('preview', 'preview-file'):
            response = self.read(missing, action, status=404)
            self.assertNotIn(self.directory.name, str(response.data))
        csv_document = self.document(name='物料.csv', content=b'001,part\n')
        self.read(csv_document, query={'page': '0'}, status=400)
        self.read(csv_document, query={'sheet': 'bad'}, status=400)
        self.read(csv_document, 'preview-file', status=400)

    def test_preview_has_no_database_or_file_side_effects(self):
        document = self.document(name='说明.csv', content=b'001,part\n')

        def snapshot():
            return (
                list(Document.all_objects.values()),
                ActionReceipt.objects.count(),
                AuditLog.objects.count(),
                list(CodeRule.objects.values()),
                {str(path): path.read_bytes() for path in Path(self.directory.name).rglob('*') if path.is_file()},
            )

        before = snapshot()
        self.read(document)
        self.read(document)
        self.read(document, query={'page': '999'}, status=400)
        self.read(document, 'preview-file', status=400)
        self.assertEqual(snapshot(), before)
