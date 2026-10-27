import csv
import io
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.business.models import BOMLine, Item, Project
from apps.core.models import ActionReceipt, AuditLog

from .test_commercial_chain import BusinessFixtures


class BOMExportTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()
        self.project = self.active_project()
        self.item.code, self.item.part_type, self.item.product_category = '000123', 'standard', '21'
        self.item.save()
        self.lines = BOMLine.objects.bulk_create(
            [
                BOMLine(project=self.project, item=self.item, assembly_unit=f'单元{i}', quantity='2.500')
                for i in range(81)
            ]
        )
        custom = Item.objects.create(code='CUSTOM', name='非标件', part_type='custom', product_category='11')
        self.other = BOMLine.objects.create(project=self.project, item=custom, quantity=3)

    def export(self, ids=None, role='purchaser', fmt='csv', **extra):
        return self.clients[role].post(
            '/api/business/bom/export/',
            {
                'project': self.project.pk,
                'ids': ids if ids is not None else [row.pk for row in self.lines],
                'file_format': fmt,
                **extra,
            },
            format='json',
        )

    def test_selected_81_rows_are_exact_without_nonstandard_or_business_writes(self):
        before = (AuditLog.objects.count(), ActionReceipt.objects.count(), BOMLine.objects.count())
        response = self.export()
        self.assertEqual(response.status_code, 200)
        rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 81)
        self.assertEqual({row['物料编码'] for row in rows}, {'000123'})
        self.assertEqual({row['物料类别'] for row in rows}, {'标准件'})
        self.assertEqual({row['产品编码类别'] for row in rows}, {'无图·标准件'})
        self.assertEqual({row['单元'] for row in rows}, {line.assembly_unit for line in self.lines})
        self.assertIn('变更说明', rows[0])
        for key in ['can_edit_bom', 'change_note', '记录ID', '项目ID', '物料ID', '单位成本']:
            self.assertNotIn(key, rows[0])
        self.assertEqual(before, (AuditLog.objects.count(), ActionReceipt.objects.count(), BOMLine.objects.count()))

    def test_xlsx_preserves_codes_quantities_units_and_escapes_formulas(self):
        self.item.name = '=HYPERLINK("bad")'
        self.item.save()
        response = self.export([self.lines[1].pk, self.lines[0].pk], fmt='xlsx')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 3)
        headers = [cell.value for cell in sheet[1]]
        rows = [dict(zip(headers, [cell.value for cell in row], strict=True)) for row in list(sheet)[1:]]
        self.assertEqual([row['单元'] for row in rows], ['单元1', '单元0'])
        self.assertEqual(rows[0]['物料编码'], '000123')
        self.assertEqual(rows[0]['需求数量'], 2.5)
        self.assertTrue(rows[0]['物料名称'].startswith("'="))
        self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))

    def test_project_scope_and_hidden_rows_fail_the_whole_selection(self):
        other_project = Project.objects.create(
            code='PRIVATE', name='另一项目', customer=self.customer, manager=self.users['admin']
        )
        outside = BOMLine.objects.create(project=other_project, item=self.item, quantity=1)
        # Even globally authorized purchasers cannot mix a different project's row into this export.
        self.assertEqual(self.export([self.lines[0].pk, outside.pk]).status_code, 409)
        self.project.members.clear()
        self.assertEqual(self.export([self.lines[0].pk], role='member').status_code, 409)
        self.project.members.add(self.users['member'])
        self.assertEqual(self.export([self.lines[0].pk], role='member').status_code, 200)
        self.lines[0].soft_delete(self.users['admin'])
        self.assertEqual(self.export([self.lines[0].pk, self.lines[1].pk]).status_code, 409)
        self.assertEqual(self.export(role='sales_manager').status_code, 403)
        self.assertIn(APIClient().post('/api/business/bom/export/', {}, format='json').status_code, (401, 403))

    def test_invalid_or_empty_selection_never_falls_back_to_whole_project(self):
        for ids in ([], 'all', [True], [self.lines[0].pk] * 2, [None], [{}]):
            with self.subTest(ids=ids):
                self.assertEqual(self.export(ids).status_code, 400)
        self.assertEqual(self.export([self.lines[0].pk], extra_field=1).status_code, 400)
        for fmt in ('exe', [], {}, None):
            self.assertEqual(self.export([self.lines[0].pk], fmt=fmt).status_code, 400)
        with patch('apps.business.services.tabular.MAX_EXPORT', 1):
            self.assertEqual(self.export([row.pk for row in self.lines[:2]]).status_code, 400)

    def test_get_export_keeps_existing_project_filter_and_uses_business_columns(self):
        response = self.clients['purchaser'].get(
            '/api/business/bom/export/', {'project': self.project.pk, 'item': self.other.item_id}
        )
        self.assertEqual(response.status_code, 200)
        rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['物料类别'], '非标件')
        self.assertEqual(rows[0]['产品编码类别'], '有图·机加')
