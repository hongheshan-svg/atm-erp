import base64
import csv
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.business.models import BOMLine, Entry, Item, Stock, StockMove, Task
from apps.business.services import transfers
from apps.business.services.import_guides import BOM_COLUMNS
from apps.core.models import ActionReceipt, AuditLog, CodeRule

from .test_commercial_chain import BusinessFixtures


class ImportFeedbackTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()
        self.project = self.active_project()

    def upload(self, path, headers, rows, *, role='admin', xlsx=False):
        if xlsx:
            book = Workbook()
            book.active.append(headers)
            for row in rows:
                book.active.append(row)
            stream = io.BytesIO()
            book.save(stream)
            book.close()
            data = stream.getvalue()
        else:
            stream = io.StringIO()
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
            data = stream.getvalue().encode()
        return self.clients[role].post(
            '/api/business/' + path,
            {
                'file': SimpleUploadedFile('input.xlsx' if xlsx else 'input.csv', data),
            },
            format='multipart',
        )

    def workbook(self, result):
        return load_workbook(io.BytesIO(base64.b64decode(result['error_file']['content'])))

    def reupload(self, path, book):
        stream = io.BytesIO()
        book.save(stream)
        return self.clients['admin'].post(
            '/api/business/' + path,
            {
                'file': SimpleUploadedFile('corrected.xlsx', stream.getvalue()),
            },
            format='multipart',
        )

    def test_wrong_template_identifies_exact_resource_before_resolving_private_rows(self):
        headers = [label for label, _ in transfers.SCHEMAS['tasks'][1]]
        rows = [['任务', '设计', 'unknown-secret-person', '', 'private-project', '']]
        for xlsx in [False, True]:
            result = self.upload('projects/import-file/', headers, rows, xlsx=xlsx)
            self.assertEqual(result.status_code, 400)
            self.assertEqual(result.data['template_hint']['code'], 'tasks')
            self.assertIn('这是项目任务模板', result.data['file'])
            self.assertNotIn('unknown-secret-person', str(result.data))
            self.assertNotIn('private-project', str(result.data))
        self.assertFalse(Task.objects.exists())
        denied = self.upload('projects/import-file/', headers, rows, role='member')
        self.assertEqual(denied.status_code, 403)
        self.assertNotIn('template_hint', denied.data)

    def test_wider_bom_template_in_task_entry_is_recognized_for_xlsx(self):
        result = self.upload('tasks/import-file/', [label for label, _ in BOM_COLUMNS], [], xlsx=True)
        self.assertEqual(result.status_code, 400)
        self.assertEqual(result.data['template_hint']['code'], 'bom')

    def test_opening_inventory_guidance_moves_to_settings_without_expanding_permissions(self):
        headers = [label for label, _ in transfers.SCHEMAS['stocks'][1]]
        result = self.upload('projects/import-file/', headers, [])
        self.assertEqual(result.status_code, 400)
        self.assertEqual(result.data['template_hint']['code'], 'stocks')
        self.assertIn('设置 → 管理员初始化', result.data['template_hint']['entry'])
        allowed = self.clients['admin'].get('/api/business/stocks/import-schema/')
        self.assertEqual(allowed.status_code, 200)
        self.assertIn('设置 → 管理员初始化', allowed.data['entry'])
        denied = self.clients['manager'].get('/api/business/stocks/import-schema/')
        self.assertEqual(denied.status_code, 403)

    def test_error_workbook_preserves_legacy_headers_blanks_and_literal_values_then_revalidates(self):
        path = 'items/import-file/'
        headers = ['物料编码', '名称', '规格', '单位']
        before = ActionReceipt.objects.count()
        counter = CodeRule.objects.get(key='item').counter
        result = self.upload(path, headers, [['0007', '', '=literal text', '件'], [], ['0008', '新物料', '', '件']])
        self.assertEqual(result.status_code, 200, result.data)
        self.assertFalse(result.data['can_import'])
        self.assertIsNone(result.data['token'])
        self.assertEqual([row['row'] for row in result.data['rows']], [2, 4])
        self.assertEqual(Item.objects.count(), 1)
        self.assertEqual(ActionReceipt.objects.count(), before)
        self.assertEqual(CodeRule.objects.get(key='item').counter, counter)
        book = self.workbook(result.data)
        self.assertEqual(book.sheetnames, ['数据'])
        self.assertEqual([cell.value for cell in book.active[1]], headers)
        self.assertEqual(book.active['A2'].value, '0007')
        self.assertEqual(book.active['C2'].value, '=literal text')
        self.assertEqual(book.active['C2'].data_type, 's')
        self.assertIsNotNone(book.active['B2'].comment)
        self.assertEqual(book.active['B3'].value, None)
        book.active['B2'] = '修正物料'
        corrected = self.reupload(path, book)
        book.close()
        self.assertEqual(corrected.status_code, 200, corrected.data)
        self.assertTrue(corrected.data['can_import'], corrected.data)
        self.assertNotIn('error_file', corrected.data)
        self.assertEqual(Item.objects.count(), 1)
        self.post('admin', 'items/import-confirm/', {'token': corrected.data['token']})
        self.assertEqual(Item.objects.get(code='0007').name, '修正物料')

    def test_legacy_bom_failed_rows_remain_visible_and_file_can_be_corrected(self):
        path = f'projects/{self.project.pk}/import-preview/'
        headers = ['物料编码', '数量', '变更说明']
        result = self.upload(path, headers, [[self.item.code, '0', '首版']])
        self.assertEqual(result.status_code, 200, result.data)
        self.assertFalse(result.data['can_import'])
        self.assertEqual(result.data['lines'][0]['row'], 2)
        self.assertEqual(result.data['lines'][0]['quantity'], '0')
        book = self.workbook(result.data)
        self.assertIn('正数', book.active['B2'].comment.text)
        book.active['B2'] = 3
        corrected = self.reupload(path, book)
        book.close()
        self.assertTrue(corrected.data['can_import'], corrected.data)
        self.assertFalse(BOMLine.objects.exists())

    def test_new_bom_errors_do_not_generate_material_codes_and_keep_template_columns(self):
        path = f'projects/{self.project.pk}/import-preview/'
        rows = [['', '装配单元', '2', '首版', '', '', '', '新物料', '', '', '', '21', '', '', '件']]
        counter = CodeRule.objects.get(key='item').counter
        result = self.upload(path, [label for label, _ in BOM_COLUMNS], rows)
        self.assertFalse(result.data['can_import'], result.data)
        self.assertNotIn('token', result.data)
        self.assertEqual(Item.objects.count(), 1)
        self.assertEqual(CodeRule.objects.get(key='item').counter, counter)
        book = self.workbook(result.data)
        self.assertEqual(book.active.max_column, len(BOM_COLUMNS))
        self.assertIsNotNone(book.active['I2'].comment)
        book.active['I2'] = '规格-A'
        corrected = self.reupload(path, book)
        book.close()
        self.assertTrue(corrected.data['can_import'], corrected.data)
        self.assertTrue(corrected.data['token'])
        self.assertEqual(Item.objects.count(), 1)
        self.assertEqual(CodeRule.objects.get(key='item').counter, counter)

    def test_bom_feedback_requires_project_write_scope(self):
        self.users['mechanical_engineer'].roles = ['mechanical_engineer', 'purchaser']
        self.users['mechanical_engineer'].save()
        result = self.upload(
            f'projects/{self.project.pk}/import-preview/',
            ['物料编码', '数量', '变更说明'],
            [[self.item.code, '-1', '']],
            role='mechanical_engineer',
        )
        self.assertIn(result.status_code, [403, 404])
        self.assertNotIn('error_file', result.data)

    def test_resolved_person_is_shown_only_after_row_authorization(self):
        self.users['member'].display_name = '张工'
        self.users['member'].save()
        headers = [label for label, _ in transfers.SCHEMAS['tasks'][1]]
        result = self.upload('tasks/import-file/', headers, [['设计任务', '设计', '张工', '', self.project.code, '']])
        self.assertTrue(result.data['can_import'], result.data)
        self.assertEqual(result.data['rows'][0]['matches']['assignee'], '张工（member）')
        self.assertFalse(Task.objects.exists())

    def test_out_of_scope_row_adds_no_resolved_names_to_preview_or_error_workbook(self):
        self.project.manager = self.users['admin']
        self.project.name = '不可访问的后台项目名称'
        self.project.save()
        self.users['member'].display_name = '后台匹配人员名称'
        self.users['member'].save()
        before = ActionReceipt.objects.count()
        headers = [label for label, _ in transfers.SCHEMAS['tasks'][1]]
        rows = [['设计任务', '设计', 'member', '', self.project.code, '']]
        result = self.upload('tasks/import-file/', headers, rows, role='manager')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertFalse(result.data['can_import'])
        self.assertIsNone(result.data['token'])
        self.assertNotIn('matches', result.data['rows'][0])
        self.assertIn('无权访问', result.data['errors'][0]['message'])
        for name in [self.project.name, self.users['member'].display_name]:
            self.assertNotIn(name, str(result.data))
        book = self.workbook(result.data)
        self.assertEqual(book.sheetnames, ['数据'])
        self.assertEqual([cell.value or '' for cell in book.active[2]], rows[0])
        comments = '\n'.join(cell.comment.text for row in book.active for cell in row if cell.comment)
        self.assertNotIn(self.project.name, comments)
        self.assertNotIn(self.users['member'].display_name, comments)
        book.close()
        self.assertFalse(Task.objects.exists())
        self.assertEqual(ActionReceipt.objects.count(), before)

    def test_mixed_error_workbook_uses_source_cells_not_authorized_display_matches(self):
        self.users['member'].display_name = '后台匹配名称'
        self.users['member'].save()
        headers = [label for label, _ in transfers.SCHEMAS['tasks'][1]]
        rows = [
            ['有效任务', '设计', 'member', '', self.project.code, ''],
            ['', '设计', 'member', '', self.project.code, ''],
        ]
        result = self.upload('tasks/import-file/', headers, rows)
        self.assertEqual(result.status_code, 200, result.data)
        self.assertFalse(result.data['can_import'])
        self.assertIsNone(result.data['token'])
        self.assertEqual(result.data['rows'][0]['matches']['assignee'], '后台匹配名称（member）')
        self.assertNotIn('matches', result.data['rows'][1])
        book = self.workbook(result.data)
        self.assertEqual(book.sheetnames, ['数据'])
        for index, source in enumerate(rows, 2):
            self.assertEqual([cell.value or '' for cell in book.active[index]], source)
        comments = '\n'.join(cell.comment.text for row in book.active for cell in row if cell.comment)
        self.assertNotIn('后台匹配名称', comments)
        book.close()
        self.assertFalse(Task.objects.exists())

    def id_preview(self, resource, data, role='admin'):
        columns = transfers.SCHEMAS[resource][1]
        return self.upload(
            f'{resource}/import-file/',
            [label for label, _ in columns],
            [[data.get(key, '') for _, key in columns]],
            role=role,
        )

    def test_id_reference_errors_name_the_column_and_use_one_wording(self):
        entry = Entry.objects.create(
            project=self.project, title='差旅', kind='expense', amount='100.00', due_date='2026-09-09'
        )
        other = self.active_project()
        foreign_task = Task.objects.create(
            project=other, title='别的项目任务', kind='design', assignee=self.users['manager']
        )
        stock = Stock.objects.create(item=Item.objects.first(), location='主仓', quantity='10.000', value='100.00')
        base = {'amount': '10', 'date': '2026-09-09', 'reason': '付款'}
        cases = [
            ('payments', {**base, 'entry': '999999'}, '款项ID', 'entry'),
            ('payments', {**base, 'entry': str(entry.pk), 'document': '999999'}, '凭证ID', 'document'),
            ('payments', {**base, 'entry': str(entry.pk), 'reconciliation': '999999'}, '对账单ID', 'reconciliation'),
            (
                'moves',
                {'stock': '999999', 'quantity': '1', 'reason': '领料', 'project': self.project.code},
                '库存ID',
                'stock',
            ),
            (
                'moves',
                {
                    'stock': str(stock.pk),
                    'quantity': '1',
                    'reason': '领料',
                    'project': self.project.code,
                    'task': str(foreign_task.pk),
                },
                '任务ID',
                'task',
            ),
            ('time', {'task': '999999', 'date': '2026-09-09', 'hours': '1', 'reason': '设计'}, '任务ID', 'task'),
        ]
        wording = set()
        for resource, data, label, key in cases:
            with self.subTest(resource=resource, key=key):
                result = self.id_preview(resource, data)
                self.assertEqual(result.status_code, 200, result.data)
                error = result.data['errors'][0]
                self.assertIn(label, error['message'])
                self.assertEqual(error['fields'], [key])
                wording.add(error['message'].split('：', 1)[1])
                book = self.workbook(result.data)
                marked = [
                    cell.column
                    for row in book.active
                    for cell in row
                    if cell.comment and cell.fill.fgColor.rgb.endswith('FDE2E2')
                ]
                self.assertEqual(len(marked), 1)
                book.close()
        self.assertEqual(wording, {'记录不存在或不在你的权限范围内。'})

    def test_person_reference_errors_name_the_column_and_keep_wording(self):
        from apps.accounts.models import User

        User.objects.filter(pk=self.users['manager'].pk).update(display_name='重名')
        User.objects.filter(pk=self.users['member'].pk).update(display_name='重名')
        inactive = self.users['mechanical_engineer']
        headers = [label for label, _ in transfers.SCHEMAS['tasks'][1]]
        keys = [key for _, key in transfers.SCHEMAS['tasks'][1]]
        ghost_rows = {'不存在的人': '未找到人员', '重名': '对应多名人员'}
        User.objects.filter(pk=inactive.pk).update(is_active=False)
        ghost_rows[inactive.username] = '已停用'
        for value, text in ghost_rows.items():
            with self.subTest(value=value):
                row = {'title': '任务', 'kind': '设计', 'assignee': value, 'project': self.project.code}
                result = self.upload('tasks/import-file/', headers, [[row.get(key, '') for key in keys]])
                self.assertEqual(result.status_code, 200, result.data)
                error = result.data['errors'][0]
                self.assertIn(text, error['message'])
                self.assertNotIn('执行人姓名或账号：', error['message'])
                self.assertEqual(error['fields'], ['assignee'])
                book = self.workbook(result.data)
                marked = [
                    cell.column
                    for row_cells in book.active
                    for cell in row_cells
                    if cell.comment and cell.fill.fgColor.rgb.endswith('FDE2E2')
                ]
                self.assertEqual(marked, [keys.index('assignee') + 1])
                book.close()

    def test_confirmation_failure_returns_located_rows_and_writes_nothing(self):
        stock = Stock.objects.create(item=Item.objects.first(), location='主仓', quantity='10.000', value='100.00')
        BOMLine.objects.create(project=self.project, item=stock.item, quantity='5.000')
        preview = self.id_preview(
            'moves', {'stock': str(stock.pk), 'quantity': '5', 'reason': '领料', 'project': self.project.code}
        )
        self.assertTrue(preview.data['can_import'], preview.data)
        stock.quantity = '1.000'
        stock.save()
        moves, receipts, audits = StockMove.objects.count(), ActionReceipt.objects.count(), AuditLog.objects.count()
        response = self.clients['admin'].post(
            '/api/business/moves/import-confirm/',
            {'token': preview.data['token']},
            format='json',
            HTTP_IDEMPOTENCY_KEY='confirm-fail',
        )
        self.assertEqual(response.status_code, 400, response.data)
        row = response.data['rows'][0]
        self.assertEqual(int(row['row']), 2)
        self.assertEqual(list(row['fields']), ['stock'])
        self.assertEqual(
            (StockMove.objects.count(), ActionReceipt.objects.count(), AuditLog.objects.count()),
            (moves, receipts, audits),
        )
        stock.refresh_from_db()
        self.assertEqual(str(stock.quantity), '1.000')

    def test_purchase_and_partner_guide_examples_are_valid_as_whole_rows(self):
        before = ActionReceipt.objects.count()
        for resource, overrides in [
            ('partners', {}),
            (
                'purchases',
                {
                    'project': self.project.code,
                    'supplier': self.supplier.code,
                    'item': self.item.code,
                },
            ),
        ]:
            with self.subTest(resource=resource):
                layout = transfers.layout(self.users['admin'], resource)
                row = {column['key']: column['example'] for column in layout['columns']}
                row.update(overrides)
                self.assertEqual(row['payment_days'], '')
                if resource == 'purchases':
                    self.assertEqual(row['payment_due_date'], '')
                result = self.id_preview(resource, row)
                self.assertTrue(result.data['can_import'], result.data)
        self.assertEqual(ActionReceipt.objects.count(), before)

    def test_expense_guidance_explains_new_expense_without_changing_compatible_headers(self):
        layout = transfers.layout(self.users['admin'], 'entries')
        self.assertEqual(
            [column['label'] for column in layout['columns']], ['项目编号', '款项', '原金额', '最近待付期限']
        )
        self.assertIn('费用名称', layout['quick_start']['minimum'])
        self.assertIn('本笔付款到期日', layout['quick_start']['minimum'])
        amount = next(column for column in layout['columns'] if column['key'] == 'amount')
        self.assertIn('本笔费用金额', amount['hint'])
        for columns in transfers.accepted_schemas('entries'):
            row = {'project': self.project.code, 'title': '差旅费', 'amount': '10', 'due_date': '2026-09-09'}
            preview = self.upload(
                'entries/import-file/', [label for label, _ in columns], [[row.get(key, '') for _, key in columns]]
            )
            self.assertTrue(preview.data['can_import'], preview.data)

    def test_out_of_scope_task_id_is_located_without_names_or_duplicate_history(self):
        other = self.active_project()
        other.manager = self.users['admin']
        other.members.clear()
        other.save()
        task = Task.objects.create(project=other, title='不可见任务名称', kind='design', assignee=self.users['admin'])
        columns = transfers.SCHEMAS['time'][1]
        row = {'task': str(task.pk), 'date': '2026-09-09', 'hours': '1', 'reason': '测试'}
        result = self.upload(
            'time/import-file/',
            [label for label, _ in columns],
            [[row.get(key, '') for _, key in columns]],
            role='member',
        )
        self.assertFalse(result.data['can_import'])
        self.assertEqual(result.data['errors'][0]['fields'], ['task'])
        self.assertIn(transfers.UNREACHABLE, result.data['errors'][0]['message'])
        self.assertNotIn(task.title, str(result.data))
        self.assertIsNone(result.data['duplicate_import'])

    def test_production_manager_template_and_chinese_task_import_are_scoped(self):
        person = self.users['production_manager']
        self.project.members.add(person)
        layout = self.clients['production_manager'].get('/api/business/tasks/import-schema/')
        self.assertEqual(layout.status_code, 200, layout.data)
        kind = next(column for column in layout.data['columns'] if column['key'] == 'kind')
        self.assertEqual(kind['options'], ['装配', '调试'])
        self.assertEqual(kind['example'], '装配')
        self.assertIn('本人参与项目', layout.data['note'])
        template = self.clients['production_manager'].get('/api/business/tasks/import-template/?file_format=xlsx')
        book = load_workbook(io.BytesIO(template.content))
        stage = next(
            control for control in book['数据'].data_validations.dataValidation if str(control.sqref).startswith('B2:')
        )
        self.assertEqual(stage.formula1, '"装配,调试"')
        book.close()
        for stage in ('装配', '调试'):
            preview = self.id_preview(
                'tasks',
                {
                    'project': self.project.code,
                    'kind': stage,
                    'title': stage + '任务',
                    'assignee': 'production_manager',
                },
                role='production_manager',
            )
            self.assertTrue(preview.data['can_import'], preview.data)
            self.post('production_manager', 'tasks/import-confirm/', {'token': preview.data['token']})
        self.assertEqual(set(Task.objects.values_list('kind', flat=True)), {'assembly', 'test'})

    def test_production_manager_mixed_design_or_foreign_project_rolls_back_entire_preview(self):
        person = self.users['production_manager']
        person.additional_roles = ['purchaser']
        person.save()
        self.project.members.add(person)
        other = self.active_project()
        columns = transfers.SCHEMAS['tasks'][1]
        good = {'project': self.project.code, 'kind': '装配', 'title': '本项目装配', 'assignee': 'production_manager'}
        before = ActionReceipt.objects.count(), AuditLog.objects.count()
        for bad in [{**good, 'kind': '设计'}, {**good, 'project': other.code}]:
            result = self.upload(
                'tasks/import-file/',
                [label for label, _ in columns],
                [[row.get(key, '') for _, key in columns] for row in (good, bad)],
                role='production_manager',
            )
            self.assertFalse(result.data['can_import'], result.data)
            self.assertIsNone(result.data['token'])
            self.assertEqual(result.data['errors'][0]['row'], 3)
            self.assertFalse(Task.objects.exists())
            self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), before)
        for role in ('mechanical_engineer', 'electrical_engineer', 'member', 'purchaser'):
            self.assertEqual(self.clients[role].get('/api/business/tasks/import-schema/').status_code, 403)

    def test_production_import_confirmation_and_replay_recheck_membership_and_stage_role(self):
        person = self.users['production_manager']
        self.project.members.add(person)
        row = {'project': self.project.code, 'kind': 'assembly', 'title': '装配', 'assignee': 'manager'}
        pending = self.id_preview('tasks', row, role='production_manager')
        self.assertTrue(pending.data['can_import'], pending.data)
        self.project.members.remove(person)
        self.post('production_manager', 'tasks/import-confirm/', {'token': pending.data['token']}, status=403)
        self.assertFalse(Task.objects.exists())
        self.project.members.add(person)
        result = self.post(
            'production_manager', 'tasks/import-confirm/', {'token': pending.data['token']}, key='production-confirm'
        )
        person.role = 'purchaser'
        person.additional_roles = ['member']
        person.save()
        self.post(
            'production_manager',
            'tasks/import-confirm/',
            {'token': pending.data['token']},
            key='production-confirm',
            status=403,
        )
        self.assertEqual(Task.objects.count(), 1)
        self.assertEqual(result['resource'], 'tasks')
        # A formerly valid design token cannot be replayed after switching to production-only authority.
        design = self.id_preview('tasks', {**row, 'kind': 'design', 'title': '设计'}, role='manager')
        self.post('manager', 'tasks/import-confirm/', {'token': design.data['token']}, key='design-confirm')
        assembly = self.id_preview('tasks', {**row, 'title': '经理导入的装配任务'}, role='manager')
        saved = self.post('manager', 'tasks/import-confirm/', {'token': assembly.data['token']}, key='assembly-confirm')
        self.users['manager'].role = 'production_manager'
        self.users['manager'].additional_roles = ['purchaser']
        self.users['manager'].save()
        self.post('manager', 'tasks/import-confirm/', {'token': design.data['token']}, key='design-confirm', status=403)
        self.post('manager', 'tasks/import-confirm/', {'token': design.data['token']}, key='new-design-key', status=403)
        for key in ('assembly-confirm', 'new-assembly-key'):
            self.assertEqual(
                self.post('manager', 'tasks/import-confirm/', {'token': assembly.data['token']}, key=key), saved
            )
        self.assertEqual(Task.objects.count(), 3)

    ITEM_HEADERS = [
        '物料编码',
        '物料名称',
        '规格',
        '图号',
        '图档版本',
        '产品编码类别',
        '品牌',
        '物料类别',
        '单位',
        '独立建码原因',
    ]

    def test_item_import_reports_every_independent_field_error_on_one_row(self):
        counts = (Item.objects.count(), CodeRule.objects.get(key='item').product_counters)
        rows = [
            ['', '机架', '', '', '', '有图·机加', '', '', '件', ''],  # 规格、图号都缺
            ['', '单错件', '', '', '', '无图·标准件', '', '', '件', ''],  # 只缺规格：单错行不变
            ['', '', '', '', '', '有图·机加', '', '', '', ''],  # 名称、规格、图号都缺
        ]
        result = self.upload('items/import-file/', self.ITEM_HEADERS, rows)
        self.assertEqual(result.status_code, 200, result.data)
        self.assertFalse(result.data['can_import'])
        by_row = {error['row']: error for error in result.data['errors']}
        self.assertEqual(set(by_row[2]['fields']), {'specification', 'drawing_number'})
        self.assertIn('规格', by_row[2]['message'])
        self.assertIn('图号', by_row[2]['message'])
        self.assertEqual(by_row[3]['fields'], ['specification'])
        self.assertNotIn('图号', by_row[3]['message'])
        self.assertEqual(set(by_row[4]['fields']), {'name', 'specification', 'drawing_number'})
        sheet = self.workbook(result.data).active
        headers = [cell.value for cell in sheet[1]]
        for row, fields in ((2, {'规格', '图号'}), (3, {'规格'}), (4, {'物料名称', '规格', '图号'})):
            marked = {headers[index] for index, cell in enumerate(sheet[row]) if cell.comment}
            self.assertEqual(marked, fields, row)
        self.assertEqual(counts, (Item.objects.count(), CodeRule.objects.get(key='item').product_counters))

    def test_item_api_returns_all_field_errors_as_dict_and_keeps_single_field_shape(self):
        client = self.clients['purchaser']
        response = client.post(
            '/api/business/items/',
            {'name': '机架', 'product_category': '11', 'part_type': 'bogus'},
            format='json',
            HTTP_IDEMPOTENCY_KEY='multi-field-1',
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.data), {'part_type', 'specification', 'drawing_number'})
        single = client.post(
            '/api/business/items/',
            {'name': '机架', 'specification': 'x', 'product_category': '11'},
            format='json',
            HTTP_IDEMPOTENCY_KEY='multi-field-2',
        )
        self.assertEqual(single.status_code, 400)
        self.assertEqual(set(single.data), {'drawing_number'})
        invalid = client.post(
            '/api/business/items/',
            {'name': '机架', 'product_category': '99'},
            format='json',
            HTTP_IDEMPOTENCY_KEY='multi-field-3',
        )
        self.assertEqual(set(invalid.data), {'product_category'})  # 类别无效时规格/图号依赖类别，不编造

    def test_bom_material_import_reports_every_new_material_field_error(self):
        from apps.business.services import bom_material_import

        row = ['', '上料', '1', '', '', '', '', '机架', '', '', '', '有图·机加', '', '', '']
        stream = io.StringIO()
        csv.writer(stream).writerows([bom_material_import.HEADERS, row])
        response = self.clients['manager'].post(
            f'/api/business/projects/{self.project.pk}/import-preview/',
            {'file': SimpleUploadedFile('bom.csv', stream.getvalue().encode())},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['can_import'])
        message = response.data['errors'][0]['message']
        for label in ('规格', '图号', '单位'):
            self.assertIn(label, message)


class CorruptSpreadsheetTests(BusinessFixtures, TestCase):
    """解包失败（deflate 数据损坏等）的 XLSX 在各导入入口返回中文 400，不再 500。"""

    def setUp(self):
        self.setup_business()
        self.project = self.active_project()

    def corrupt_xlsx(self, headers):
        import zipfile

        book = Workbook()
        book.active.append(headers)
        for index in range(50):
            book.active.append([f'I{index}'] + [''] * (len(headers) - 1))
        stream = io.BytesIO()
        book.save(stream)
        data = bytearray(stream.getvalue())
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
            offset = archive.getinfo('xl/worksheets/sheet1.xml').header_offset
        name_length = int.from_bytes(data[offset + 26 : offset + 28], 'little')
        extra_length = int.from_bytes(data[offset + 28 : offset + 30], 'little')
        start = offset + 30 + name_length + extra_length
        data[start : start + 8] = b'\xff' * 8  # 非法 deflate 块类型，解压时抛 zlib.error
        return bytes(data)

    def test_corrupt_xlsx_returns_chinese_400_for_generic_and_bom_import(self):
        from apps.business.services import bom_material_import

        cases = [
            ('admin', '/api/business/projects/import-file/', [label for label, _ in transfers.SCHEMAS['projects'][1]]),
            ('admin', '/api/business/items/import-file/', [label for label, _ in transfers.SCHEMAS['items'][1]]),
            ('manager', f'/api/business/projects/{self.project.pk}/import-preview/', bom_material_import.HEADERS),
        ]
        for role, url, headers in cases:
            with self.subTest(url=url):
                client = self.clients[role]
                client.raise_request_exception = False
                response = client.post(
                    url, {'file': SimpleUploadedFile('broken.xlsx', self.corrupt_xlsx(headers))}, format='multipart'
                )
                self.assertEqual(response.status_code, 400, getattr(response, 'data', response.content[:200]))
                self.assertIn('表格损坏、加密或格式不支持', str(response.data))
        self.assertFalse(BOMLine.objects.filter(project=self.project).exists())
