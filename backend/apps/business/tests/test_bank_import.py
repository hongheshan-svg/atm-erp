import hashlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import date

from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from django.test import TestCase, TransactionTestCase
from openpyxl import Workbook, load_workbook

from apps.business.models import BankMatch, BankRecord, Payment
from apps.business.services import bank_import, finance
from apps.core.models import ActionReceipt, AuditLog, Company

from .test_commercial_chain import BusinessFixtures


def file(rows=None, name='bank.xlsx', missing=False, income='100.00'):
    book = Workbook()
    sheet = book.active
    rows = rows or [
        ['[HISTORYDETAIL]'],
        bank_import.ICBC,
        [
            '000000000',
            '001234567890',
            '2026-01-01 12:00:00',
            '贷',
            '' if missing else '客户',
            '',
            '货款',
            '',
            '',
            '业务发生账号:123456789012',
            income,
            '',
            '',
            '1000.00',
        ],
        [
            '000000000',
            '009876543210',
            '2026-01-02 12:00:00',
            '借',
            '员工',
            '',
            '备用金',
            '',
            '',
            '',
            '',
            '20.00',
            '',
            '980.00',
        ],
    ]
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile(name, stream.getvalue())


def reexport(stamp='2026-1-1 12:00:00', *, duplicate=False, missing=False):
    book = load_workbook(file(missing=missing))
    sheet = book.active
    if duplicate:
        sheet.append([cell.value for cell in sheet[3]])
        sheet['C5'] = stamp
    else:
        sheet['C3'] = stamp
        sheet['K3'] = '100.000'
        sheet['N3'] = '1,000.0'
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return SimpleUploadedFile('重新导出.xlsx', stream.getvalue())


def huaxia(core='000123'):
    rows = [
        ['账号：', '123456789012'],
        ['币种：', '人民币'],
        [],
        [],
        [],
        ['收入总金额：', '100', '收入总笔数：', '1'],
        ['支出总金额：', '0', '支出总笔数：', '0'],
        bank_import.HXB,
        ['1', '2026-01-01', '12:00:00', '', '100', '1000', '001234567890', '客户', '银行', core, '货款'],
    ]
    return file(rows, name='华夏.xlsx')


def legacy_fingerprint(account, value, source):
    identity = [
        source['bank'],
        account,
        source['timestamp'],
        str(value),
        source['balance'],
        source['counterparty_account'],
        source['original_reference'],
    ]
    return hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode()).hexdigest()


class BankImportTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()

    def preview(self, upload=None, role='finance', status=200):
        result = self.clients[role].post(
            '/api/business/bank-records/import-file/', {'file': upload or file()}, format='multipart'
        )
        self.assertEqual(result.status_code, status, result.data)
        return result.data

    def confirm(self, token, role='finance', status=200):
        return self.post(role, 'bank-records/import-confirm/', {'token': token}, status=status)

    def test_native_preview_commit_repeat_and_reexport(self):
        preview = self.preview()
        self.assertEqual(BankRecord.objects.count(), 0)
        self.assertEqual(preview['summary']['income'], '100.00')
        self.assertEqual(preview['summary']['expense'], '20.00')
        self.assertTrue(preview['can_import'])
        self.assertEqual(
            [column['label'] for column in preview['columns']],
            ['银行日期', '账户', '流水号 / 导入标识', '对方户名', '收入 / 支出', '说明'],
        )
        self.assertEqual(preview['row_count'], 2)
        self.assertEqual(preview['rows'][1]['data']['amount'], '-20.00')
        self.assertEqual(self.confirm(preview['token']), {'count': 2, 'skipped': 0})
        self.assertEqual(self.confirm(preview['token']), {'count': 2, 'skipped': 0})
        # New filename, preview and user still deduplicate by transaction, not batch nonce.
        another = self.preview(file(name='renamed.xlsx'), role='admin')
        self.assertEqual(self.confirm(another['token'], role='admin'), {'count': 0, 'skipped': 2})
        self.assertEqual(BankRecord.objects.count(), 2)
        self.assertEqual(Payment.objects.count(), 0)
        self.assertEqual(BankRecord.objects.first().source['original_reference'], '000000000')
        self.assertEqual(len(set(BankRecord.objects.values_list('reference', flat=True))), 2)

    def test_equivalent_timestamp_and_money_reexport_cannot_create_second_payment(self):
        self.confirm(self.preview()['token'])
        repeated = self.preview(reexport())
        self.assertTrue(repeated['can_import'], repeated['errors'])
        self.assertTrue(all(row['status'] == '已导入（跳过）' for row in repeated['rows']))
        self.assertEqual(self.confirm(repeated['token']), {'count': 0, 'skipped': 2})
        self.assertEqual(BankRecord.objects.count(), 2)
        bank = BankRecord.objects.get(amount='100')
        entry = self.active_project().entries.get(kind='receivable')
        balance = finance.balance(entry)
        payload = {'entry': entry.pk, 'amount': '100', 'reason': '认领原始银行收入'}
        self.post('finance', f'bank-records/{bank.pk}/allocate/', payload)
        counts = (
            Payment.objects.count(),
            BankMatch.objects.count(),
            ActionReceipt.objects.count(),
            AuditLog.objects.count(),
        )
        self.post('finance', f'bank-records/{bank.pk}/allocate/', payload, status=409)
        self.assertEqual(Payment.objects.count(), 1)
        self.assertEqual(BankMatch.objects.count(), 1)
        self.assertEqual(finance.balance(entry), balance - 100)
        self.assertEqual(
            (
                Payment.objects.count(),
                BankMatch.objects.count(),
                ActionReceipt.objects.count(),
                AuditLog.objects.count(),
            ),
            counts,
        )

    def test_historical_raw_timestamp_fingerprint_is_recognized_without_changing_locked_evidence(self):
        self.confirm(self.preview(reexport(missing=True))['token'])
        bank = BankRecord.objects.get(amount='100')
        self.post('finance', f'bank-records/{bank.pk}/review/', {'counterparty': '已核实客户', 'reason': '核实旧流水'})
        entry = self.active_project().entries.get(kind='receivable')
        self.post(
            'finance',
            f'bank-records/{bank.pk}/allocate/',
            {'entry': entry.pk, 'amount': '100', 'reason': '原记录已认领'},
        )
        source = bank.source
        legacy = legacy_fingerprint(bank.account, bank.amount, source)
        legacy_reference = f'IMPORT-ICBC-{legacy[:40]}'
        # Represent a pre-upgrade row whose fingerprint included the original date formatting.
        BankRecord.objects.filter(pk=bank.pk).update(import_fingerprint=legacy, reference=legacy_reference)
        self.lock_history()
        preview = self.preview(file(missing=True))
        self.assertTrue(preview['can_import'], preview['errors'])
        self.assertEqual(self.confirm(preview['token']), {'count': 0, 'skipped': 2})
        self.assertEqual(BankRecord.objects.count(), 2)
        bank.refresh_from_db()
        self.assertEqual((bank.import_fingerprint, bank.reference, bank.source), (legacy, legacy_reference, source))
        self.assertEqual(bank.counterparty, '已核实客户')
        self.assertFalse(bank.needs_review)
        self.assertEqual(Payment.objects.count(), 1)
        self.assertEqual(BankMatch.objects.get().bank_id, bank.pk)

    def test_equivalent_timestamp_rows_in_one_file_reject_whole_batch(self):
        preview = self.preview(reexport(duplicate=True))
        self.assertFalse(preview['can_import'])
        self.assertIsNone(preview['token'])
        self.assertIn('文件内有完全重复的交易', str(preview['errors']))
        self.assertFalse(BankRecord.objects.exists())

    def test_missing_historical_evidence_and_ambiguous_matches_block_import(self):
        self.confirm(self.preview()['token'])
        bank = BankRecord.objects.get(amount='100')
        source = bank.source
        BankRecord.objects.filter(pk=bank.pk).update(source={})
        preview = self.preview(reexport())
        self.assertFalse(preview['can_import'])
        self.assertIsNone(preview['token'])
        self.assertIn('原始证据不足', str(preview['errors']))
        BankRecord.objects.filter(pk=bank.pk).update(source=source)
        source = {**source, 'timestamp': '2026-1-1 12:00:00'}
        old = legacy_fingerprint(bank.account, bank.amount, source)
        BankRecord.objects.create(
            account=bank.account,
            date=bank.date,
            amount=bank.amount,
            counterparty=bank.counterparty,
            reference=f'IMPORT-ICBC-{old[:40]}',
            import_fingerprint=old,
            source=source,
            reason='模拟升级前重复记录',
        )
        preview = self.preview(reexport())
        self.assertFalse(preview['can_import'])
        self.assertIsNone(preview['token'])
        self.assertIn('存在多条相同交易', str(preview['errors']))
        self.assertEqual(BankRecord.objects.count(), 3)
        self.assertFalse(Payment.objects.exists())

    def test_old_unfinished_token_normalizes_inside_execute_and_completed_payload_stays_stable(self):
        def old_token():
            batch = signing.loads(self.preview(reexport())['token'], salt='bank-import')
            for row in batch['records']:
                old = legacy_fingerprint(row['data']['account'], row['data']['amount'], row['source'])
                row['fingerprint'] = old
                row['data']['reference'] = f'IMPORT-ICBC-{old[:40]}'
            return signing.dumps(batch, salt='bank-import', compress=True), batch

        token, batch = old_token()
        pending, _ = old_token()
        self.assertEqual(self.confirm(token), {'count': 2, 'skipped': 0})
        canonical = self.preview()['rows'][0]['fingerprint']
        bank = BankRecord.objects.get(amount='100')
        self.assertEqual(bank.import_fingerprint, canonical)
        self.assertEqual(bank.reference, f'IMPORT-ICBC-{canonical[:40]}')
        receipt = ActionReceipt.objects.get(key='bank-import-batch:' + batch['nonce'])
        encoded = json.dumps(batch, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)
        self.assertEqual(receipt.fingerprint, hashlib.sha256(encoded.encode()).hexdigest())
        self.lock_history()
        self.assertEqual(self.confirm(pending), {'count': 0, 'skipped': 2})
        self.assertEqual(self.confirm(token), {'count': 2, 'skipped': 0})
        self.assertEqual(BankRecord.objects.count(), 2)

    def test_different_transaction_time_balance_account_or_original_reference_is_not_collapsed(self):
        self.confirm(self.preview()['token'])
        for cell, value in [
            ('C3', '2026-01-01 12:00:01'),
            ('N3', '1000.01'),
            ('B3', '001234567891'),
            ('A3', '000000001'),
        ]:
            with self.subTest(cell=cell):
                book = load_workbook(file())
                book.active[cell] = value
                stream = io.BytesIO()
                book.save(stream)
                book.close()
                preview = self.preview(SimpleUploadedFile('不同交易.xlsx', stream.getvalue()))
                self.assertTrue(preview['can_import'], preview['errors'])
                self.assertEqual(self.confirm(preview['token']), {'count': 1, 'skipped': 1})
        self.assertEqual(BankRecord.objects.count(), 6)

    def test_missing_name_preserved_and_review_required(self):
        preview = self.preview(file(missing=True))
        self.confirm(preview['token'])
        bank = BankRecord.objects.get(needs_review=True)
        self.assertEqual(bank.source['counterparty'], '')
        from apps.business.services.banking import check_review
        from apps.core.api import Conflict

        with self.assertRaises(Conflict):
            check_review(bank)
        self.post('finance', f'bank-records/{bank.pk}/review/', {'counterparty': '核实客户', 'reason': '银行回单核对'})
        bank.refresh_from_db()
        self.assertFalse(bank.needs_review)
        self.assertEqual(bank.source['counterparty'], '')
        self.assertTrue(AuditLog.objects.filter(operation='bank.review').exists())
        self.confirm(self.preview(file(missing=True))['token'])
        bank.refresh_from_db()
        self.assertEqual(bank.counterparty, '核实客户')
        self.assertFalse(bank.needs_review)

    def test_legacy_import_is_enriched_without_duplicate_money(self):
        preview = self.preview()
        record = preview['rows'][0]
        self.post('finance', 'bank-records/', {**record['data'], 'reference': record['legacy_reference']}, status=201)
        self.assertEqual(self.confirm(preview['token']), {'count': 1, 'skipped': 1})
        self.assertEqual(BankRecord.objects.count(), 2)
        self.assertEqual(BankRecord.objects.filter(import_fingerprint__isnull=True).count(), 0)

    def manual(self, reference, amount='100.00'):
        data = {
            'amount': amount,
            'date': '2026-01-01',
            'account': '123456789012',
            'reference': reference,
            'counterparty': '客户',
            'reason': '先手工登记到账',
        }
        return self.post('finance', 'bank-records/', data, status=201)['id']

    def test_manual_record_with_same_bank_reference_is_the_imported_transaction(self):
        bank = self.manual('000123')
        preview = self.preview(huaxia())
        self.assertTrue(preview['can_import'], preview['errors'])
        self.assertEqual(preview['rows'][0]['status'], '已导入（跳过）')
        self.assertEqual(self.confirm(preview['token']), {'count': 0, 'skipped': 1})
        record = BankRecord.objects.get()
        self.assertEqual(record.pk, bank)
        self.assertEqual(record.source['original_reference'], '000123')
        again = self.preview(huaxia())
        self.assertEqual(self.confirm(again['token']), {'count': 0, 'skipped': 1})
        self.assertEqual(BankRecord.objects.count(), 1)
        self.post('finance', f'bank-records/{bank}/void/', {'reason': '测试同号不同金额'})
        self.manual('000123', amount='90.00')
        changed = self.preview(huaxia())
        self.assertFalse(changed['can_import'])
        self.assertIn('日期或金额不同', str(changed['errors']))

    def test_manual_record_without_comparable_reference_blocks_import_until_voided(self):
        # ICBC voucher numbers are shared by rows, so the manual reference cannot identify the transfer.
        bank = self.manual('000000000')
        preview = self.preview()
        self.assertFalse(preview['can_import'])
        self.assertIsNone(preview['token'])
        self.assertEqual([error['row'] for error in preview['errors']], [3])
        self.assertIn(f'手工银行记录 ID {bank}', preview['errors'][0]['message'])
        self.assertIn('疑似同一笔', preview['errors'][0]['message'])
        self.assertIn('作废误录的手工记录', preview['errors'][0]['message'])
        self.post('finance', f'bank-records/{bank}/void/', {'reason': '与银行导入重复'})
        preview = self.preview()
        self.assertTrue(preview['can_import'], preview['errors'])
        self.assertEqual(self.confirm(preview['token']), {'count': 2, 'skipped': 0})
        self.assertEqual(BankRecord.objects.filter(void_reason='').count(), 2)

    def test_manual_suspect_added_after_preview_blocks_confirmation_without_side_effects(self):
        preview = self.preview()
        self.manual('客户转账')
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        self.assertIn('疑似同一笔', str(self.confirm(preview['token'], status=400)))
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)

    def test_huaxia_blank_manual_reference_blocks_until_a_different_reference_is_filled(self):
        legacy = BankRecord.objects.create(
            account='123456789012',
            date='2026-01-01',
            amount='100.00',
            counterparty='客户',
            reference='',
            reason='旧版无流水号手工记录',
        )
        preview = self.preview(huaxia())
        self.assertFalse(preview['can_import'])
        self.assertIn(f'手工银行记录 ID {legacy.pk}', str(preview['errors']))
        self.assertIn('补填银行流水号', str(preview['errors']))
        BankRecord.objects.filter(pk=legacy.pk).update(reference='000999')
        preview = self.preview(huaxia())
        self.assertTrue(preview['can_import'], preview['errors'])
        self.assertEqual(self.confirm(preview['token']), {'count': 1, 'skipped': 0})
        self.assertEqual(BankRecord.objects.count(), 2)

    def test_permissions_token_binding_and_tamper(self):
        self.preview(role='member', status=403)
        preview = self.preview()
        self.confirm(preview['token'], role='admin', status=400)
        self.confirm(preview['token'] + 'x', status=400)
        self.confirm(preview['token'], role='member', status=403)
        self.assertEqual(BankRecord.objects.count(), 0)

    def test_bad_money_formula_and_file_rejected(self):
        result = self.preview(file(income='100.001'))
        self.assertFalse(result['can_import'])
        self.assertIn('最多两位小数', result['errors'][0]['message'])
        self.assertNotIn('ErrorDetail', result['errors'][0]['message'])
        self.preview(file(income='=10+20'), status=400)
        self.preview(SimpleUploadedFile('bad.xls', b'not excel'), status=400)
        self.assertEqual(BankRecord.objects.count(), 0)

    def test_invalid_date_and_long_fields_have_chinese_row_errors(self):
        for cell, value, expected in [('C3', 'bad-date', '交易日期或时间格式不正确'), ('E3', '户名' * 76, '对方户名')]:
            with self.subTest(cell=cell):
                book = load_workbook(file())
                book.active[cell] = value
                stream = io.BytesIO()
                book.save(stream)
                book.close()
                result = self.preview(SimpleUploadedFile('bank.xlsx', stream.getvalue()))
                self.assertFalse(result['can_import'])
                self.assertIsNone(result['token'])
                self.assertEqual(result['errors'][0]['row'], 3)
                self.assertIn(expected, result['errors'][0]['message'])
                self.assertNotIn('ErrorDetail', result['errors'][0]['message'])
                self.assertNotIn('counterparty', result['errors'][0]['message'])
                self.assertFalse(BankRecord.objects.exists())

    def test_huaxia_header_mismatch_blocks_import(self):
        rows = [
            ['账号：', '123456789012'],
            ['币种：', '人民币'],
            [],
            [],
            [],
            ['收入总金额：', '101', '收入总笔数：', '1'],
            ['支出总金额：', '0', '支出总笔数：', '0'],
            bank_import.HXB,
            ['1', '2026-01-01', '12:00:00', '', '100', '100', '001234567890', '客户', '银行', '000123', '货款'],
        ]
        preview = self.preview(file(rows))
        self.assertFalse(preview['can_import'])
        self.assertIsNone(preview['token'])
        rows[5][1] = '100'
        valid = self.preview(file(rows))
        self.assertTrue(valid['can_import'])
        self.assertEqual(valid['rows'][0]['source']['original_reference'], '000123')

    def test_huaxia_equivalent_timestamp_and_money_reexport_is_skipped(self):
        rows = [
            ['账号：', '123456789012'],
            ['币种：', '人民币'],
            [],
            [],
            [],
            ['收入总金额：', '100', '收入总笔数：', '1'],
            ['支出总金额：', '0', '支出总笔数：', '0'],
            bank_import.HXB,
            ['1', '2026-01-01', '12:00:00', '', '100', '1000', '001234567890', '客户', '银行', '000123', '货款'],
        ]
        self.assertEqual(self.confirm(self.preview(file(rows))['token']), {'count': 1, 'skipped': 0})
        rows[8][1], rows[8][2], rows[8][4], rows[8][5] = '2026-1-1', '12:0:0', '100.000', '1,000.0'
        preview = self.preview(file(rows, name='华夏重导.xlsx'))
        self.assertTrue(preview['can_import'], preview['errors'])
        self.assertEqual(self.confirm(preview['token']), {'count': 0, 'skipped': 1})
        self.assertEqual(BankRecord.objects.count(), 1)

    def test_conflict_at_confirmation_rolls_back_whole_batch(self):
        preview = self.preview()
        last = preview['rows'][1]
        self.post('finance', 'bank-records/', {**last['data'], 'amount': '-99'}, status=201)
        self.confirm(preview['token'], status=400)
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual(Payment.objects.count(), 0)

    def test_matching_pending_name_rejected_before_payment(self):
        project = self.active_project()
        entry = project.entries.get(kind='receivable')
        self.confirm(self.preview(file(missing=True))['token'])
        bank = BankRecord.objects.get(needs_review=True)
        self.post(
            'finance',
            f'bank-records/{bank.pk}/allocate/',
            {'entry': entry.pk, 'amount': '100', 'reason': '试图认领'},
            status=409,
        )
        self.assertEqual(Payment.objects.count(), 0)

    def lock_history(self):
        company, _ = Company.objects.get_or_create(pk=1)
        company.locked_through = date(2026, 1, 31)
        company.save()

    def test_locked_period_import_creates_raw_evidence_once_without_payment(self):
        self.lock_history()
        preview = self.preview()
        self.assertTrue(preview['can_import'])
        self.assertFalse(BankRecord.objects.exists())
        self.assertEqual(self.confirm(preview['token']), {'count': 2, 'skipped': 0})
        self.assertEqual(self.confirm(preview['token']), {'count': 2, 'skipped': 0})
        self.assertEqual(
            self.confirm(self.preview(file(name='历史流水副本.xlsx'))['token']), {'count': 0, 'skipped': 2}
        )
        self.assertEqual(BankRecord.objects.count(), 2)
        self.assertFalse(Payment.objects.exists())
        self.assertFalse(BankMatch.objects.exists())
        for bank in BankRecord.objects.all():
            self.assertTrue(bank.import_fingerprint)
            self.assertEqual(bank.source['original_reference'], '000000000')
            self.assertIn('timestamp', bank.source)
        self.assertEqual(AuditLog.objects.filter(operation='bank.create').count(), 2)
        self.assertEqual(AuditLog.objects.filter(operation='bank.import-evidence').count(), 2)

    def test_locked_raw_bank_record_can_be_registered_but_cannot_post_or_be_changed(self):
        project = self.active_project()
        entry = project.entries.get(kind='receivable')
        data = self.preview()['rows'][0]['data']
        self.lock_history()
        bank = self.post('finance', 'bank-records/', data, status=201)['id']
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        for path, body in [
            (f'bank-records/{bank}/allocate/', {'entry': entry.pk, 'amount': '100', 'reason': '禁止锁期认领入账'}),
            (f'entries/{entry.pk}/pay/', {'amount': '100', 'date': '2026-01-01', 'reason': '禁止锁期补录收款'}),
            (f'bank-records/{bank}/void/', {'reason': '禁止修改锁期原始记录'}),
        ]:
            with self.subTest(path=path):
                result = self.post('finance', path, body, status=409)
                self.assertIn('锁账', str(result))
                self.assertFalse(Payment.objects.exists())
                self.assertFalse(BankMatch.objects.exists())
                self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)
        self.assertEqual(BankRecord.objects.get(pk=bank).void_reason, '')

    def test_locked_import_retains_missing_name_and_still_blocks_review_edit(self):
        self.lock_history()
        self.confirm(self.preview(file(missing=True))['token'])
        bank = BankRecord.objects.get(needs_review=True)
        self.assertEqual(bank.source['counterparty'], '')
        self.post(
            'finance', f'bank-records/{bank.pk}/review/', {'counterparty': '客户', 'reason': '历史核实'}, status=409
        )
        bank.refresh_from_db()
        self.assertTrue(bank.needs_review)
        self.assertEqual(bank.source['counterparty'], '')

    def test_locked_legacy_record_enrichment_rejects_and_rolls_back_new_rows(self):
        preview = self.preview()
        last = preview['rows'][1]
        bank = self.post(
            'finance', 'bank-records/', {**last['data'], 'reference': last['legacy_reference']}, status=201
        )['id']
        self.lock_history()
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        self.confirm(preview['token'], status=409)
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertIsNone(BankRecord.objects.get(pk=bank).import_fingerprint)
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)


class BankImportConcurrencyTests(BusinessFixtures, TransactionTestCase):
    def setUp(self):
        self.setup_business()

    def test_two_users_import_same_transactions_concurrently(self):
        self.concurrent_imports([file(), file()])

    def test_two_users_import_equivalent_timestamp_formats_concurrently(self):
        self.concurrent_imports([file(), reexport()])

    def test_concurrent_confirmations_recheck_manual_bank_records(self):
        from rest_framework.exceptions import ValidationError

        tokens = {
            name: [
                (self.users[role], bank_import.preview(self.users[role], build())['token'])
                for role in ['admin', 'finance']
            ]
            for name, build in [('icbc', file), ('huaxia', huaxia)]
        }
        manual = {
            'amount': '100.00',
            'date': '2026-01-01',
            'account': '123456789012',
            'reference': '000123',
            'counterparty': '客户',
            'reason': '先手工登记',
        }
        self.post('finance', 'bank-records/', manual, status=201)
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())

        def run(job):
            close_old_connections()
            try:
                return bank_import.confirm(*job)
            except ValidationError as exc:
                return str(exc.detail)
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(run, tokens['icbc']))
        self.assertTrue(all('疑似同一笔' in result for result in results), results)
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(run, tokens['huaxia']))
        self.assertEqual(sorted(result['skipped'] for result in results), [1, 1], results)
        self.assertEqual(sum(result['count'] for result in results), 0)
        self.assertEqual(BankRecord.objects.get().source['original_reference'], '000123')

    def concurrent_imports(self, uploads):
        jobs = [
            (self.users[role], bank_import.preview(self.users[role], upload)['token'])
            for role, upload in zip(['admin', 'finance'], uploads, strict=True)
        ]

        def run(job):
            close_old_connections()
            try:
                return bank_import.confirm(*job)
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(run, jobs))
        self.assertEqual(sum(result['count'] for result in results), 2)
        self.assertEqual(sum(result['skipped'] for result in results), 2)
        self.assertEqual(BankRecord.objects.count(), 2)
