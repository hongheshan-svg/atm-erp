import csv
import io
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from decimal import Decimal
from threading import Barrier
from unittest.mock import patch
from urllib.parse import quote

from django.db import IntegrityError, close_old_connections, connection, transaction
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from apps.business.models import (
    BankMatch,
    BankRecord,
    Delivery,
    Document,
    Entry,
    Partner,
    Payment,
    Project,
    PurchaseLine,
    PurchaseOrder,
    PurchaseWarranty,
    Reconciliation,
    SalesOrder,
    Stock,
    StockMove,
)
from apps.business.services import banking, finance, reconciliation
from apps.core.api import Conflict
from apps.core.models import ActionReceipt, AuditLog, CodeRule, Company

from .test_commercial_chain import TODAY, BusinessFixtures


class ReconciliationTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()
        self.project = self.active_project()

    def request(self, role, path, data, status=200, key=None):
        response = self.clients[role].post(
            '/api/business/' + path, data, format='json', HTTP_IDEMPOTENCY_KEY=key or str(uuid.uuid4())
        )
        self.assertEqual(response.status_code, status, response.data)
        return response.data

    def purchase_entry(self, receive=True):
        purchase = self.purchase(self.project)
        if receive:
            self.post(
                'warehouse',
                f'purchases/{purchase.pk}/receive/',
                {'lines': [{'line': purchase.lines.get().pk, 'quantity': '5'}], 'reason': '收货核对'},
            )
        return purchase.entry

    def statement(self, entry, amount='500', kind='settlement', counterparty=None, confirm=True):
        result = self.request(
            'finance',
            'reconciliations/',
            {
                'entry': entry.pk,
                'kind': kind,
                'approved_amount': amount,
                'counterparty_balance': counterparty if counterparty is not None else str(finance.balance(entry)),
                'basis': 'HT-001 合同约定预付额度',
                'reason': '双方核对',
            },
            201,
        )
        if confirm:
            self.request(
                'manager' if kind == 'prepayment' else 'finance',
                f'reconciliations/{result["id"]}/confirm/',
                {'reason': '逐项核准'},
            )
        return result['id']

    def pay(self, entry, amount='500', statement=None, refund=False, status=200, key=None):
        return self.request(
            'finance',
            f'entries/{entry.pk}/{"refund" if refund else "pay"}/',
            {
                'amount': amount,
                'date': TODAY,
                'reason': '实际收付',
                **({'reconciliation': statement} if statement else {}),
            },
            status,
            key,
        )

    def bank(self, amount='500', project=True):
        return self.request(
            'finance',
            'bank-records/',
            {
                'amount': amount,
                'date': TODAY,
                'account': 'BANK-A',
                'reference': str(uuid.uuid4()),
                'counterparty': '客户银行户名',
                'reason': '银行实际记录',
                **({'project': self.project.pk} if project else {}),
            },
            201,
        )['id']

    def test_gate_difference_partial_payments_and_replay(self):
        entry = self.purchase_entry()
        self.pay(entry, status=409)
        bad = self.statement(entry, counterparty='501', confirm=False)
        self.request('finance', f'reconciliations/{bad}/confirm/', {'reason': '有差异'}, 409)
        self.request('finance', f'reconciliations/{bad}/void/', {'reason': '更正对方账单'})
        statement = self.statement(entry)
        key = str(uuid.uuid4())
        first = self.pay(entry, '300', statement, key=key)
        self.assertEqual(self.pay(entry, '300', statement, key=key), first)
        self.pay(entry, '250', statement, status=409)
        self.pay(entry, '200', statement)
        self.assertEqual(entry.payments.count(), 2)
        detail = self.clients['finance'].get(f'/api/business/reconciliations/{statement}/').data
        self.assertTrue(detail['valid'])
        self.assertEqual(detail['remaining_amount'], '0.00')
        self.assertEqual(finance.balance(entry), 0)

    def test_invalid_kind_shape_is_rejected_without_statement_or_receipt(self):
        entry = self.project.entries.get(kind='receivable')
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        for kind in (['settlement'], {'value': 'settlement'}):
            with self.subTest(kind=kind):
                response = self.request(
                    'finance',
                    'reconciliations/',
                    {
                        'entry': entry.pk,
                        'kind': kind,
                        'counterparty_balance': '10000',
                        'approved_amount': '10000',
                        'reason': '异常类型校验',
                    },
                    400,
                )
                self.assertIn('kind', response)
                self.assertFalse(Reconciliation.objects.exists())
                self.assertEqual(ActionReceipt.objects.count(), receipts)
                self.assertEqual(AuditLog.objects.count(), audits)
                entry.refresh_from_db()
                self.assertEqual(finance.balance(entry), 10000)

    def test_prepayment_requires_contract_basis_and_manager_confirmation(self):
        entry = self.purchase_entry(receive=False)
        self.request(
            'finance',
            'reconciliations/',
            {
                'entry': entry.pk,
                'kind': 'settlement',
                'approved_amount': '500',
                'counterparty_balance': '500',
                'reason': '未收货不能普通结算',
            },
            409,
        )
        self.request(
            'finance',
            'reconciliations/',
            {
                'entry': entry.pk,
                'kind': 'prepayment',
                'approved_amount': '200',
                'counterparty_balance': '500',
                'reason': '缺条款',
            },
            400,
        )
        statement = self.statement(entry, amount='200', kind='prepayment', confirm=False)
        self.request('finance', f'reconciliations/{statement}/confirm/', {'reason': '财务不能核准预付款'}, 403)
        self.request('manager', f'reconciliations/{statement}/confirm/', {'reason': '合同约定预付40%'})
        self.pay(entry, '201', statement, status=409)
        self.pay(entry, '200', statement)

    def test_partial_receipt_allowance_is_checked_on_application_without_side_effects(self):
        entry = self.purchase_entry(receive=False)
        self.post(
            'warehouse',
            f'purchases/{entry.purchase_id}/receive/',
            {'lines': [{'line': entry.purchase.lines.get().pk, 'quantity': '2'}], 'reason': '部分收货'},
        )
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        counters = list(CodeRule.objects.order_by('pk').values('key', 'counter'))
        result = self.request(
            'finance',
            'reconciliations/',
            {
                'entry': entry.pk,
                'kind': 'settlement',
                'approved_amount': '500',
                'counterparty_balance': '500',
                'reason': '额度超出实际收货',
            },
            409,
        )
        self.assertIn('200.00', str(result['approved_amount']))
        self.assertNotIn('amount', result)
        self.assertFalse(Reconciliation.objects.exists())
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)
        self.assertEqual(list(CodeRule.objects.order_by('pk').values('key', 'counter')), counters)
        statement = self.statement(entry, amount='200')
        self.pay(entry, '200', statement)
        self.assertEqual(finance.balance(entry), 300)

    def test_confirmation_rechecks_allowance_for_existing_over_limit_draft(self):
        entry = self.purchase_entry()
        statement = self.statement(entry, confirm=False)
        # A draft saved before application validation was introduced still needs confirmation validation.
        draft = Reconciliation.objects.get(pk=statement)
        draft.approved_amount = Decimal('500.01')
        draft.save()
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        result = self.request('finance', f'reconciliations/{statement}/confirm/', {'reason': '旧草稿'}, 409)
        self.assertIn('500.00', str(result['approved_amount']))
        draft.refresh_from_db()
        self.assertEqual(draft.status, 'draft')
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)

    def test_invalid_settlement_month_errors_target_request_field(self):
        entry = self.purchase_entry()
        for month in ['2026-13', '9999-01']:
            with self.subTest(month=month):
                result = self.request(
                    'finance',
                    'reconciliations/',
                    {
                        'entry': entry.pk,
                        'settlement_month': month,
                        'approved_amount': '500',
                        'counterparty_balance': '500',
                        'reason': '月份检查',
                    },
                    400,
                )
                self.assertIn('settlement_month', result)
                self.assertNotIn('month', result)
                self.assertFalse(Reconciliation.objects.exists())

    def test_receipt_change_invalidates_confirmed_authorization(self):
        entry = self.purchase_entry(receive=False)
        statement = self.statement(entry, kind='prepayment')
        purchase = entry.purchase
        self.post(
            'warehouse',
            f'purchases/{purchase.pk}/receive/',
            {'lines': [{'line': purchase.lines.get().pk, 'quantity': '1'}], 'reason': '新增收货'},
        )
        self.pay(entry, '100', statement, status=409)
        self.assertFalse(Payment.objects.exists())

    def patch(self, role, path, data):
        response = self.clients[role].patch(
            '/api/business/' + path, data, format='json', HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4())
        )
        self.assertEqual(response.status_code, 200, response.data)

    def detail(self, statement):
        return self.clients['finance'].get(f'/api/business/reconciliations/{statement}/').data

    def test_master_data_renames_keep_confirmed_purchase_settlement_usable(self):
        entry = self.purchase_entry()
        statement = self.statement(entry)
        self.patch('purchaser', f'partners/{self.supplier.pk}/', {'name': '供应商（更正名称）'})
        self.patch('purchaser', f'items/{self.item.pk}/', {'name': '伺服电机（更正名称）'})
        detail = self.detail(statement)
        self.assertTrue(detail['valid'])
        # The statement keeps showing what finance actually reconciled.
        self.assertEqual(detail['snapshot']['purchase']['supplier'], '供应商')
        self.assertEqual(detail['snapshot']['lines'][0]['item__name'], '伺服电机')
        self.pay(entry, '500', statement)
        self.assertEqual(finance.balance(entry), 0)

    def test_project_text_edit_keeps_customer_authorization_valid(self):
        entry = self.project.entries.get(kind='receivable')
        statement = self.statement(entry, amount='100')
        sale = SalesOrder.objects.get(project=self.project)
        self.post(
            'manager',
            f'projects/{self.project.pk}/edit/',
            {'name': '自动装配线（改名）', 'requirements': '补充一行技术说明', 'reason': '完善需求'},
        )
        self.assertNotEqual(SalesOrder.objects.get(pk=sale.pk).updated_at, sale.updated_at)
        detail = self.detail(statement)
        self.assertTrue(detail['valid'])
        self.assertEqual(detail['snapshot']['project']['name'], '自动装配线')
        self.pay(entry, '100', statement)

    def test_contract_amendment_amount_still_invalidates_customer_authorization(self):
        entry = self.project.entries.get(kind='receivable')
        statement = self.statement(entry, amount='100')
        sale = SalesOrder.objects.get(project=self.project)
        document = Document.objects.create(
            project=self.project, category='contract', original_name='补充协议.pdf', file='isolated/a.pdf', size=1
        )
        self.post(
            'manager',
            f'sales/{sale.pk}/amend/',
            {
                'expected_updated_at': sale.updated_at.isoformat(),
                'reason': '增加调试内容',
                'date': TODAY,
                'document': document.pk,
                'amount': '12000.00',
                'equipment_quantity': 1,
                'warranty_months': 12,
                'milestones': [{'title': '增补款', 'amount': '2000.00', 'due_date': TODAY}],
            },
        )
        entry.refresh_from_db()
        self.assertEqual(entry.amount, 10000)
        self.assertFalse(self.detail(statement)['valid'])
        self.pay(entry, '100', statement, status=409)
        self.assertFalse(Payment.objects.exists())

    def test_purchase_line_and_counterparty_facts_still_invalidate_authorization(self):
        entry = self.purchase_entry()
        statement = self.statement(entry)
        line = entry.purchase.lines.get()
        other = Partner.objects.create(code='S2', name='其他供应商', kind='supplier')
        for model, pk, change in [
            (PurchaseLine, line.pk, {'quantity': '6'}),
            (PurchaseLine, line.pk, {'unit_price': '101'}),
            (PurchaseOrder, entry.purchase_id, {'supplier': other}),
            (PurchaseOrder, entry.purchase_id, {'payment_term': 'month30'}),
            (PurchaseOrder, entry.purchase_id, {'payment_days': 45}),
            (PurchaseOrder, entry.purchase_id, {'payment_due_date': timezone.localdate() + timedelta(days=10)}),
            (Entry, entry.pk, {'due_date': timezone.localdate() + timedelta(days=10)}),
            (Entry, entry.pk, {'amount': '501'}),
            (Entry, entry.pk, {'credit_amount': '1'}),
            (Entry, entry.pk, {'cancelled': True}),
            (Project, self.project.pk, {'customer': other}),
        ]:
            with self.subTest(change=change), transaction.atomic():
                model.objects.filter(pk=pk).update(**change)
                self.assertFalse(reconciliation.current(Reconciliation.objects.get(pk=statement)))
                transaction.set_rollback(True)
        self.assertTrue(reconciliation.current(Reconciliation.objects.get(pk=statement)))

    def test_titles_and_purchase_note_keep_original_snapshot_and_authorization(self):
        entry = self.purchase_entry()
        statement = self.statement(entry)
        original = self.detail(statement)['snapshot']
        Entry.objects.filter(pk=entry.pk).update(title='修正款项标题')
        PurchaseOrder.objects.filter(pk=entry.purchase_id).update(note='补充非金额备注')
        self.assertTrue(self.detail(statement)['valid'])
        self.assertEqual(self.detail(statement)['snapshot'], original)
        self.pay(entry, '500', statement)

    def test_other_payment_invalidates_customer_authorization(self):
        entry = self.project.entries.get(kind='receivable')
        statement = self.statement(entry, amount='100')
        self.pay(entry, '1')
        self.assertFalse(self.detail(statement)['valid'])
        self.pay(entry, '100', statement, status=409)
        self.assertEqual(Payment.objects.count(), 1)

    def test_historical_bank_evidence_exception_does_not_allow_locked_refund_or_reversal(self):
        entry = self.project.entries.get(kind='receivable')
        original = Payment.objects.create(entry=entry, amount='10000', date='2026-01-01', reason='历史已登记收款')
        entry.credit_amount = Decimal('100')
        entry.save()
        statement = self.statement(entry, amount='100', kind='refund')
        company, _ = Company.objects.get_or_create(pk=1)
        company.locked_through = '2026-01-31'
        company.save()
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        for path, data in [
            (
                f'entries/{entry.pk}/refund/',
                {'amount': '100', 'date': '2026-01-02', 'reason': '锁定期退款', 'reconciliation': statement},
            ),
            (f'payments/{original.pk}/reverse/', {'date': '2026-01-02', 'reason': '锁定期冲销'}),
        ]:
            with self.subTest(path=path):
                response = self.request('finance', path, data, 409)
                self.assertIn('锁账', str(response))
                self.assertEqual(Payment.objects.count(), 1)
                self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)
                self.assertEqual(finance.balance(entry), -100)

    def test_purchase_return_invalidates_confirmed_authorization(self):
        entry = self.purchase_entry()
        statement = self.statement(entry)
        receipt = entry.purchase.lines.get().stockmove_set.get(kind='receipt')
        self.post('warehouse', f'moves/{receipt.pk}/return-purchase/', {'quantity': '1', 'reason': '退回一件'})
        self.assertFalse(self.detail(statement)['valid'])
        self.pay(entry, '100', statement, status=409)
        self.assertFalse(Payment.objects.exists())

    def test_refund_requires_current_reconciliation_and_keeps_original(self):
        entry = self.purchase_entry(receive=False)
        self.pay(entry, statement=self.statement(entry, kind='prepayment'))
        self.post('purchaser', f'purchases/{entry.purchase_id}/cancel-remainder/', {'reason': '停止采购'})
        entry.refresh_from_db()
        self.pay(entry, '100', refund=True, status=409)
        statement = self.statement(entry, kind='refund')
        self.pay(entry, '100', statement, refund=True)
        self.pay(entry, '400', statement, refund=True)
        self.assertEqual(list(entry.payments.order_by('id').values_list('amount', flat=True)), [500, -100, -400])
        self.assertEqual(finance.balance(entry), 0)

    def test_bank_allocation_is_partial_idempotent_and_does_not_double_cash(self):
        entry = self.project.entries.get(kind='receivable')
        bank = self.bank(project=False)
        payload = {'entry': entry.pk, 'amount': '200', 'reason': '认领合同款'}
        key = str(uuid.uuid4())
        first = self.request('finance', f'bank-records/{bank}/allocate/', payload, key=key)
        self.assertEqual(self.request('finance', f'bank-records/{bank}/allocate/', payload, key=key), first)
        self.assertEqual(Payment.objects.count(), 1)
        self.assertEqual(BankMatch.objects.count(), 1)
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 300)
        self.assertEqual(finance.balance(entry), 9800)
        self.request('finance', f'bank-records/{bank}/allocate/', {**payload, 'amount': '301'}, 409)
        self.assertEqual(Payment.objects.count(), 1)
        with self.assertRaises(Conflict):
            banking.check_close(self.project)
        self.request('finance', f'bank-records/{bank}/allocate/', {**payload, 'amount': '300'})
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 0)

    def test_bank_match_does_not_create_payment_and_unmatch_preserves_history(self):
        entry = self.project.entries.get(kind='receivable')
        bank_id = self.bank()
        bank = BankRecord.objects.get(pk=bank_id)
        payment = self.request(
            'finance',
            f'entries/{entry.pk}/pay/',
            {
                'amount': '500',
                'date': TODAY,
                'reason': '先登记到账',
                'method': 'bank',
                'account': bank.account,
                'reference': bank.reference,
            },
        )['id']
        match = self.request('finance', f'bank-records/{bank_id}/match/', {'payment': payment, 'reason': '核对银行'})[
            'id'
        ]
        self.assertEqual(Payment.objects.count(), 1)
        self.request('finance', f'bank-records/{bank_id}/match/', {'payment': payment, 'reason': '重复匹配'}, 409)
        self.request('finance', f'payments/{payment}/reverse/', {'date': TODAY, 'reason': '先解除匹配'}, 409)
        self.request('finance', f'bank-records/{bank_id}/unmatch/', {'match': match, 'reason': '认领项目有误'})
        self.request('finance', f'payments/{payment}/reverse/', {'date': TODAY, 'reason': '纠错原收款'})
        self.assertEqual(BankMatch.objects.count(), 2)
        self.assertEqual(Payment.objects.count(), 2)
        self.assertEqual(banking.remaining(bank), 500)
        self.assertEqual(finance.balance(entry), 10000)

    def bank_with(self, amount, reference, account='A1', project=True):
        return self.request(
            'finance',
            'bank-records/',
            {
                'amount': amount,
                'date': TODAY,
                'account': account,
                'reference': reference,
                'counterparty': '客户银行户名',
                'reason': '银行实际记录',
                **({'project': self.project.pk} if project else {}),
            },
            201,
        )['id']

    def bank_pay(self, entry, amount, reference, account='A1', status=200, key=None, **extra):
        return self.request(
            'finance',
            f'entries/{entry.pk}/pay/',
            {
                'amount': amount,
                'date': TODAY,
                'reason': '手工登记到账',
                'method': 'bank',
                'account': account,
                'reference': reference,
                **extra,
            },
            status,
            key,
        )

    def test_allocation_rejects_bank_transaction_already_registered_manually(self):
        entry = self.project.entries.get(kind='receivable')
        payment = self.bank_pay(entry, '3000', 'R1')['id']
        bank = self.bank_with('3000', 'R1')
        receipts = ActionReceipt.objects.count()
        rejected = self.request(
            'finance', f'bank-records/{bank}/allocate/', {'entry': entry.pk, 'amount': '3000', 'reason': '认领'}, 409
        )
        self.assertIn('匹配已有收付款', str(rejected['detail']))
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(Payment.objects.count(), 1)
        self.assertFalse(BankMatch.objects.exists())
        self.assertEqual(finance.balance(entry), 7000)
        duplicate = self.bank_pay(entry, '1000', 'R1', status=409)
        self.assertIn('重复登记', str(duplicate['detail']))
        self.assertEqual(Payment.objects.count(), 1)
        self.request('finance', f'bank-records/{bank}/match/', {'payment': payment, 'reason': '核对银行'})
        self.assertEqual(Payment.objects.count(), 1)
        self.assertEqual(finance.balance(entry), 7000)
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 0)

    def test_matched_manual_payment_lets_bank_remainder_be_split_to_other_entries(self):
        entry = self.project.entries.get(kind='receivable')
        other = self.active_project().entries.get(kind='receivable')
        bank = self.bank_with('5000', 'R1', project=False)
        payment = self.bank_pay(entry, '3000', 'R1')['id']
        claim = {'entry': other.pk, 'amount': '2000', 'reason': '认领剩余'}
        self.request('finance', f'bank-records/{bank}/allocate/', claim, 409)
        self.request('finance', f'bank-records/{bank}/match/', {'payment': payment, 'reason': '核对银行'})
        key = str(uuid.uuid4())
        first = self.request('finance', f'bank-records/{bank}/allocate/', claim, key=key)
        self.assertEqual(self.request('finance', f'bank-records/{bank}/allocate/', claim, key=key), first)
        self.assertEqual(Payment.objects.filter(account='A1', reference='R1').count(), 2)
        self.assertEqual(BankMatch.objects.count(), 2)
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 0)
        self.assertEqual(finance.balance(entry), 7000)
        self.assertEqual(finance.balance(other), 8000)

    def test_claimed_transfer_remainder_may_be_registered_manually_within_bank_balance(self):
        # 一笔到账先认领部分到款项，剩余部分手工登记后再匹配，是合法拆分；银行剩余额度限制继续登记。
        entry = self.project.entries.get(kind='receivable')
        other = self.active_project().entries.get(kind='receivable')
        bank = self.bank_with('1100', 'R7')
        self.request(
            'finance', f'bank-records/{bank}/allocate/', {'entry': entry.pk, 'amount': '400', 'reason': '认领'}
        )
        manual = self.bank_pay(entry, '600', 'R7')['id']
        self.request('finance', f'bank-records/{bank}/match/', {'payment': manual, 'reason': '核对银行'})
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 100)
        receipts = ActionReceipt.objects.count()
        rejected = self.bank_pay(entry, '200', 'R7', status=409)
        self.assertIn('尚未匹配的金额只剩 100.00 元', str(rejected['detail']))
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.bank_pay(other, '100', 'R7')
        self.bank_pay(other, '50', 'R7', status=409)
        self.assertEqual(finance.balance(entry), 9000)
        self.assertEqual(finance.balance(other), 9900)

    def test_manual_bank_reference_is_unique_per_entry_and_direction_but_may_split(self):
        entry = self.project.entries.get(kind='receivable')
        other = self.active_project().entries.get(kind='receivable')
        key = str(uuid.uuid4())
        first = self.bank_pay(entry, '3000', 'R1', key=key)
        self.assertEqual(self.bank_pay(entry, '3000', 'R1', key=key), first)
        receipts = ActionReceipt.objects.count()
        rejected = self.bank_pay(entry, '1000', 'R1', status=409)
        self.assertIn('核对流水号', str(rejected['detail']))
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.bank_pay(entry, '1000', 'R2')
        self.bank_pay(entry, '1000', 'R1', account='A2')
        self.bank_pay(other, '1000', 'R1')
        self.pay(entry, '100')
        self.pay(entry, '100')
        self.assertEqual(finance.balance(entry), 4800)
        self.assertEqual(finance.balance(other), 9000)
        purchase = self.purchase_entry(receive=False)
        self.bank_pay(purchase, '500', 'R9', **{'reconciliation': self.statement(purchase, kind='prepayment')})
        self.post('purchaser', f'purchases/{purchase.purchase_id}/cancel-remainder/', {'reason': '停止采购'})
        purchase.refresh_from_db()
        statement = self.statement(purchase, kind='refund')
        refund = {'amount': '100', 'date': TODAY, 'reason': '供应商退回', 'method': 'bank', 'account': 'A1'}
        self.request(
            'finance', f'entries/{purchase.pk}/refund/', {**refund, 'reference': 'R9', 'reconciliation': statement}
        )
        self.assertEqual(finance.balance(purchase), -400)

    def test_over_balance_message_follows_entry_direction(self):
        receivable = self.project.entries.get(kind='receivable')
        self.assertEqual(str(self.pay(receivable, '10000.01', status=409)['detail']), '收款超过当前未收余额。')
        payable = self.purchase_entry()
        self.assertEqual(str(self.pay(payable, '500.01', status=409)['detail']), '付款超过当前未付余额。')
        self.assertFalse(Payment.objects.exists())

    def test_reversed_payment_no_longer_blocks_reregistration_or_allocation(self):
        entry = self.project.entries.get(kind='receivable')
        first = self.bank_pay(entry, '3000', 'R1')['id']
        self.request('finance', f'payments/{first}/reverse/', {'date': TODAY, 'reason': '登记金额有误'})
        second = self.bank_pay(entry, '3000', 'R1')['id']
        bank = self.bank_with('3000', 'R1')
        claim = {'entry': entry.pk, 'amount': '3000', 'reason': '认领'}
        self.request('finance', f'bank-records/{bank}/allocate/', claim, 409)
        self.request('finance', f'payments/{second}/reverse/', {'date': TODAY, 'reason': '改为银行认领'})
        self.request('finance', f'bank-records/{bank}/allocate/', claim)
        self.assertEqual(finance.balance(entry), 7000)
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=bank)), 0)

    def test_bank_direction_account_scope_and_permissions(self):
        bank = self.bank('-500')
        entry = self.project.entries.get(kind='receivable')
        self.request(
            'finance',
            f'bank-records/{bank}/allocate/',
            {'entry': entry.pk, 'amount': '100', 'reason': '支出不能认领成收入'},
            409,
        )
        for role in [
            'manager',
            'sales_manager',
            'purchaser',
            'warehouse',
            'member',
            'purchase_manager',
            'mechanical_engineer',
            'electrical_engineer',
        ]:
            self.assertEqual(self.clients[role].get('/api/business/bank-records/').status_code, 403)
        for role in [
            'sales_manager',
            'purchaser',
            'warehouse',
            'member',
            'purchase_manager',
            'mechanical_engineer',
            'electrical_engineer',
        ]:
            self.assertEqual(self.clients[role].get('/api/business/reconciliations/').status_code, 403)
        other = self.active_project()
        other_entry = other.entries.get()
        self.request(
            'finance',
            f'bank-records/{self.bank()}/allocate/',
            {'entry': other_entry.pk, 'amount': '100', 'reason': '错项目'},
            409,
        )
        self.assertFalse(Payment.objects.exists())

    def test_changed_role_cannot_replay_payment_or_confirmation(self):
        entry = self.purchase_entry()
        statement = self.statement(entry)
        key = str(uuid.uuid4())
        self.pay(entry, statement=statement, key=key)
        self.users['finance'].role = 'member'
        self.users['finance'].save()
        self.pay(entry, statement=statement, key=key, status=403)

    def test_customer_receipt_is_not_blocked_by_unconfirmed_statement(self):
        entry = self.project.entries.get()
        self.statement(entry, amount='100', counterparty='9999', confirm=False)
        self.pay(entry, '100')
        self.assertEqual(finance.balance(entry), 9900)
        with self.assertRaises(Conflict):
            banking.check_close(self.project)

    def test_void_bank_and_reenter_corrected_record_preserves_original(self):
        bank = BankRecord.objects.get(pk=self.bank())
        payload = {
            'amount': '450',
            'date': TODAY,
            'account': bank.account,
            'reference': bank.reference,
            'counterparty': bank.counterparty,
            'reason': '金额录入更正',
        }
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        duplicate = self.request('finance', 'bank-records/', payload, 409)
        self.assertIn('reference', duplicate)
        self.assertIn('此账户和银行流水号已登记', str(duplicate['reference']))
        self.assertNotIn('amount', duplicate)
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)
        bank.refresh_from_db()
        self.assertEqual(bank.amount, 500)
        self.request('finance', f'bank-records/{bank.pk}/void/', {'reason': '录入金额有误'})
        self.request('finance', 'bank-records/', payload, 201)
        self.assertEqual(BankRecord.objects.count(), 2)
        self.assertEqual(Reconciliation.objects.count(), 0)

    def test_public_bank_request_cannot_supply_internal_import_evidence(self):
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        result = self.request(
            'finance',
            'bank-records/',
            {
                'amount': '100',
                'date': TODAY,
                'account': 'B',
                'reference': 'R',
                'counterparty': '客户',
                'reason': '不能伪造导入凭据',
                'import_evidence': {'import_fingerprint': 'forged'},
            },
            400,
        )
        self.assertIn('import_evidence', result)
        self.assertFalse(BankRecord.objects.exists())
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)

    def test_reference_conflict_detected_by_full_clean_after_stale_precheck_uses_same_field(self):
        bank = BankRecord.objects.get(pk=self.bank())
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        # Model validation can observe a competitor committed after the initial existence check.
        with patch.object(BankRecord.objects, 'filter') as precheck:
            precheck.return_value.exists.return_value = False
            response = self.request(
                'finance',
                'bank-records/',
                {
                    'amount': '100',
                    'date': TODAY,
                    'account': bank.account,
                    'reference': bank.reference,
                    'counterparty': bank.counterparty,
                    'reason': '预检后发生竞争',
                },
                409,
            )
        self.assertIn('reference', response)
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)

    def test_unrelated_real_integrity_errors_are_not_reported_as_duplicate_bank_reference(self):
        payload = {
            'amount': '100',
            'date': TODAY,
            'account': 'B',
            'reference': 'R',
            'counterparty': '客户',
            'reason': '核对',
        }
        existing = BankRecord.objects.create(
            amount='1',
            date=TODAY,
            account='OLD',
            reference='OLD',
            counterparty='客户',
            reason='原始记录',
            import_fingerprint='a' * 64,
        )
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())
        for expected, change in [
            ('23514', {'amount': 0}),
            ('23505', {'import_fingerprint': existing.import_fingerprint}),
        ]:
            with self.subTest(sqlstate=expected):

                def invalid_save(bank, actor):
                    for key, value in change.items():
                        setattr(bank, key, value)
                    bank.save()
                    return bank

                with patch('apps.business.services.banking.save', side_effect=invalid_save):
                    with self.assertRaises(IntegrityError) as caught:
                        banking.create(self.users['finance'], str(uuid.uuid4()), payload)
                self.assertEqual(caught.exception.__cause__.pgcode, expected)
                self.assertNotEqual(caught.exception.__cause__.diag.constraint_name, 'lean_bank_reference_unique')
                self.assertEqual(BankRecord.objects.count(), 1)
                self.assertEqual((ActionReceipt.objects.count(), AuditLog.objects.count()), counts)

    def test_unclaimed_overpayment_can_be_returned_without_fake_business_entry(self):
        incoming, outgoing = self.bank('500'), self.bank('-100')
        entry = self.project.entries.get(kind='receivable')
        self.request(
            'finance',
            f'bank-records/{incoming}/allocate/',
            {'entry': entry.pk, 'amount': '400', 'reason': '实际合同回款'},
        )
        payload = {'returned': outgoing, 'amount': '100', 'reason': '多汇部分实际退回'}
        key = str(uuid.uuid4())
        result = self.request('finance', f'bank-records/{incoming}/return-unclaimed/', payload, key=key)
        self.assertEqual(
            self.request('finance', f'bank-records/{incoming}/return-unclaimed/', payload, key=key), result
        )
        for bank_id in [incoming, outgoing]:
            detail = self.clients['finance'].get(f'/api/business/bank-records/{bank_id}/').data
            self.assertEqual(detail['remaining_amount'], '0.00')
        self.assertEqual(Payment.objects.count(), 1)
        self.assertEqual(Entry.objects.count(), 1)
        self.assertEqual(finance.balance(entry), 9600)
        self.request('finance', f'bank-records/{incoming}/return-unclaimed/', payload, 409)
        self.request(
            'finance', f'bank-records/{incoming}/reverse-return/', {'offset': result['id'], 'reason': '关联错了'}
        )
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=incoming)), 100)
        self.assertEqual(banking.remaining(BankRecord.objects.get(pk=outgoing)), 100)
        self.assertEqual(Payment.objects.count(), 1)

    def test_payment_import_cannot_bypass_gate_or_consume_limit_during_preview(self):
        import csv
        import io

        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.business.services.transfers import SCHEMAS

        entry = self.purchase_entry()

        def preview(statement=''):
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow([label for label, _ in SCHEMAS['payments'][1]])
            writer.writerow([entry.pk, '500', TODAY, '批量付款', '', '', '', '', statement])
            return self.clients['finance'].post(
                '/api/business/payments/import-file/',
                {'file': SimpleUploadedFile('payments.csv', output.getvalue().encode())},
                format='multipart',
            )

        denied = preview()
        self.assertEqual(denied.status_code, 200)
        self.assertFalse(denied.data['can_import'])
        statement = self.statement(entry)
        valid = preview(statement)
        self.assertTrue(valid.data['can_import'], valid.data)
        self.assertFalse(Payment.objects.exists())
        self.request('finance', 'payments/import-confirm/', {'token': valid.data['token']})
        self.assertEqual(Payment.objects.get().reconciliation_id, statement)


class BankReferenceConcurrencyTests(BusinessFixtures, TransactionTestCase):
    def setUp(self):
        self.setup_business()

    def test_concurrent_bank_creation_serializes_and_rejects_duplicate(self):
        barrier = Barrier(2)
        payload = {
            'amount': '100',
            'date': TODAY,
            'account': 'SERIALIZED',
            'reference': 'SAME',
            'counterparty': '客户',
            'reason': '同时登记同一银行事实',
        }
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())

        def run(role):
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                try:
                    banking.create(self.users[role], str(uuid.uuid4()), payload)
                    return 201
                except Conflict:
                    return 409
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(run, ['admin', 'finance'])), [201, 409])
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual(ActionReceipt.objects.count(), counts[0] + 1)
        self.assertEqual(AuditLog.objects.count(), counts[1] + 1)
        self.assertFalse(Payment.objects.exists())

    def test_real_unique_constraint_race_returns_reference_conflict_without_side_effects(self):
        barrier = Barrier(2)
        original_save = BankRecord.save
        payload = {
            'amount': '100',
            'date': TODAY,
            'account': 'RACE',
            'reference': 'SAME',
            'counterparty': '客户',
            'reason': '同时登记同一银行事实',
        }
        counts = (ActionReceipt.objects.count(), AuditLog.objects.count())

        def synchronized_save(bank, *args, **kwargs):
            # Both full_clean/precheck calls finish before either INSERT, proving the real 23505 path.
            if bank._state.adding:
                barrier.wait(timeout=10)
            return original_save(bank, *args, **kwargs)

        def run(role):
            close_old_connections()
            try:
                try:
                    return 201, banking.create(self.users[role], str(uuid.uuid4()), payload)
                except Conflict as exc:
                    self.assertEqual(exc.__cause__.__cause__.pgcode, '23505')
                    self.assertEqual(exc.__cause__.__cause__.diag.constraint_name, 'lean_bank_reference_unique')
                    return 409, exc.detail
            finally:
                close_old_connections()

        # Exercise the DB-constraint fallback for an unsynchronized/legacy writer. Normal API account
        # serialization is covered separately; otherwise the barrier inside INSERT would deadlock by design.
        with (
            patch('apps.business.services.bank_identity.lock_accounts'),
            patch.object(BankRecord, 'save', synchronized_save),
            ThreadPoolExecutor(max_workers=2) as pool,
        ):
            results = list(pool.map(run, ['admin', 'finance']))
        self.assertEqual(sorted(status for status, _ in results), [201, 409])
        failure = next(detail for status, detail in results if status == 409)
        self.assertIn('reference', failure)
        self.assertIn('此账户和银行流水号已登记', str(failure['reference']))
        self.assertEqual(BankRecord.objects.count(), 1)
        self.assertEqual(ActionReceipt.objects.count(), counts[0] + 1)
        self.assertEqual(AuditLog.objects.count(), counts[1] + 1)
        self.assertFalse(Payment.objects.exists())


class DataCheckTests(BusinessFixtures, TestCase):
    """Administrator read-only checklist of historical records that today's rules would reject."""

    CATEGORIES = ('duplicate-payments', 'claim-manual', 'duplicate-banks', 'warranty-links', 'deleted-open')

    def setUp(self):
        self.setup_business()
        self.project = self.active_project()
        self.today = timezone.localdate()
        self.stock = Stock.objects.create(item=self.item, location='主仓', quantity=100, value=10000)

    def checks(self, category=None, *, role='admin', status=200, **params):
        response = self.clients[role].get(
            '/api/business/data-checks/', {**({'category': category} if category else {}), **params}
        )
        self.assertEqual(response.status_code, status, getattr(response, 'data', response))
        return response

    def counts(self):
        return {row['key']: row['count'] for row in self.checks().data['categories']}

    def rows(self, category):
        return self.checks(category, page_size=200).data['results']

    def entry(self, title, *, kind='receivable', amount='1000'):
        return Entry.objects.create(project=self.project, kind=kind, title=title, amount=amount, due_date=self.today)

    def pay(self, entry, amount, reference='', account='基本户', **extra):
        return Payment.objects.create(
            entry=entry,
            amount=amount,
            date=self.today,
            reason=extra.pop('reason', '历史登记'),
            account=account,
            reference=reference,
            **extra,
        )

    def bank(self, reference, amount, account='基本户', **extra):
        return BankRecord.objects.create(
            amount=amount,
            date=extra.pop('date', self.today),
            account=account,
            reference=reference,
            counterparty='客户',
            reason='到账',
            **extra,
        )

    def match(self, bank, payment, amount):
        return BankMatch.objects.create(bank=bank, payment=payment, amount=amount, reason='历史匹配')

    def removed(self, obj):
        return AuditLog.objects.create(
            actor=self.users['admin'],
            operation='record.delete',
            resource=f'{obj._meta.model_name}:{obj.pk}',
            detail={'resource': 'history', 'record': {}, 'reason': '历史删除', 'mode': 'archive'},
        )

    def receipt(self, code, quantity='5'):
        purchase = PurchaseOrder.objects.create(
            code=code, project=self.project, supplier=self.supplier, status='received', due_date=self.today
        )
        line = PurchaseLine.objects.create(
            purchase=purchase, item=self.item, quantity=quantity, unit_price='100', received_quantity=quantity
        )
        move = StockMove.objects.create(
            stock=self.stock,
            project=self.project,
            purchase_line=line,
            kind='receipt',
            quantity=quantity,
            value=Decimal(quantity) * 100,
            received_date=self.today,
            reason='采购收货',
        )
        return purchase, move

    def give_back(self, receipt, quantity):
        return StockMove.objects.create(
            stock=self.stock,
            project=self.project,
            purchase_line=receipt.purchase_line,
            source=receipt,
            kind='purchase_return',
            quantity=-Decimal(quantity),
            value=-Decimal(quantity) * 100,
            reason='采购退货',
        )

    def warranty(self, receipt, quantity, **links):
        return PurchaseWarranty.objects.create(
            receipt=receipt, date=self.today, quantity=quantity, description='报修', **links
        )

    def group(self, index):
        """One hit for every category; groups never share an identity with each other."""
        duplicated = self.entry(f'重复登记款项{index}')
        self.pay(duplicated, '100', f'DUP-{index}')
        self.pay(duplicated, '100', f'DUP-{index}')

        amount = Decimal(300 + index)
        claimed = self.bank(f'CLM-{index}', amount)
        self.match(claimed, self.pay(self.entry(f'认领款项{index}'), amount, f'CLM-{index}'), amount)
        self.pay(self.entry(f'手工款项{index}'), amount, f'CLM-{index}')

        self.bank(f'M-{index}', '500', account=f'核查户{index}')
        self.bank(f'IMPORT-HXB-{index}', '500', account=f'核查户{index}', import_fingerprint=f'fp-{index}')

        _, receipt = self.receipt(f'PO-DC-{index}')
        returned = self.give_back(receipt, '2')
        expense = self.entry(f'质保费用{index}', kind='expense', amount='50')
        self.warranty(receipt, '1', returned=returned, expense=expense)
        self.warranty(receipt, '2', returned=returned, expense=expense)

        self.removed(self.entry(f'已删未结款项{index}', amount='100'))
        zero = self.entry(f'已删待确认对账款项{index}', amount='0')
        Reconciliation.objects.create(
            code=f'REC-Z-{index}', entry=zero, kind='settlement', approved_amount='10', reason='待确认'
        )
        self.removed(zero)
        self.removed(self.pay(self.entry(f'未结款项{index}'), '100'))
        self.removed(
            Reconciliation.objects.create(
                code=f'REC-D-{index}',
                entry=self.entry(f'对账款项{index}'),
                kind='settlement',
                approved_amount='10',
                reason='待确认',
            )
        )
        self.removed(self.bank(f'DEL-{index}', '700', account=f'删除户{index}'))
        open_bank = self.bank(f'DM-{index}', '800', account=f'匹配户{index}')
        self.removed(self.match(open_bank, self.pay(self.entry(f'匹配款项{index}'), '100'), '100'))
        self.removed(
            Delivery.objects.create(code=f'DEL-DC-{index}', project=self.project, quantity=1, shipped_date=self.today)
        )

    def test_every_category_lists_its_hits_with_links_and_reasons(self):
        self.group(1)
        self.assertEqual(
            self.counts(),
            {'duplicate-payments': 2, 'claim-manual': 1, 'duplicate-banks': 2, 'warranty-links': 2, 'deleted-open': 7},
        )
        body = self.checks().data
        self.assertEqual(body['category'], 'duplicate-payments')
        self.assertIn('只读核查，不会修改任何数据', body['notice'])
        first = body['results'][0]
        self.assertEqual(
            set(first),
            {
                'category',
                'category_label',
                'record_type',
                'record_type_label',
                'record_id',
                'title',
                'project_code',
                'amount',
                'problem',
                'path',
            },
        )
        self.assertEqual(first['project_code'], self.project.code)
        self.assertEqual(first['amount'], '100.00')
        self.assertIn('可能重复登记', first['problem'])
        self.assertEqual(
            first['path'], f'/finance?section=payments&resource=payments&focus={first["record_id"]}&project=all'
        )

        manual = self.rows('claim-manual')[0]
        self.assertEqual(Payment.objects.get(pk=manual['record_id']).entry.title, '手工款项1')
        self.assertIn('已有 1 条收付款匹配了银行流水', manual['problem'])
        self.assertIn('可能重复计收', manual['problem'])

        banks = self.rows('duplicate-banks')
        self.assertEqual({row['record_type'] for row in banks}, {'bank-records'})
        self.assertIn('本条为手工录入', banks[0]['problem'])
        self.assertIn('本条为导入记录', banks[1]['problem'])
        self.assertTrue(banks[0]['path'].startswith('/finance?section=bank&resource=bank-records&focus='))

        warranties = self.rows('warranty-links')
        self.assertIn('同时关联了 2 张质保单', warranties[0]['problem'])
        self.assertIn('关联退货数量 2.000 超过本质保单报修数量 1.000', warranties[0]['problem'])
        self.assertIn('费用「质保费用1」同时关联了 2 张质保单', warranties[0]['problem'])
        self.assertNotIn('超过本质保单报修数量', warranties[1]['problem'])
        purchase = PurchaseOrder.objects.get(code='PO-DC-1')
        self.assertEqual(warranties[0]['path'], f'/purchases?resource=purchases&focus={purchase.pk}')

        deleted = self.rows('deleted-open')
        self.assertEqual(
            [row['record_type'] for row in deleted],
            ['entries', 'entries', 'payments', 'reconciliations', 'bank-records', 'bank-matches', 'deliveries'],
        )
        self.assertTrue(all(row['problem'].startswith('已移入“已删除记录”，但按现行在办规则仍在办') for row in deleted))
        self.assertIn('账款余额 100.00 元尚未结清', deleted[0]['problem'])
        self.assertIn('待确认的对账单', deleted[1]['problem'])
        self.assertIn('余额 900.00 元尚未结清', deleted[2]['problem'])
        self.assertIn('对账单待确认', deleted[3]['problem'])
        self.assertIn('还有 700.00 元未认领或未匹配', deleted[4]['problem'])
        self.assertIn('还有 700.00 元未认领或未匹配', deleted[5]['problem'])
        self.assertIn('交付批次尚未验收', deleted[6]['problem'])
        delivery = Delivery.objects.get(code='DEL-DC-1')
        self.assertEqual(
            deleted[6]['path'],
            f'/projects/{self.project.pk}?tab=deliveries&section=deliveries&resource=deliveries&focus={delivery.pk}',
        )

    def test_records_that_current_rules_accept_are_not_listed(self):
        entry = self.entry('正常款项')
        reversed_payment = self.pay(entry, '100', 'R-1')
        self.pay(entry, '-100', 'R-1', reversal_of=reversed_payment, reason='冲销')
        self.pay(entry, '100', 'R-1')
        self.pay(entry, '100', '')
        self.pay(entry, '100', '')
        # Same reference in another entry is a split of one transfer; a refund is the other direction.
        self.pay(self.entry('拆分款项'), '100', 'R-1')
        self.pay(entry, '-50', 'R-1')
        # All registrations of a claimed transfer are matched.
        bank = self.bank('ALL-MATCHED', '200')
        for title in ('认领一', '认领二'):
            self.match(bank, self.pay(self.entry(title), '100', 'ALL-MATCHED'), '100')
        # Bank groups: two imported rows, a voided manual row, or a different date.
        self.bank('A', '600', account='导入户', import_fingerprint='fp-a')
        self.bank('B', '600', account='导入户', source={'bank': 'ICBC'})
        self.bank('C', '600', account='作废户', void_reason='误录')
        self.bank('D', '600', account='作废户')
        self.bank('E', '600', account='日期户')
        self.bank('F', '600', account='日期户', date=self.today - timedelta(days=1))
        # Warranty linked to its own partial return and expense.
        _, receipt = self.receipt('PO-OK')
        self.warranty(receipt, '2', returned=self.give_back(receipt, '1'), expense=self.entry('质保', kind='expense'))
        # Deleted but ended: settled entry, void statement, accepted delivery out of warranty, restored record.
        settled = self.entry('已结清', amount='100')
        self.pay(settled, '100')
        self.removed(settled)
        self.removed(
            Reconciliation.objects.create(
                code='REC-V', entry=settled, kind='settlement', approved_amount='10', reason='作废', status='void'
            )
        )
        self.removed(
            Delivery.objects.create(
                code='DEL-OK',
                project=self.project,
                quantity=1,
                shipped_date=self.today,
                accepted_date=self.today,
                warranty_until=self.today - timedelta(days=1),
            )
        )
        restored = self.entry('已恢复未结款项', amount='100')
        self.removed(restored)
        AuditLog.objects.create(
            actor=self.users['admin'], operation='record.restore', resource=f'entry:{restored.pk}', detail={}
        )
        # In progress but never deleted.
        self.entry('未删除未结款项', amount='100')
        self.assertEqual(self.counts(), dict.fromkeys(self.CATEGORIES, 0))

    def test_only_administrators_can_read_or_export(self):
        self.group(1)
        for role in ('finance', 'manager', 'purchaser', 'member'):
            self.checks(role=role, status=403)
            self.checks('deleted-open', role=role, status=403, file_format='csv')
        actor = self.users['finance']
        actor.additional_roles = ['admin']
        actor.save()
        self.assertEqual(self.checks(role='finance').data['categories'][0]['count'], 2)
        self.checks('unknown', status=400)
        self.checks(status=400, file_format='pdf')
        self.assertIn(self.client.get('/api/business/data-checks/').status_code, {401, 403})

    def test_export_uses_chinese_name_and_blocks_formulas(self):
        entry = self.entry('=HYPERLINK("http://bad")')
        self.pay(entry, '100', '+SUM(1)')
        self.pay(entry, '100', '+SUM(1)')
        response = self.checks('duplicate-payments', file_format='csv')
        self.assertIn("filename*=utf-8''", response['Content-Disposition'])
        self.assertIn(quote('历史数据核查_同一款项重复登记的银行收付款_'), response['Content-Disposition'])
        lines = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(
            lines[0], ['类别', '记录类型', '记录ID', '编号/标题', '项目编号', '金额', '问题说明', '原单路径']
        )
        self.assertEqual(len(lines), 3)
        self.assertEqual(lines[1][0], '同一款项重复登记的银行收付款')
        self.assertTrue(lines[1][3].startswith('\'=HYPERLINK("http://bad")'), lines[1][3])
        self.assertEqual(lines[1][5], '100.00')
        self.assertTrue(lines[1][7].startswith('/finance?section=payments'))
        workbook = load_workbook(io.BytesIO(self.checks('duplicate-payments', file_format='xlsx').content))
        values = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(values[1][5], 100)
        self.assertTrue(values[1][3].startswith("'="))

    def test_checks_are_read_only(self):
        self.group(1)
        models = (Entry, Payment, BankRecord, BankMatch, Reconciliation, PurchaseWarranty, Delivery, StockMove)

        def snapshot():
            return (
                [
                    (model.all_objects.count(), model.all_objects.order_by('-updated_at').values('updated_at').first())
                    for model in models
                ],
                ActionReceipt.objects.count(),
                AuditLog.objects.count(),
            )

        before = snapshot()
        for category in self.CATEGORIES:
            self.checks(category)
            self.checks(category, file_format='csv')
            self.checks(category, file_format='xlsx')
        self.assertEqual(snapshot(), before)

    def test_query_count_does_not_grow_with_problem_records(self):
        def measured():
            total = 0
            for category in self.CATEGORIES:
                with CaptureQueriesContext(connection) as queries:
                    self.checks(category, page_size=200)
                total += len(queries)
            return total

        for index in range(2):
            self.group(index)
        small = measured()
        self.assertEqual(self.counts()['deleted-open'], 14)
        for index in range(2, 10):
            self.group(index)
        self.assertEqual(self.counts()['deleted-open'], 70)
        self.assertEqual(measured(), small)
