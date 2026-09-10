import io
import tempfile
import uuid
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError

from apps.business.models import BOMLine, Document, SalesOrder
from apps.core.models import ActionReceipt

from .test_commercial_chain import BusinessFixtures


class ImportTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()
        self.project = self.active_project()

    def preview(self, content, filename='bom.csv', role='manager', status=200):
        response = self.clients[role].post(
            f'/api/business/projects/{self.project.pk}/import-preview/',
            {'file': SimpleUploadedFile(filename, content)},
            format='multipart',
        )
        self.assertEqual(response.status_code, status, response.data)
        return response.data

    def test_csv_preview_and_confirm_reuse_bom_validation(self):
        preview = self.preview(('\ufeff物料编码,数量,变更说明\r\n' + self.item.code + ',2.125,首版\r\n').encode())
        self.assertTrue(preview['can_import'])
        self.assertFalse(BOMLine.objects.exists())
        payload = {
            'expected_revision': preview['expected_revision'],
            'lines': [{field: row[field] for field in ('item', 'quantity', 'change_note')} for row in preview['lines']],
        }
        self.post('manager', f'projects/{self.project.pk}/revise-bom/', payload)
        self.assertEqual(str(BOMLine.objects.get().quantity), '2.125')
        self.post('manager', f'projects/{self.project.pk}/revise-bom/', payload, status=409)

    def test_preview_reports_row_errors_without_python_error_repr(self):
        preview = self.preview(
            f'物料编码,数量,变更说明\nUNKNOWN,1,\n{self.item.code},NaN,\n{self.item.code},2,\n'.encode()
        )
        self.assertFalse(preview['can_import'])
        self.assertEqual([row['row'] for row in preview['errors']], [2, 3, 4])
        self.assertNotIn('ErrorDetail', str(preview))
        self.assertFalse(BOMLine.objects.exists())

    def test_xlsx_values_supported_and_formulas_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['物料编码', '数量', '变更说明'])
        sheet.append([self.item.code, 2, '首版'])
        content = io.BytesIO()
        workbook.save(content)
        self.assertTrue(self.preview(content.getvalue(), 'bom.xlsx')['can_import'])
        sheet['B2'] = '=1+1'
        content = io.BytesIO()
        workbook.save(content)
        self.preview(content.getvalue(), 'bom.xlsx', status=400)
        workbook.close()

    def test_xlsx_extra_columns_are_not_silently_ignored(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['物料编码', '数量', '变更说明'])
        sheet.append([self.item.code, 2, '首版'])
        sheet['Z2'] = 999
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        self.preview(content.getvalue(), 'bom.xlsx', status=400)

    def test_header_size_and_row_limits(self):
        for content in [
            b'',
            b'wrong,header\nI1,2\n',
            b'a' * (5 * 1024 * 1024 + 1),
            ('物料编码,数量,变更说明\n' + '\n' * 1001).encode(),
        ]:
            self.preview(content, status=400)

    def test_xlsx_entities_and_decompression_limits_are_rejected(self):
        for name, content in [
            ('evil.xml', b'<!DOCTYPE x [<!ENTITY e "expanded">]><x>&e;</x>'),
            ('large.bin', b'x' * (20 * 1024 * 1024 + 1)),
        ]:
            archive = io.BytesIO()
            with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as zipped:
                zipped.writestr(name, content)
            self.preview(archive.getvalue(), 'bom.xlsx', status=400)

    def test_role_gate_and_template(self):
        self.preview(f'物料编码,数量,变更说明\n{self.item.code},2,\n'.encode(), role='member', status=403)
        response = self.clients['manager'].get('/api/business/projects/bom-template/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode('utf-8-sig').strip(), '物料编码,单元,需求数量,变更说明')

    def test_new_bom_template_order_and_readable_preview(self):
        from openpyxl import load_workbook

        response = self.clients['manager'].get('/api/business/projects/bom-template/', {'file_format': 'xlsx'})
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual([cell.value for cell in book.active[1]], ['物料编码', '单元', '需求数量', '变更说明'])
        book.active.append([self.item.code, '装配单元', '2.125', '初版'])
        content = io.BytesIO()
        book.save(content)
        book.close()
        preview = self.preview(content.getvalue(), 'bom.xlsx')
        self.assertTrue(preview['can_import'])
        line = preview['lines'][0]
        self.assertEqual(
            (line['item_code'], line['item_name'], line['assembly_unit'], line['quantity']),
            (self.item.code, self.item.name, '装配单元', '2.125'),
        )
        self.post(
            'manager',
            f'projects/{self.project.pk}/revise-bom/',
            {
                'expected_revision': preview['expected_revision'],
                'lines': [{key: line[key] for key in ['item', 'assembly_unit', 'quantity', 'change_note']}],
            },
        )
        self.assertEqual(BOMLine.objects.get().assembly_unit, '装配单元')

    def test_existing_bom_requires_reason_and_observes_ordered_floor(self):
        self.post(
            'manager',
            f'projects/{self.project.pk}/revise-bom/',
            {
                'expected_revision': self.bom_version(self.project.pk),
                'lines': [{'item': self.item.pk, 'quantity': '5'}],
            },
        )
        self.assertFalse(self.preview(f'物料编码,数量,变更说明\n{self.item.code},5,\n'.encode())['can_import'])
        self.purchase(self.project)
        self.assertFalse(self.preview(f'物料编码,数量,变更说明\n{self.item.code},4,缩减\n'.encode())['can_import'])


class DocumentTests(BusinessFixtures, TestCase):
    def owner_upload(self, role, owner, owner_id, *, key=None, status=201, category='contract'):
        response = self.clients[role].post(
            '/api/business/documents/',
            {
                owner: owner_id,
                'category': category,
                'file': SimpleUploadedFile('单据合同.pdf', b'private contract'),
            },
            format='multipart',
            HTTP_IDEMPOTENCY_KEY=key or str(uuid.uuid4()),
        )
        self.assertEqual(response.status_code, status, response.data)
        return response.data

    def draft_sale(self):
        return self.post(
            'sales_manager',
            'sales/',
            {
                'name': '未签约销售',
                'customer': self.customer.pk,
                'manager': self.users['sales_manager'].pk,
            },
            status=201,
        )['id']

    def test_draft_sales_contract_replay_and_project_aggregation_after_signing(self):
        from .test_commercial_chain import TODAY

        sale = self.draft_sale()
        uploaded = self.owner_upload('sales_manager', 'sale', sale, key='contract')
        self.assertEqual(uploaded, self.owner_upload('sales_manager', 'sale', sale, key='contract'))
        doc = Document.objects.get(pk=uploaded['id'])
        self.assertIsNone(doc.project_id)
        self.assertEqual(doc.sale_id, sale)
        self.post('sales_manager', f'sales/{sale}/quote/', {'amount': '100', 'reason': '确认'})
        signed = self.post(
            'sales_manager',
            f'sales/{sale}/sign/',
            {
                'date': TODAY,
                'manager': self.users['manager'].pk,
                'milestones': [{'title': '合同款', 'amount': '100', 'due_date': TODAY}],
            },
        )
        listing = self.clients['manager'].get(f'/api/business/documents/?project={signed["project"]}').data
        self.assertEqual([row['id'] for row in listing['results']], [doc.pk])
        self.assertEqual(listing['results'][0]['project'], signed['project'])
        self.assertEqual(Document.objects.count(), 1)
        self.assertEqual(len([p for p in Path(self.directory.name).rglob('*') if p.is_file()]), 1)
        self.assertEqual(
            self.clients['sales_manager'].get(f'/api/business/documents/{doc.pk}/download/').status_code, 200
        )

    def test_order_attachment_permissions_apply_to_listing_export_and_download(self):
        sale = SalesOrder.objects.get(project=self.project)
        sales_doc = self.owner_upload('manager', 'sale', sale.pk, category='other')['id']
        purchase = self.purchase(self.project)
        purchase_doc = self.owner_upload('purchaser', 'purchase', purchase.pk)['id']
        project_doc = self.upload(role='manager', category='contract')['id']
        for role, visible in [
            ('manager', {sales_doc, purchase_doc, project_doc}),
            ('finance', {sales_doc, purchase_doc, project_doc}),
            ('purchaser', {purchase_doc}),
            ('warehouse', set()),
            ('member', set()),
            ('sales_manager', set()),
        ]:
            client = self.clients[role]
            listing = client.get(f'/api/business/documents/?project={self.project.pk}')
            self.assertEqual({row['id'] for row in listing.data['results']}, visible)
            for doc_id in [sales_doc, purchase_doc, project_doc]:
                self.assertEqual(
                    client.get(f'/api/business/documents/{doc_id}/download/').status_code,
                    200 if doc_id in visible else 404,
                )
            exported = client.get('/api/business/documents/export/?file_format=csv')
            self.assertEqual(exported.status_code, 200)
            if not visible:
                self.assertNotIn('单据合同.pdf', exported.content.decode('utf-8-sig'))
        self.owner_upload('sales_manager', 'sale', sale.pk, status=403)
        self.owner_upload('warehouse', 'purchase', purchase.pk, status=403)
        self.owner_upload('finance', 'purchase', purchase.pk, status=403)

    def test_order_replay_rechecks_current_owner_and_role(self):
        sale_id = self.draft_sale()
        self.owner_upload('sales_manager', 'sale', sale_id, key='own-doc')
        SalesOrder.objects.filter(pk=sale_id).update(manager=self.users['manager'])
        self.owner_upload('sales_manager', 'sale', sale_id, key='own-doc', status=403)
        self.assertEqual(self.clients['sales_manager'].get('/api/business/documents/').data['count'], 0)
        purchase = self.purchase(self.project)
        self.owner_upload('purchaser', 'purchase', purchase.pk, key='po-doc')
        self.users['purchaser'].role = 'warehouse'
        self.users['purchaser'].save(update_fields=['role'])
        self.owner_upload('purchaser', 'purchase', purchase.pk, key='po-doc', status=403)
        self.assertEqual(Document.objects.count(), 2)

    def test_order_upload_rejects_ambiguous_owner_and_wrong_category(self):
        sale_id = self.draft_sale()
        self.owner_upload('sales_manager', 'sale', sale_id, category='receipt', status=400)
        response = self.clients['admin'].post(
            '/api/business/documents/',
            {
                'project': self.project.pk,
                'sale': sale_id,
                'category': 'contract',
                'file': SimpleUploadedFile('a.pdf', b'a'),
            },
            format='multipart',
            HTTP_IDEMPOTENCY_KEY='invalid-owner',
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Document.objects.exists())
        self.post('sales_manager', f'sales/{sale_id}/cancel/', {'reason': '取消'})
        self.owner_upload('sales_manager', 'sale', sale_id, status=409)

    def test_sales_contract_can_support_amendment_but_purchase_contract_cannot(self):
        from .test_commercial_chain import TODAY

        sale = SalesOrder.objects.get(project=self.project)
        purchase = self.purchase(self.project)
        purchase_doc = self.owner_upload('purchaser', 'purchase', purchase.pk)['id']
        sales_doc = self.owner_upload('manager', 'sale', sale.pk)['id']
        payload = {
            'expected_updated_at': sale.updated_at.isoformat(),
            'document': purchase_doc,
            'reason': '调整设备数量',
            'date': TODAY,
            'amount': '10000.00',
            'equipment_quantity': 2,
            'warranty_months': 12,
        }
        self.post('manager', f'sales/{sale.pk}/amend/', payload, status=404)
        self.post('manager', f'sales/{sale.pk}/amend/', {**payload, 'document': sales_doc})
        sale.refresh_from_db()
        self.assertEqual(sale.equipment_quantity, 2)

    def setUp(self):
        self.setup_business()
        self.project = self.active_project()
        self.directory = tempfile.TemporaryDirectory(prefix='lean-documents-test-')
        self.addCleanup(self.directory.cleanup)
        settings = override_settings(MEDIA_ROOT=self.directory.name)
        settings.enable()
        self.addCleanup(settings.disable)

    def upload(
        self,
        role='member',
        category='drawing',
        content=b'protected drawing',
        *,
        key=None,
        status=201,
        name='assembly.pdf',
    ):
        response = self.clients[role].post(
            '/api/business/documents/',
            {'project': self.project.pk, 'category': category, 'file': SimpleUploadedFile(name, content)},
            format='multipart',
            HTTP_IDEMPOTENCY_KEY=key or str(uuid.uuid4()),
        )
        self.assertEqual(response.status_code, status, response.data)
        return response.data

    def test_project_attachment_is_protected_and_not_exposed_as_public_media(self):
        document_id = self.upload()['id']
        document = Document.objects.get(pk=document_id)
        response = self.clients['member'].get(f'/api/business/documents/{document_id}/')
        self.assertNotIn('file', response.data)
        self.assertNotIn('protected/', str(response.data))
        download = self.clients['member'].get(response.data['download_url'])
        self.assertEqual(download.status_code, 200)
        self.assertIn('attachment;', download['Content-Disposition'])
        self.assertEqual(download['Cache-Control'], 'private, no-store')
        self.assertEqual(b''.join(download.streaming_content), b'protected drawing')
        self.project.members.clear()
        self.assertEqual(
            self.clients['member'].get(f'/api/business/documents/{document_id}/download/').status_code, 404
        )
        self.assertEqual(self.clients['member'].get('/api/business/documents/').data['count'], 0)
        self.assertEqual(self.clients['manager'].get('/media/' + document.file.name).status_code, 404)

    def test_contract_and_receipt_categories_enforce_roles(self):
        self.upload(category='contract', status=403)
        contract = self.upload(role='manager', category='contract')['id']
        self.upload(role='manager', category='receipt', status=403)
        receipt = self.upload(role='finance', category='receipt')['id']
        for role in ('member', 'warehouse', 'purchaser'):
            self.assertEqual(self.clients[role].get('/api/business/documents/').data['count'], 0)
            self.assertEqual(self.clients[role].get(f'/api/business/documents/{contract}/download/').status_code, 404)
        self.assertEqual(self.clients['finance'].get(f'/api/business/documents/{receipt}/').status_code, 200)

    def test_replayed_upload_writes_one_file_and_changed_content_conflicts(self):
        key = str(uuid.uuid4())
        first = self.upload(key=key)
        self.assertEqual(self.upload(key=key), first)
        self.upload(key=key, content=b'different bytes', status=409)
        self.assertEqual(Document.objects.count(), 1)
        self.assertEqual(len([path for path in Path(self.directory.name).rglob('*') if path.is_file()]), 1)

    def test_database_failure_removes_written_file_and_receipt(self):
        key = str(uuid.uuid4())
        with patch('apps.business.services.documents.audit', side_effect=ValidationError('模拟事务失败')):
            self.upload(key=key, status=400)
        self.assertFalse(Document.objects.exists())
        self.assertFalse(ActionReceipt.objects.filter(key=key).exists())
        self.assertFalse(any(path.is_file() for path in Path(self.directory.name).rglob('*')))

    def test_missing_file_returns_safe_not_found(self):
        document = Document.objects.get(pk=self.upload()['id'])
        document.file.storage.delete(document.file.name)
        response = self.clients['member'].get(f'/api/business/documents/{document.pk}/download/')
        self.assertEqual(response.status_code, 404)
        self.assertNotIn(self.directory.name, str(response.data))

    def test_large_attachment_uses_file_upload_limits(self):
        self.upload(content=b'x' * (6 * 1024 * 1024))
        self.upload(content=b'x' * (20 * 1024 * 1024 + 1), status=400)
        self.assertEqual(Document.objects.count(), 1)

    def test_cancelled_project_accepts_refund_receipts_only(self):
        self.post('manager', f'projects/{self.project.pk}/cancel/', {'reason': '客户取消'})
        self.upload(status=409)
        self.upload(role='finance', category='receipt')
