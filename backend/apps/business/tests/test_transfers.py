import csv
import io
import zipfile
from unittest.mock import patch
from xml.etree import ElementTree

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.accounts.models import User
from apps.business.models import (
    BOMLine,
    Item,
    Partner,
    Payment,
    Project,
    PurchaseOrder,
    SalesOrder,
    Stock,
    Task,
    TimeEntry,
)
from apps.business.services import transfers
from apps.business.services.import_guides import GUIDE_MARKER, GUIDE_SHEET
from apps.core.models import ActionReceipt, AuditLog, CodeRule

from .test_commercial_chain import TODAY, BusinessFixtures


class TransferTests(BusinessFixtures, TestCase):
    def setUp(self):
        self.setup_business()

    def preview(self, resource, rows, role='admin', xlsx=False):
        headers = [label for label, _ in transfers.LEGACY_SCHEMAS[resource][1]]
        if xlsx:
            output = io.BytesIO()
            book = Workbook()
            book.active.append(headers)
            for row in rows:
                book.active.append(row)
            book.save(output)
            content = output.getvalue()
        else:
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(headers)
            writer.writerows(rows)
            content = output.getvalue().encode()
        return self.clients[role].post(
            f'/api/business/{resource}/import-file/',
            {'file': SimpleUploadedFile('rows.xlsx' if xlsx else 'rows.csv', content)},
            format='multipart',
        )

    def commit(self, resource, preview, role='admin', status=200):
        return self.post(role, resource + '/import-confirm/', {'token': preview.data['token']}, status=status)

    def test_preview_rollback_confirmation_idempotency_and_changed_data(self):
        counter = CodeRule.objects.get(key='item').counter
        receipts = ActionReceipt.objects.count()
        preview = self.preview('items', [['NEW', '新物料', '规格', '件'], ['', '自动编码', '', '台']], xlsx=True)
        self.assertTrue(preview.data['can_import'], preview.data)
        self.assertEqual(Item.objects.count(), 1)
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(CodeRule.objects.get(key='item').counter, counter)
        result = self.commit('items', preview)
        self.assertEqual(result['count'], 2)
        self.assertEqual(self.commit('items', preview), result)
        self.assertEqual(Item.objects.count(), 3)
        later = self.preview('items', [['A', '一', '', '件'], ['B', '二', '', '件']])
        Item.objects.create(code='B', name='被其他操作占用')
        self.commit('items', later, status=400)
        self.assertFalse(Item.objects.filter(code='A').exists())

    def test_invalid_duplicate_token_permission_and_expiry(self):
        preview = self.preview('items', [['X', '一', '', '件'], ['X', '二', '', '件']])
        self.assertFalse(preview.data['can_import'])
        self.assertEqual(preview.data['errors'][0]['row'], 3)
        self.assertFalse(Item.objects.filter(code='X').exists())
        self.assertEqual(self.preview('items', [['X', '一', '', '件']], role='member').status_code, 403)
        valid = self.preview('items', [['X', '一', '', '件']])
        self.commit('items', valid, role='manager', status=400)
        self.post('admin', 'items/import-confirm/', {'token': valid.data['token'] + 'bad'}, status=400)
        with patch('django.core.signing.time.time', return_value=9999999999):
            self.commit('items', valid, status=400)
        self.users['admin'].role = 'member'
        self.users['admin'].save()
        self.commit('items', valid, status=403)

    def test_masterdata_sales_projects_and_grouped_purchase_import(self):
        self.users['manager'].display_name = '刘经理'
        self.users['manager'].save()
        self.users['member'].display_name = '吴兴广'
        self.users['member'].save()
        partners = self.preview('partners', [['往来', 'both', '联系人', '001234', '地址']])
        self.commit('partners', partners)
        self.assertEqual(Partner.objects.get(name='往来').phone, '001234')
        detail = ['导入需求', 'C1', '刘经理', '说明', TODAY, '2', '12']
        self.commit('sales', self.preview('sales', [detail]))
        self.assertEqual(SalesOrder.objects.get(name='导入需求').status, 'draft')
        self.commit('projects', self.preview('projects', [detail + ['吴兴广；mechanical_engineer']]))
        project = Project.objects.get(name='导入需求')
        self.assertEqual(project.manager_id, self.users['manager'].pk)
        self.assertEqual(SalesOrder.objects.get(name='导入需求').manager_id, self.users['manager'].pk)
        self.assertTrue(project.members.filter(username='member').exists())
        self.assertTrue(project.members.filter(username='mechanical_engineer').exists())
        other = Item.objects.create(code='I2', name='第二物料')
        rows = [
            ['P1', project.code, 'S1', TODAY, '', 'I1', '2', '10'],
            ['P1', project.code, 'S1', TODAY, '', other.code, '3', '20'],
        ]
        preview = self.preview('purchases', rows, role='purchaser')
        self.assertTrue(preview.data['can_import'], preview.data)
        self.commit('purchases', preview, role='purchaser')
        purchase = PurchaseOrder.objects.get()
        self.assertEqual(purchase.status, 'draft')
        self.assertEqual(purchase.lines.count(), 2)
        self.assertFalse(project.entries.exists())
        exported = self.clients['purchaser'].get('/api/business/purchases/export/')
        exported_rows = list(csv.reader(io.StringIO(exported.content.decode('utf-8-sig'))))
        self.assertEqual(len(exported_rows), 3)
        self.assertIn('物料编码', exported_rows[0])

    def test_operational_imports_reuse_finance_inventory_and_execution(self):
        self.users['member'].display_name = '吴兴广'
        self.users['member'].save()
        self.users['manager'].display_name = '刘经理'
        self.users['manager'].save()
        project = self.active_project()
        self.commit('tasks', self.preview('tasks', [[project.code, 'design', '设计任务', '', '吴兴广', TODAY]]))
        task = Task.objects.get(title='设计任务')
        self.assertEqual(task.assignee_id, self.users['member'].pk)
        timed = self.preview('time', [[str(task.pk), '吴兴广', TODAY, '2', '设计']], role='member')
        self.assertTrue(timed.data['can_import'], timed.data)
        self.commit('time', timed, role='member')
        self.assertEqual(TimeEntry.objects.get(task=task).user_id, self.users['member'].pk)
        bad = self.preview('time', [[str(task.pk), '刘经理', TODAY, '2', '越权']], role='member')
        self.assertFalse(bad.data['can_import'])
        self.commit(
            'entries', self.preview('entries', [[project.code, '差旅', '100', TODAY]], role='finance'), role='finance'
        )
        entry = project.entries.get(kind='expense')
        payment = self.preview('payments', [[str(entry.pk), '30', TODAY, '支付']], role='finance')
        self.commit('payments', payment, role='finance')
        self.assertEqual(Payment.objects.get(entry=entry).amount, 30)
        excessive = self.preview('payments', [[str(entry.pk), '80', TODAY, '超付']], role='finance')
        self.assertFalse(excessive.data['can_import'])
        self.commit('stocks', self.preview('stocks', [['I1', '主仓', '2', '10', '期初']]))
        stock = Stock.objects.get(item=self.item)
        BOMLine.objects.create(project=project, item=self.item, quantity=1)
        self.commit(
            'moves',
            self.preview('moves', [[project.code, str(stock.pk), '', '1', '领料']], role='warehouse'),
            role='warehouse',
        )
        for stage in ['design', 'assembly', 'test']:
            Task.objects.update_or_create(
                project=project,
                kind=stage,
                defaults={'title': stage, 'assignee': self.users['member'], 'status': 'done'},
            )
        delivery = self.preview('deliveries', [[project.code, '1', TODAY, '吴兴广（member）', '刘经理', '发货']])
        self.assertTrue(delivery.data['can_import'], delivery.data)
        self.commit('deliveries', delivery)
        self.assertEqual(project.deliveries.count(), 1)
        self.assertEqual(project.tasks.get(kind='install').assignee_id, self.users['member'].pk)
        self.assertEqual(project.tasks.get(kind='acceptance').assignee_id, self.users['manager'].pk)

    def test_task_import_accepts_name_username_and_explicit_identity(self):
        user = self.users['member']
        user.username = 'wuxingguang'
        user.display_name = '吴兴广'
        user.save()
        project = self.active_project()
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        rows = [
            [project.code, 'design', '按账号指派', '', user.username, TODAY],
            [project.code, 'design', '按姓名指派', '', '  吴兴广  ', TODAY],
            [project.code, 'design', '明确同名人员', '', '吴兴广（wuxingguang）', TODAY],
            [project.code, 'design', '英文括号', '', '吴兴广(wuxingguang)', TODAY],
        ]
        preview = self.preview('tasks', rows, xlsx=True)
        self.assertTrue(preview.data['can_import'], preview.data)
        self.assertFalse(project.tasks.exists())
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(AuditLog.objects.count(), audits)
        result = self.commit('tasks', preview)
        self.assertEqual(result['count'], 4)
        self.assertEqual(set(project.tasks.values_list('assignee_id', flat=True)), {user.pk})
        self.assertEqual(self.commit('tasks', preview), result)
        self.assertEqual(project.tasks.count(), 4)

    def test_duplicate_names_and_account_name_collisions_require_explicit_identity(self):
        project = self.active_project()
        user = self.users['member']
        user.display_name = '吴兴广'
        user.save()
        other = User.objects.create_user(username='other-wu', display_name=user.display_name, role='member')
        for active in (True, False):
            other.is_active = active
            other.save()
            ambiguous = self.preview('tasks', [[project.code, 'design', '不能猜测', '', '吴兴广', TODAY]])
            self.assertFalse(ambiguous.data['can_import'], ambiguous.data)
            self.assertIsNone(ambiguous.data['token'])
            self.assertIn('对应多名人员', ambiguous.data['errors'][0]['message'])
        # The same string can be one person's account and another person's name.
        other.display_name = user.username
        other.save()
        collision = self.preview('tasks', [[project.code, 'design', '账号姓名冲突', '', user.username, TODAY]])
        self.assertFalse(collision.data['can_import'], collision.data)
        explicit = self.preview('tasks', [[project.code, 'design', '明确本人', '', '吴兴广（member）', TODAY]])
        self.assertTrue(explicit.data['can_import'], explicit.data)
        self.commit('tasks', explicit)
        self.assertEqual(project.tasks.get().assignee_id, user.pk)

    def test_missing_partial_and_mismatched_person_references_are_rejected(self):
        project = self.active_project()
        self.users['member'].display_name = '吴兴广'
        self.users['member'].save()
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        for value in ('吴兴', '不存在', '错误姓名（member）', '吴兴广（missing）'):
            with self.subTest(value=value):
                preview = self.preview('tasks', [[project.code, 'design', '错误人员', '', value, TODAY]])
                self.assertFalse(preview.data['can_import'], preview.data)
                self.assertIn('未找到人员', preview.data['errors'][0]['message'])
        self.assertFalse(project.tasks.exists())
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(AuditLog.objects.count(), audits)

    def test_confirmation_rejects_identity_reassignment_and_rolls_back_whole_batch(self):
        project = self.active_project()
        user = self.users['member']
        user.display_name = '吴兴广'
        user.save()
        preview = self.preview(
            'tasks',
            [
                [project.code, 'design', '前行也要回滚', '', 'manager', TODAY],
                [project.code, 'design', '不得换人', '', '吴兴广', TODAY],
            ],
        )
        self.assertTrue(preview.data['can_import'], preview.data)
        user.display_name = '更名后的员工'
        user.save()
        other = User.objects.create_user(username='replacement', display_name='吴兴广', role='member')
        project.members.add(other)
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        rejected = self.commit('tasks', preview, status=400)
        self.assertIn('人员对应关系已变化', str(rejected))
        self.assertFalse(project.tasks.exists())
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(AuditLog.objects.count(), audits)

    def test_names_do_not_bypass_project_manager_role_or_sales_ownership(self):
        user = self.users['member']
        user.display_name = '吴兴广'
        user.save()
        invalid = self.preview('projects', [['无经理权限', 'C1', '吴兴广', '', TODAY, '1', '12', '']])
        self.assertFalse(invalid.data['can_import'], invalid.data)
        salesperson = self.users['sales_manager']
        salesperson.display_name = '销售张经理'
        salesperson.save()
        own = self.preview('sales', [['本人销售', 'C1', '销售张经理', '', TODAY, '1', '12']], role='sales_manager')
        self.assertTrue(own.data['can_import'], own.data)
        self.commit('sales', own, role='sales_manager')
        self.assertEqual(SalesOrder.objects.get(name='本人销售').manager_id, salesperson.pk)
        denied = self.preview('sales', [['他人销售', 'C1', 'manager', '', TODAY, '1', '12']], role='sales_manager')
        self.assertFalse(denied.data['can_import'], denied.data)
        self.assertFalse(SalesOrder.objects.filter(name='他人销售').exists())

    def test_task_import_accepts_exact_chinese_username(self):
        user = self.users['member']
        user.username = '中文登录账号'
        user.display_name = '不同的显示姓名'
        user.save()
        project = self.active_project()
        preview = self.preview('tasks', [[project.code, 'design', '中文账号指派', '', user.username, TODAY]])
        self.assertTrue(preview.data['can_import'], preview.data)
        self.assertFalse(project.tasks.exists())
        self.commit('tasks', preview)
        self.assertEqual(project.tasks.get().assignee_id, user.pk)

    def test_task_import_distinguishes_inactive_and_nonparticipant_accounts(self):
        project = self.active_project()
        user = self.users['member']
        user.display_name = '吴兴广'
        user.save()
        rows = [[project.code, 'design', '指派任务', '', user.display_name, TODAY]]
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        valid = self.preview('tasks', rows)
        self.assertTrue(valid.data['can_import'], valid.data)
        user.is_active = False
        user.save()
        inactive = self.preview('tasks', rows)
        self.assertFalse(inactive.data['can_import'], inactive.data)
        self.assertIsNone(inactive.data['token'])
        self.assertEqual(inactive.data['errors'][0]['row'], 2)
        self.assertIn('人员“吴兴广（member）”已停用', inactive.data['errors'][0]['message'])
        rejected = self.commit('tasks', valid, status=400)
        self.assertIn('已停用', str(rejected))
        self.assertEqual(str(rejected['rows'][0]['row']), '2')
        user.is_active = True
        user.save()
        project.members.remove(user)
        outside = self.preview('tasks', rows)
        self.assertFalse(outside.data['can_import'], outside.data)
        self.assertIsNone(outside.data['token'])
        self.assertIn('不是该项目的负责人或成员', outside.data['errors'][0]['message'])
        self.assertFalse(project.tasks.exists())
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(AuditLog.objects.count(), audits)

    def test_exports_all_filtered_rows_sensitive_scope_and_formulas(self):
        for index in range(25):
            Item.objects.create(code=f'E{index}', name='=HYPERLINK("bad")')
        response = self.clients['member'].get('/api/business/items/export/', {'search': 'E', 'page_size': 1})
        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 26)
        self.assertTrue(rows[1][2].startswith("'="))
        xlsx = self.clients['admin'].get('/api/business/items/export/', {'file_format': 'xlsx'})
        book = load_workbook(io.BytesIO(xlsx.content))
        self.assertFalse(any(cell.data_type == 'f' for row in book.active for cell in row))
        project = self.active_project()
        project.members.clear()
        response = self.clients['member'].get('/api/business/projects/export/')
        self.assertNotIn(project.code, response.content.decode())
        project.members.add(self.users['member'])
        response = self.clients['member'].get('/api/business/projects/export/')
        self.assertIn(project.code, response.content.decode())
        self.assertNotIn('合同金额', response.content.decode())
        self.assertEqual(self.clients['member'].get('/api/business/payments/export/').status_code, 403)
        self.assertEqual(
            self.clients['manager'].get('/api/business/reports/', {'file_format': 'xlsx'}).status_code, 403
        )
        report = self.clients['admin'].get('/api/business/reports/', {'file_format': 'csv', 'search': project.code})
        self.assertIn(project.code, report.content.decode())
        self.assertIn('当前筛选全部项目', report.content.decode())
        with patch('apps.business.services.tabular.MAX_EXPORT', 1):
            self.assertEqual(self.clients['admin'].get('/api/business/items/export/').status_code, 400)

    def test_templates_and_malformed_files(self):
        for resource in transfers.SCHEMAS:
            for fmt in ['csv', 'xlsx']:
                response = self.clients['admin'].get(f'/api/business/{resource}/import-template/', {'file_format': fmt})
                self.assertEqual(response.status_code, 200, resource)
        formula = self.preview('items', [['X', '=1+1', '', '件']], xlsx=True)
        self.assertEqual(formula.status_code, 400)
        missing = self.preview('payments', [['999999', '1', TODAY, '不存在']])
        self.assertEqual(missing.status_code, 200)
        self.assertFalse(missing.data['can_import'])

    def test_model_validation_errors_keep_file_rows_and_roll_back_preview(self):
        counter = CodeRule.objects.get(key='partner').counter
        receipts = ActionReceipt.objects.count()
        audits = AuditLog.objects.count()
        template = self.clients['admin'].get('/api/business/partners/import-template/')
        headers = next(csv.reader(io.StringIO(template.content.decode('utf-8-sig'))))
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(
            [
                {'往来单位': '错误类型一', '类型': 'unknown'},
                {'往来单位': '有效但只预览', '类型': '供应商'},
                {'往来单位': '错误类型二', '类型': 'client'},
            ]
        )
        preview = self.clients['admin'].post(
            '/api/business/partners/import-file/',
            {'file': SimpleUploadedFile('partners.csv', stream.getvalue().encode())},
            format='multipart',
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertFalse(preview.data['can_import'])
        self.assertIsNone(preview.data['token'])
        self.assertEqual([error['row'] for error in preview.data['errors']], [2, 4])
        self.assertTrue(all('类型' in error['message'] for error in preview.data['errors']))
        self.assertTrue(all('kind' not in error['message'] for error in preview.data['errors']))
        self.assertFalse(Partner.objects.filter(name__in=['错误类型一', '有效但只预览', '错误类型二']).exists())
        self.assertEqual(CodeRule.objects.get(key='partner').counter, counter)
        self.assertEqual(ActionReceipt.objects.count(), receipts)
        self.assertEqual(AuditLog.objects.count(), audits)

    def test_import_reference_and_field_errors_use_chinese_names(self):
        project = self.active_project()
        cases = [
            ('stocks', [['MISSING', '主仓', '1', '1', '']], '物料'),
            ('projects', [['项目', 'MISSING', 'manager', '', TODAY, '1', '12', '']], '往来单位'),
            ('tasks', [['MISSING', 'design', '设计', '', 'member', TODAY]], '项目'),
            ('tasks', [[project.code, 'design', '设计', '', 'member', '20']], '期限'),
            ('stocks', [['I1', '主仓', 'bad', '1', '']], '数量'),
        ]
        for resource, rows, label in cases:
            with self.subTest(resource=resource, label=label):
                preview = self.preview(resource, rows)
                self.assertFalse(preview.data['can_import'], preview.data)
                error = preview.data['errors'][0]
                self.assertEqual(error['row'], 2)
                self.assertIn(label, error['message'])
                for internal in ['item：', 'partner：', 'project：', 'due_date：', 'quantity：']:
                    self.assertNotIn(internal, error['message'])
        self.assertFalse(project.tasks.exists())
        self.assertFalse(Stock.objects.exists())

    def downloaded_preview(self, resource, rows, fmt='csv', *, previous=False):
        response = self.clients['admin'].get(f'/api/business/{resource}/import-template/', {'file_format': fmt})
        self.assertEqual(response.status_code, 200)
        if fmt == 'csv':
            headers = next(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
            stream = io.StringIO()
            writer = csv.writer(stream)
            writer.writerow([label for label, _ in transfers.PREVIOUS_SCHEMAS[resource]] if previous else headers)
            writer.writerows([[row.get(label, '') for label in headers] for row in rows])
            content = stream.getvalue().encode()
        else:
            book = load_workbook(io.BytesIO(response.content))
            headers = [cell.value for cell in book['数据'][1]]
            for row in rows:
                book['数据'].append([row.get(label, '') for label in headers])
            if previous:
                for column, (label, _) in enumerate(transfers.PREVIOUS_SCHEMAS[resource], 1):
                    book['数据'].cell(1, column).value = label
            stream = io.BytesIO()
            book.save(stream)
            book.close()
            content = stream.getvalue()
        response = self.clients['admin'].post(
            f'/api/business/{resource}/import-file/',
            {'file': SimpleUploadedFile('filled.' + fmt, content)},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data['can_import'], response.data)
        self.assertEqual([column['label'] for column in response.data['columns']], headers)
        return response

    def test_downloaded_templates_match_schema_and_preview_for_every_resource(self):
        project = self.active_project(equipment=2)
        task = Task.objects.create(project=project, title='模板工时任务', kind='design', assignee=self.users['admin'])
        self.commit('entries', self.preview('entries', [[project.code, '待付费用', '100', TODAY]]))
        entry = project.entries.get(kind='expense')
        self.commit('stocks', self.preview('stocks', [['I1', '领料仓', '5', '2', '期初']]))
        stock = Stock.objects.get(location='领料仓')
        BOMLine.objects.create(project=project, item=self.item, quantity=2)
        for stage in ['design', 'assembly', 'test']:
            Task.objects.create(project=project, title=stage, kind=stage, assignee=self.users['admin'], status='done')
        task.status = 'done'
        task.save()
        self.commit('moves', self.preview('moves', [[project.code, str(stock.pk), '', '1', '发货准备']]))
        samples = {
            'items': {
                '物料编码': 'NEW-TEMPLATE',
                '物料名称': '模板物料',
                '规格': '规格A',
                '品牌': '品牌B',
                '物料类别': '标准件',
                '单位': '台',
            },
            'partners': {'往来单位': '模板供应商', '类型': '供应商', '采购账期': '月结30天', '电话': '00123'},
            'sales': {'销售名称': '模板销售', '客户编码': 'C1', '负责人姓名或账号': 'manager', '设备数量': '1'},
            'projects': {'项目名称': '模板项目', '客户编码': 'C1', '负责人姓名或账号': 'manager', '设备数量': '1'},
            'purchases': {
                '分组号': 'G1',
                '项目编号': project.code,
                '供应商编码': 'S1',
                '交期': TODAY,
                '物料编码': 'I1',
                '数量': '1',
                '含税单价': '12.34',
                '采购账期': '现付',
            },
            'tasks': {'任务名称': '新设计', '阶段': '设计', '执行人姓名或账号': 'member', '项目编号': project.code},
            'stocks': {'物料编码': 'I1', '库位': '新期初仓', '库存数量': '2', '单位成本': '3.45', '原因': '期初'},
            'entries': {'项目编号': project.code, '款项': '模板差旅', '原金额': '10.25', '最近待付期限': TODAY},
            'payments': {
                '款项ID': str(entry.pk),
                '金额': '10',
                '日期': TODAY,
                '说明': '实际支付',
                '结算方式': '银行转账',
            },
            'time': {'任务ID': str(task.pk), '人员姓名或账号': 'admin', '日期': TODAY, '工时': '1.25', '说明': '设计'},
            'deliveries': {
                '项目编号': project.code,
                '设备数量': '1',
                '发货日': TODAY,
                '安装人姓名或账号': 'member',
                '验收人姓名或账号': 'manager',
            },
            'moves': {'库存ID': str(stock.pk), '数量': '1', '说明': '生产领料', '项目编号': project.code},
        }
        self.assertEqual(set(samples), set(transfers.SCHEMAS))
        for resource, row in samples.items():
            for fmt in ['csv', 'xlsx']:
                with self.subTest(resource=resource, fmt=fmt):
                    layout = self.clients['admin'].get(f'/api/business/{resource}/import-schema/')
                    self.assertEqual(layout.status_code, 200)
                    preview = self.downloaded_preview(resource, [row], fmt)
                    self.assertEqual(preview.data['columns'], layout.data['columns'])
                    self.assertEqual(preview.data['row_count'], 1)
                    self.assertEqual(preview.data['rows'][0]['row'], 2)
                    if any(key in transfers.PERSON_LABELS for _, key in transfers.SCHEMAS[resource][1]):
                        old = self.downloaded_preview(resource, [row], fmt, previous=True)
                        self.assertEqual(old.data['rows'], preview.data['rows'])
        # Confirm the new order does not swap brand/category/unit, including Chinese enums.
        result = self.commit('items', self.downloaded_preview('items', [samples['items']], 'xlsx'))
        item = Item.objects.get(pk=result['results'][0]['id'])
        self.assertEqual((item.brand, item.part_type, item.unit), ('品牌B', 'standard', '台'))
        self.assertEqual(self.clients['member'].get('/api/business/items/import-schema/').status_code, 403)

    def test_every_template_includes_downloadable_field_guidance_without_example_records(self):
        for resource in transfers.SCHEMAS:
            with self.subTest(resource=resource):
                layout = self.clients['admin'].get(f'/api/business/{resource}/import-schema/').data
                html = self.clients['admin'].get(f'/api/business/{resource}/import-template/', {'file_format': 'html'})
                self.assertEqual(html.status_code, 200)
                self.assertContains(html, layout['title'] + '导入模板填写说明')
                self.assertContains(html, '第一次填写，照这三步做')
                has_options = any(column['options'] for column in layout['columns'])
                self.assertContains(html, '表头不显示下拉箭头' if has_options else '本模板没有固定选项')
                for key in ('prepare', 'minimum', 'example'):
                    self.assertTrue(layout['quick_start'][key], resource)
                    self.assertContains(html, layout['quick_start'][key])
                self.assertContains(html, layout['entry'])
                self.assertIn('attachment;', html['Content-Disposition'])
                response = self.clients['admin'].get(
                    f'/api/business/{resource}/import-template/', {'file_format': 'xlsx'}
                )
                with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                    root = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
                    stored_validations = root.findall('.//{*}dataValidation')
                with io.BytesIO(response.content) as stream:
                    book = load_workbook(stream)
                    self.assertEqual(book.sheetnames, ['数据', GUIDE_SHEET])
                    self.assertEqual(book['数据'].max_row, 1)
                    self.assertEqual(book[GUIDE_SHEET]['A1'].value, GUIDE_MARKER)
                    self.assertEqual(book.active.title, GUIDE_SHEET)
                    self.assertEqual(book[GUIDE_SHEET]['B6'].hyperlink.location, "'数据'!A2")
                    self.assertEqual(book[GUIDE_SHEET]['B4'].value, layout['quick_start']['prepare'])
                    self.assertEqual(book[GUIDE_SHEET]['B5'].value, layout['quick_start']['minimum'])
                    self.assertIn(
                        '再点右侧▼' if has_options else '本模板没有固定下拉选项', book[GUIDE_SHEET]['B7'].value
                    )
                    validations = list(book['数据'].data_validations.dataValidation)
                    self.assertEqual(len(validations), len(layout['columns']))
                    for index, column in enumerate(layout['columns'], 1):
                        for key in ('required', 'hint', 'example'):
                            self.assertTrue(column[key], (resource, column))
                            self.assertIn(column[key], book['数据'].cell(1, index).comment.text)
                        self.assertContains(html, column['label'])
                        if column['key'] in transfers.PERSON_LABELS:
                            self.assertIn('姓名或账号', column['label'])
                            self.assertIn('可填姓名或登录账号', column['hint'])
                            self.assertIn('姓名（用户名）', column['hint'])
                            letter = book['数据'].cell(1, index).column_letter
                            self.assertEqual(book['数据'].column_dimensions[letter].number_format, '@')
                        validation = validations[index - 1]
                        letter = book['数据'].cell(1, index).column_letter
                        self.assertEqual(column['cell'], f'{letter}2')
                        self.assertTrue(column['source'])
                        self.assertEqual(book[GUIDE_SHEET].cell(13 + index, 1).hyperlink.location, f"'数据'!{letter}2")
                        self.assertIn(column['source'], book[GUIDE_SHEET].cell(13 + index, 2).value)
                        if column['key'] in ('code', 'customer', 'supplier', 'item', 'project', 'phone', 'reference'):
                            self.assertEqual(book['数据'].column_dimensions[letter].number_format, '@')
                        self.assertEqual(str(validation.sqref), f'{letter}2:{letter}1001')
                        self.assertTrue(validation.showInputMessage)
                        self.assertIn(column['required'], validation.prompt)
                        self.assertLessEqual(len(validation.prompt), 255)
                        # openpyxl's alias default overwrites allowBlank on reload; verify the native file flag.
                        self.assertEqual(stored_validations[index - 1].get('allowBlank'), '1')
                        if column['options']:
                            self.assertEqual(validation.type, 'list')
                            self.assertFalse(validation.showDropDown)  # False keeps the Excel arrow visible.
                            self.assertEqual(validation.formula1, '"' + ','.join(column['options']) + '"')
                            self.assertLessEqual(len(validation.formula1), 255)
                            self.assertEqual(validation.showErrorMessage, not column['allow_custom'])
                            self.assertEqual(validation.errorStyle, 'stop')
                        else:
                            self.assertEqual(validation.type, 'custom')
                            self.assertFalse(validation.showErrorMessage)
                    self.assertTrue(book['数据'].conditional_formatting)
                    book.close()
        for fmt in ('html', 'xlsx', 'csv'):
            self.assertEqual(
                self.clients['member'].get('/api/business/projects/import-template/', {'file_format': fmt}).status_code,
                403,
            )

    def test_fixed_dropdown_choices_are_importable_and_custom_units_remain_allowed(self):
        expected = {
            ('items', 'product_category'): list(Item.ProductCategory.labels),
            ('items', 'part_type'): ['标准件', '非标件'],
            ('partners', 'kind'): ['客户', '供应商', '客户及供应商'],
            ('partners', 'payment_term'): [
                '指定日期',
                '现付',
                '月结30天',
                '月结60天',
                '月结90天',
                '月结120天',
                '自定义月结',
            ],
            ('purchases', 'payment_term'): [
                '指定日期',
                '现付',
                '月结30天',
                '月结60天',
                '月结90天',
                '月结120天',
                '自定义月结',
            ],
            ('tasks', 'kind'): ['设计', '装配', '调试'],
            ('payments', 'method'): ['银行转账', '现金', '其他'],
        }
        for (resource, key), options in expected.items():
            with self.subTest(resource=resource, key=key):
                layout = self.clients['admin'].get(f'/api/business/{resource}/import-schema/').data
                column = next(column for column in layout['columns'] if column['key'] == key)
                self.assertEqual(column['options'], options)
                self.assertFalse(column['allow_custom'])
        samples = [
            {
                '物料名称': f'下拉选项验证{code}',
                '规格': '实际型号规格',
                '产品编码类别': label,
                '图号': f'TEST-{code}' if code.startswith('1') else '',
                '物料类别': '非标件' if code.startswith('1') else '标准件',
                '单位': '托盘',
            }
            for code, label in Item.ProductCategory.choices
        ]
        initial = Item.objects.count()
        preview = self.downloaded_preview('items', samples, 'xlsx')
        self.assertEqual(preview.data['row_count'], 8)
        self.assertEqual(Item.objects.count(), initial)
        result = self.commit('items', preview)
        items = Item.objects.filter(pk__in=[row['id'] for row in result['results']])
        self.assertEqual(set(items.values_list('product_category', flat=True)), set(Item.ProductCategory.values))
        self.assertEqual(set(items.values_list('unit', flat=True)), {'托盘'})

    def test_purchase_preview_retains_file_rows_and_grouped_confirmation(self):
        project = self.active_project()
        row = {
            '分组号': 'GROUP',
            '项目编号': project.code,
            '供应商编码': 'S1',
            '交期': TODAY,
            '物料编码': 'I1',
            '数量': '2',
            '含税单价': '10',
        }
        other = Item.objects.create(code='SECOND', name='第二物料')
        preview = self.downloaded_preview(
            'purchases', [row, {**row, '物料编码': other.code, '数量': '3', '含税单价': '20'}]
        )
        self.assertEqual((preview.data['count'], preview.data['row_count']), (1, 2))
        self.assertEqual([r['row'] for r in preview.data['rows']], [2, 3])
        self.assertNotIn('lines', preview.data['rows'][0]['data'])
        self.assertEqual(preview.data['rows'][1]['data']['item'], 'SECOND')
        self.commit('purchases', preview)
        purchase = PurchaseOrder.objects.get()
        self.assertEqual(purchase.status, 'draft')
        self.assertEqual(list(purchase.lines.order_by('id').values_list('quantity', 'unit_price')), [(2, 10), (3, 20)])
